"""
LawFlow 智律工作台 - 阶段一：立案文书自动化 + 中国司法实务计算闭环
核心接口：
1. POST /api/v1/phase1/generate-litigation-package:
   一站式生成三大高频案由全套诉讼文书（起诉状、委托书、法定代表人身份证明、送达地址确认书、财产保全申请书、仲裁申请书及补偿金清单），
   调用宿主机 gemma4:12b 自动改写为引用《民法典》或《劳动合同法》的严谨【事实与理由】，并输出 .docx 与 .pdf。
2. 法律精算引擎接口：
   - /api/v1/phase1/calc/rmb-capital: 人民币大写（基于 cn2an）
   - /api/v1/phase1/calc/idcard: 18位身份证校验及法定退休年龄核定
   - /api/v1/phase1/calc/uscc: 18位统一社会信用代码 GB 32100-2015 校验
   - /api/v1/phase1/calc/lpr-interest: 央行 LPR 历期计息及起诉状利息明细说明
   - /api/v1/phase1/calc/labor-arbitration: 劳动争议 N / N+1 / 2N 补偿金精算
   - /api/v1/phase1/calc/court-fees: 诉讼费用交纳办法分段计费
"""

import os
import re
import shutil
import subprocess
import time
import uuid
import zipfile
import logging
from datetime import datetime
from typing import List, Optional, Dict, Any, Literal, Tuple
from fastapi import APIRouter, HTTPException, Query, Request
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field, field_validator, model_validator
from docxtpl import DocxTemplate
import httpx
import pymupdf as fitz

from gateway.utils.cn_legal_calc import (
    rmb_to_capital,
    validate_and_parse_idcard,
    validate_and_parse_uscc,
    calculate_court_fees,
    calculate_preservation_insurance,
    calculate_lpr_interest,
    calculate_labor_compensation,
    calculate_statute_of_limitations,
    calculate_preservation_renewal,
    verify_statutory_validity,
    calculate_settlement_npv_and_recovery,
    calculate_appeal_deadline,
    calculate_delayed_performance_interest,
    calculate_penalty_adjustment_reference,
)
from gateway.utils.text_utils import (
    normalize_court_name,
    normalize_arbitration_committee,
    extract_agreed_annual_rate,
    extract_first_date_iso,
    format_date_cn,
)


def _http_422(detail: str) -> HTTPException:
    return HTTPException(status_code=422, detail=detail)
from gateway.core.ollama_client import ollama_client
from gateway.core.security import sanitize_filename, safe_join, validate_download_filename
from gateway.utils.legal_prompts import FACTS_POLISH_SYSTEM_PROMPT
from gateway.utils.appeal_and_precedent_engine import analyze_first_instance_judgment_and_precedents

logger = logging.getLogger("lawflow.phase1")
router = APIRouter(tags=["阶段一：立案文书自动化与法律精算"])

from gateway.core.paths import get_base_dir, get_generated_dir

PLACEHOLDER = "【待填写】"


def _templates_dir() -> str:
    return str(get_base_dir() / "templates")

STIRLING_PDF_URL = os.getenv("STIRLING_PDF_URL", "http://127.0.0.1:8081").rstrip("/")


_SOFFICE_CACHED_PATH: Optional[str] = None
_SOFFICE_SEARCHED: bool = False

def _find_soffice() -> Optional[str]:
    """定位 LibreOffice soffice 可执行文件（带极速内存缓存防多次遍历 PATH 超时）。"""
    global _SOFFICE_CACHED_PATH, _SOFFICE_SEARCHED
    if _SOFFICE_SEARCHED:
        return _SOFFICE_CACHED_PATH

    for static_cand in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/usr/bin/soffice",
        "/usr/bin/libreoffice",
    ):
        if os.path.isfile(static_cand):
            _SOFFICE_CACHED_PATH = static_cand
            _SOFFICE_SEARCHED = True
            return static_cand

    for cand in (
        shutil.which("soffice"),
        shutil.which("libreoffice"),
    ):
        if cand and os.path.isfile(cand):
            _SOFFICE_CACHED_PATH = cand
            _SOFFICE_SEARCHED = True
            return cand

    _SOFFICE_CACHED_PATH = None
    _SOFFICE_SEARCHED = True
    return None


def _convert_docx_via_soffice(docx_path: str, pdf_path: str) -> bool:
    soffice = _find_soffice()
    if not soffice:
        return False
    out_dir = os.path.dirname(os.path.abspath(pdf_path)) or "."
    try:
        proc = subprocess.run(
            [
                soffice,
                "--headless",
                "--nologo",
                "--nofirststartwizard",
                "--convert-to",
                "pdf",
                "--outdir",
                out_dir,
                os.path.abspath(docx_path),
            ],
            capture_output=True,
            text=True,
            timeout=120,
        )
        if proc.returncode != 0:
            logger.warning(f"LibreOffice 转换失败: {proc.stderr}")
            return False
        produced = os.path.join(out_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
        if os.path.abspath(produced) != os.path.abspath(pdf_path):
            if os.path.isfile(produced):
                shutil.move(produced, pdf_path)
        return os.path.isfile(pdf_path) and os.path.getsize(pdf_path) > 100
    except Exception as e:
        logger.warning(f"LibreOffice 转换异常: {e}")
        return False


def _extract_docx_paragraphs(docx_path: str, text_summary: str = "") -> List[str]:
    import docx as python_docx
    d = python_docx.Document(docx_path)
    lines: List[str] = []
    for p in d.paragraphs:
        if p.text.strip():
            lines.append(p.text)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append("  ".join(cells))
    if not lines:
        if text_summary:
            lines = [text_summary]
        else:
            lines = ["司法文书内容"]
    return lines


def _pymupdf_fit_lines(doc: "fitz.Document", text_rect, lines: List[str], fontsize: int = 11) -> int:
    """二分查找能完整写入当前新页的最大行数；写入后返回消耗行数。"""
    if not lines:
        return 0
    page_width, page_height = 595, 842
    lo, hi = 1, len(lines)
    best = 1
    while lo <= hi:
        mid = (lo + hi) // 2
        probe = doc.new_page(width=page_width, height=page_height)
        rc = probe.insert_textbox(
            text_rect, "\n".join(lines[:mid]), fontsize=fontsize, fontname="china-s", align=0
        )
        doc.delete_page(-1)
        if rc >= 0:
            best = mid
            lo = mid + 1
        else:
            hi = mid - 1
    page = doc.new_page(width=page_width, height=page_height)
    page.insert_textbox(
        text_rect, "\n".join(lines[:best]), fontsize=fontsize, fontname="china-s", align=0
    )
    return best


def _pymupdf_paginated_render(paragraphs: List[str], pdf_path: str) -> bool:
    """
    按行分页写入 PDF：检查 insert_textbox 返回值，溢出则新开一页继续，直至写完。
    """
    page_width, page_height = 595, 842
    margin_x, margin_top, margin_bottom = 54, 72, 55
    fontsize = 11
    all_lines: List[str] = []
    for p in paragraphs:
        text = p if p is not None else ""
        if not text:
            all_lines.append("")
            continue
        for i in range(0, len(text), 38):
            all_lines.append(text[i : i + 38])
        all_lines.append("")

    doc = fitz.open()
    text_rect = fitz.Rect(margin_x, margin_top, page_width - margin_x, page_height - margin_bottom)
    remaining = all_lines
    while remaining:
        consumed = _pymupdf_fit_lines(doc, text_rect, remaining, fontsize=fontsize)
        if consumed <= 0:
            # 保底：强制写一行避免死循环
            page = doc.new_page(width=page_width, height=page_height)
            page.insert_textbox(text_rect, remaining[0], fontsize=fontsize, fontname="china-s", align=0)
            consumed = 1
        remaining = remaining[consumed:]

    total_pages = doc.page_count
    for pno in range(total_pages):
        page = doc[pno]
        from gateway.utils.disclaimer import disclaimer_pdf_footer_line
        footer_rect = fitz.Rect(margin_x, page_height - 48, page_width - margin_x, page_height - 12)
        footer_str = disclaimer_pdf_footer_line(page_no=pno + 1, total_pages=total_pages)
        page.insert_textbox(footer_rect, footer_str, fontsize=6.5, fontname="china-s", align=1)

    doc.save(pdf_path)
    doc.close()
    return total_pages > 0 and os.path.isfile(pdf_path)


# ==============================================================================
# 文书转换辅助函数：将 Word 转换为 PDF
# ==============================================================================
# 最近一次 convert_docx_to_pdf 使用的引擎（供一键成套返回 engine_status）
_LAST_PDF_ENGINE: str = "local_pymupdf"


def get_last_pdf_engine() -> str:
    return _LAST_PDF_ENGINE


async def convert_docx_to_pdf(docx_path: str, pdf_path: str, title: str = "司法公文", text_summary: str = "") -> bool:
    """
    优先 LibreOffice (soffice)；其次 Stirling-PDF（仅配置的一个地址，且需探测缓存可用）；
    最后 PyMuPDF 按行分页兜底。全部失败时返回 False，不得生成缺页/空壳 PDF。
    """
    global _LAST_PDF_ENGINE
    from gateway.utils.disclaimer import stamp_disclaimer_on_pdf
    from gateway.core.service_probe import is_stirling_available, stirling_pdf_url

    # 1. LibreOffice
    if _convert_docx_via_soffice(docx_path, pdf_path):
        logger.info(f"LibreOffice 成功将 {docx_path} 转换为 {pdf_path}")
        stamp_disclaimer_on_pdf(pdf_path)
        _LAST_PDF_ENGINE = "libreoffice"
        return True

    # 2. Stirling-PDF：仅尝试配置的 STIRLING_PDF_URL；探测不可用则跳过，避免 10s×多地址超时
    if is_stirling_available():
        url = (STIRLING_PDF_URL or stirling_pdf_url()).rstrip("/")
        try:
            async with httpx.AsyncClient(timeout=5.0) as client:
                with open(docx_path, "rb") as f:
                    files = {
                        "fileInput": (
                            os.path.basename(docx_path),
                            f,
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        )
                    }
                    resp = await client.post(f"{url}/api/v1/convert/file/pdf", files=files)
                    if resp.status_code == 200 and len(resp.content) > 100:
                        with open(pdf_path, "wb") as pdf_file:
                            pdf_file.write(resp.content)
                        logger.info(f"Stirling-PDF 成功将 {docx_path} 转换为 {pdf_path}")
                        stamp_disclaimer_on_pdf(pdf_path)
                        _LAST_PDF_ENGINE = "stirling"
                        return True
        except Exception as e:
            logger.warning("Stirling-PDF 转换失败，回退本地引擎: %s", e)

    # 3. PyMuPDF 按行分页兜底
    try:
        paragraphs = _extract_docx_paragraphs(docx_path, text_summary=text_summary)
        ok = _pymupdf_paginated_render(paragraphs, pdf_path)
        if ok:
            logger.info(f"PyMuPDF 分页渲染成功: {pdf_path}")
            stamp_disclaimer_on_pdf(pdf_path)
            _LAST_PDF_ENGINE = "local_pymupdf"
            return True
        logger.error("PyMuPDF 分页渲染失败：未生成有效 PDF")
        _LAST_PDF_ENGINE = "local_pymupdf"
        return False
    except Exception as e:
        logger.error(f"PDF 转换异常: {e}")
        _LAST_PDF_ENGINE = "local_pymupdf"
        if os.path.isfile(pdf_path):
            try:
                os.remove(pdf_path)
            except OSError:
                pass
        return False


def _highlight_placeholders_in_docx(docx_path: str) -> None:
    """将文档中【待填写】文字标为黄色高亮。"""
    try:
        import docx as python_docx
        from docx.enum.text import WD_COLOR_INDEX
        d = python_docx.Document(docx_path)
        targets = list(d.paragraphs)
        for table in d.tables:
            for row in table.rows:
                for cell in row.cells:
                    targets.extend(cell.paragraphs)
        for p in targets:
            for run in p.runs:
                if PLACEHOLDER in (run.text or ""):
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        d.save(docx_path)
    except Exception as e:
        logger.warning(f"黄色高亮占位符失败: {e}")


def _or_placeholder(value: Optional[str], missing_fields: List[str], field_name: str) -> str:
    if value is None or str(value).strip() == "":
        missing_fields.append(field_name)
        return PLACEHOLDER
    s = str(value).strip()
    # 禁止字段名充当内容
    if s == field_name or s in ("法定代表人", "主要负责人", "性别", "出生日期"):
        missing_fields.append(field_name)
        return PLACEHOLDER
    return s


def _resolve_party(party: "PartyInfo", role: str, missing_fields: List[str]) -> Dict[str, Any]:
    """按 party_type 解析当事人展示字段。"""
    pid = (party.id_card_or_uscc or "").strip()
    party_type = party.party_type
    id_info = validate_and_parse_idcard(pid) if pid else {"valid": False}
    uscc_info = validate_and_parse_uscc(pid) if pid else {"valid": False}

    if party_type is None:
        if uscc_info.get("valid"):
            party_type = "legal"
        elif id_info.get("valid"):
            party_type = "natural"
        else:
            # 18 位且含字母倾向法人；否则自然人占位
            if len(pid) == 18 and any(c.isalpha() for c in pid):
                party_type = "legal"
            else:
                party_type = "natural"

    result: Dict[str, Any] = {
        "name": _or_placeholder(party.name, missing_fields, f"{role}.name"),
        "address": _or_placeholder(party.address, missing_fields, f"{role}.address"),
        "phone": _or_placeholder(party.phone, missing_fields, f"{role}.phone"),
        "party_type": party_type,
        "ethnicity": party.ethnicity or "汉族",
        "id_card_or_uscc": pid or PLACEHOLDER,
    }
    if not pid:
        missing_fields.append(f"{role}.id_card_or_uscc")

    if party_type == "natural":
        parsed_gender = id_info.get("gender") if id_info.get("valid") else None
        parsed_birth = None
        if id_info.get("valid") and id_info.get("birth_date"):
            # 1990-03-07 → 1990年3月7日
            parsed_birth = format_date_cn(id_info["birth_date"])

        if party.gender and parsed_gender and party.gender != parsed_gender:
            raise _http_422(
                f"{role} 性别填写「{party.gender}」与身份证解析结果「{parsed_gender}」冲突"
            )
        if party.birth and parsed_birth:
            # 宽松比较：去掉 0 填充差异
            def _norm_birth(s: str) -> str:
                return re.sub(r"年0*", "年", re.sub(r"月0*", "月", s.replace("日", "")))
            if _norm_birth(party.birth) != _norm_birth(parsed_birth):
                raise _http_422(
                    f"{role} 出生日期填写「{party.birth}」与身份证解析结果「{parsed_birth}」冲突"
                )

        result["gender"] = party.gender or parsed_gender or PLACEHOLDER
        result["birth"] = party.birth or parsed_birth or PLACEHOLDER
        if result["gender"] == PLACEHOLDER:
            missing_fields.append(f"{role}.gender")
        if result["birth"] == PLACEHOLDER:
            missing_fields.append(f"{role}.birth")
        result["id_label"] = "公民身份号码"
        result["legal_rep"] = ""
        result["rep_position"] = ""
    else:
        result["gender"] = ""
        result["birth"] = ""
        result["id_label"] = "统一社会信用代码"
        result["legal_rep"] = _or_placeholder(party.legal_rep, missing_fields, f"{role}.legal_rep")
        result["rep_position"] = _or_placeholder(party.rep_position, missing_fields, f"{role}.rep_position")

    return result


def _build_claims_for_cause(
    req: "GenerateLitigationPackageReq",
    agreed_rate: Optional[float],
    start_date_iso: str,
    start_date_display: str,
    interest_meta: Dict[str, Any],
) -> List[str]:
    """用户诉请原样写入；利息诉请按约定/上限格式化；不追加暂计金额。"""
    claims: List[str] = []
    for raw in req.claims or []:
        c = str(raw).strip()
        if not c:
            continue
        # 劳动争议用语规范化
        if "劳动" in req.cause_of_action:
            c = c.replace("判令被告", "裁决被申请人").replace("判令", "裁决")
            c = c.replace("原告", "申请人").replace("被告", "被申请人")
        # 利息诉请：有约定利率时改写为规范表述
        if "利息" in c and ("借贷" in req.cause_of_action or agreed_rate is not None):
            rate = interest_meta.get("applied_rate")
            if rate is not None:
                note = ""
                if interest_meta.get("capped"):
                    note = f"（约定年利率{agreed_rate:g}%超过法定上限，按上限年利率{rate:g}%主张）"
                c = (
                    f"判令被告支付利息（以{req.claim_amount:,.2f}元为基数，按年利率{rate:g}%，"
                    f"自{start_date_display}起计算至实际清偿之日止）{note}"
                )
        # 律师费旁注
        if "律师费" in c and "需有合同约定" not in c:
            c = f"{c.rstrip('；;。.')}（需有合同约定或法律依据）"
        claims.append(c)
    return claims



# ==============================================================================
# Pydantic 模型定义
# ==============================================================================
class RMBCapitalReq(BaseModel):
    amount: float = Field(..., examples=[123000.50])

class IDCardReq(BaseModel):
    id_card: str = Field(..., examples=["110101199003072375"])
    is_cadre: bool = Field(default=False, description="针对女性，是否为管理/干部岗位（干部55岁退休，工人50岁）")

class USCCReq(BaseModel):
    uscc: str = Field(..., examples=["91440300708461136T"])

class CourtFeeReq(BaseModel):
    claim_amount: float = Field(..., ge=0, examples=[250000.0])
    case_type: str = Field(default="property")
    is_summary_procedure: bool = Field(default=False)
    divorce_base_fee: Optional[float] = Field(
        default=100.0,
        ge=50,
        le=300,
        description="离婚案件基础规费（50至300元之间）",
    )
    preservation_amount: Optional[float] = Field(
        default=None,
        ge=0,
        description="保全申请金额，留空默认等于诉讼标的额",
    )


class PreservationInsuranceReq(BaseModel):
    preservation_amount: float = Field(..., ge=0, examples=[500000.0])
    min_rate: Optional[float] = Field(default=0.0008, ge=0, le=1, description="保费最低费率（默认0.08%）")
    max_rate: Optional[float] = Field(default=0.0015, ge=0, le=1, description="保费最高费率（默认0.15%）")
    min_premium: Optional[float] = Field(default=500.0, ge=0, description="起保最低保费（默认500.00元）")

class LPRInterestReq(BaseModel):
    principal: float = Field(..., examples=[100000.0])
    start_date: str = Field(..., examples=["2023-01-01"])
    end_date: str = Field(..., examples=["2024-01-01"])
    multiplier: float = Field(default=1.0)
    term: Literal["1Y", "5Y"] = Field(default="1Y")
    fixed_annual_rate: Optional[float] = Field(
        default=None,
        ge=0,
        le=100,
        description="年利率，单位为百分数，例如 12 表示年利率 12%",
    )
    rule_mode: str = Field(default="commercial_floating", examples=["commercial_floating"])
    contract_date: Optional[str] = Field(default=None, examples=["2022-03-01"])
    filing_date: Optional[str] = Field(
        default=None,
        examples=["2021-12-01"],
        description="起诉日期（2020-08-20 前成立的民间借贷合同必填）",
    )
    markup_ratio: Optional[float] = Field(
        default=None,
        description="买卖合同固定 LPR 加计比例，须在 0.3–0.5",
    )
    today_date: Optional[str] = Field(default=None)

    @field_validator("term", mode="before")
    @classmethod
    def _normalize_term(cls, v):
        if v is None:
            return "1Y"
        s = str(v).strip().upper()
        if s not in ("1Y", "5Y"):
            raise ValueError('term 仅允许 "1Y" 或 "5Y"')
        return s

    @field_validator("fixed_annual_rate")
    @classmethod
    def _reject_decimal_rate(cls, v):
        if v is not None and 0 < float(v) < 1:
            raise ValueError("疑似填写了小数，请输入百分数（例如 0.24 应写为 24）")
        return v

    @model_validator(mode="after")
    def _check_date_order(self):
        try:
            d0 = datetime.strptime(self.start_date, "%Y-%m-%d").date()
            d1 = datetime.strptime(self.end_date, "%Y-%m-%d").date()
        except ValueError as e:
            raise ValueError(f"日期格式无效: {e}") from e
        if d1 < d0:
            raise ValueError("计息截止日期不能早于起始日期")
        return self

class LaborCalcReq(BaseModel):
    monthly_salary: float = Field(..., examples=[12000.0])
    hire_date: str = Field(..., examples=["2021-03-01"])
    termination_date: str = Field(..., examples=["2024-05-15"])
    local_avg_salary: Optional[float] = Field(default=None)
    termination_reason: Optional[Literal[
        "expiry_termination",
        "employer_dismissal_with_notice",
        "mutual_agreement",
        "employee_resign_employer_fault",
        "unlawful_dismissal",
    ]] = Field(default=None, description="解除/终止事由")

class PartyInfo(BaseModel):
    name: str = Field(..., examples=["张三"])
    id_card_or_uscc: str = Field(default="", examples=["110101199003072375"])
    gender: Optional[str] = Field(default=None, examples=["男"])
    birth: Optional[str] = Field(default=None, examples=["1990年3月7日"])
    ethnicity: Optional[str] = Field(default="汉族", examples=["汉族"])
    address: str = Field(default="", examples=["北京市朝阳区建国门外大街1号"])
    phone: Optional[str] = Field(default=None, examples=["13800000000"])
    legal_rep: Optional[str] = Field(default=None, examples=["李某某"])
    rep_position: Optional[str] = Field(default=None, examples=["执行董事"])
    party_type: Optional[Literal["natural", "legal"]] = Field(
        default=None,
        description="natural=自然人 / legal=法人或非法人组织；未传时自动判断",
    )

    @model_validator(mode="before")
    @classmethod
    def _alias_id_fields(cls, data):
        if isinstance(data, dict):
            if not data.get("id_card_or_uscc"):
                for alt in ("id_number", "id_card", "id", "uscc"):
                    if data.get(alt):
                        data["id_card_or_uscc"] = data[alt]
                        break
        return data

    @model_validator(mode="after")
    def _auto_party_type(self):
        if self.party_type is not None:
            return self
        pid = (self.id_card_or_uscc or "").strip()
        if len(pid) == 18:
            if validate_and_parse_uscc(pid).get("valid"):
                object.__setattr__(self, "party_type", "legal")
            elif validate_and_parse_idcard(pid).get("valid"):
                object.__setattr__(self, "party_type", "natural")
        return self


class EvidenceItem(BaseModel):
    name: str = Field(..., description="证据名称")
    source: Optional[str] = Field(default="", description="证据来源")
    witness_name: Optional[str] = Field(default="", description="证人姓名")
    witness_address: Optional[str] = Field(default="", description="证人住所")
    category: Optional[str] = Field(default="书证")
    purpose: Optional[str] = Field(default="")
    pages: Optional[str] = Field(default="")
    original_or_copy: Optional[str] = Field(default="复印件")


SubDocType = Literal[
    "power_of_attorney",
    "service_address",
    "legal_rep_certificate",
    "property_preservation",
    "evidence_catalog",
    "law_firm_letter",
    "id_credential_attachment",
    "lawyer_license_attachment",
    "evidence_bundle_cover",
]

class GenerateLitigationPackageReq(BaseModel):
    cause_of_action: str = Field(
        default="民间借贷纠纷",
        description="案由：民间借贷纠纷 / 劳动人事争议 / 买卖合同纠纷 / 通用民事纠纷",
        examples=["民间借贷纠纷"]    )
    plaintiff: PartyInfo
    defendant: PartyInfo
    claims: List[str] = Field(
        default_factory=list,
        description="诉讼请求列表（原样写入文书）",
    )
    claim_amount: float = Field(default=200000.0)
    raw_case_notes: str = Field(
        ...,
        description="律师碎片化案情笔记（AI将自动提炼并引用民法典或劳动合同法法条）",
        examples=["2022年3月5日，被告李四因公司资金周转，向原告借款20万元，原告当天通过招商银行网银全额转入被告账户。被告出具借条承诺1年内还清，年化利率15%。借款到期后原告多次通过微信及电话催促还款，被告拒不偿还本金和利息。"]    )
    court_name: str = Field(default="北京市朝阳区人民法院", examples=["北京市朝阳区人民法院"])
    arbitration_committee: Optional[str] = Field(
        default=None,
        description="劳动争议致送机关全称；未填时按地区补全为××劳动人事争议仲裁委员会",
    )
    lawyer_name: str = Field(default="王律师", examples=["王律师"])
    law_firm: str = Field(default="北京市京师律师事务所", examples=["北京市京师律师事务所"])
    lawyer_phone: str = Field(default="13900008888", examples=["13900008888"])
    lawyer_license: str = Field(default="11101201510123456", examples=["11101201510123456"])
    auto_generate_sub_docs: bool = Field(default=True, description="是否一并生成授权委托书、送达地址确认书、身份证明或保全申请书（兼容参数）")
    sub_documents: Optional[List[SubDocType]] = Field(
        default=None,
        description="选择生成的附属文书类型列表，可选：power_of_attorney, service_address, legal_rep_certificate, property_preservation, evidence_catalog, law_firm_letter, id_credential_attachment, lawyer_license_attachment, evidence_bundle_cover",
    )
    archive_format: Literal["spc_mobile", "flat"] = Field(
        default="spc_mobile",
        description="压缩包结构格式：spc_mobile（最高人民法院移动微法院4类目录标准）或 flat（传统单层扁平结构）",
    )
    evidences: Optional[List[EvidenceItem]] = Field(default=None, description="证据和证据来源，证人姓名和住所")
    copies: Optional[int] = Field(default=None, description="起诉状副本份数，默认等于被告人数")
    agreed_annual_rate: Optional[float] = Field(default=None, description="约定年利率（百分数）")
    contract_date: Optional[str] = Field(default=None, description="合同成立日 YYYY-MM-DD")


# ==============================================================================
# 计算接口
# ==============================================================================
@router.post("/api/v1/phase1/calc/rmb-capital", summary="人民币大写转换 (cn2an)")
@router.post("/api/calc/rmb-capital", include_in_schema=False)
async def api_rmb_capital(req: RMBCapitalReq):
    return {"amount": req.amount, "rmb_capital": rmb_to_capital(req.amount)}


@router.post("/api/v1/phase1/calc/idcard", summary="18位身份证解析与法定退休年龄判断")
@router.post("/api/calc/idcard", include_in_schema=False)
async def api_idcard(req: IDCardReq):
    return validate_and_parse_idcard(req.id_card, is_cadre=req.is_cadre)


@router.post("/api/v1/phase1/calc/uscc", summary="18位统一社会信用代码校验与解析 (GB 32100-2015)")
async def api_uscc(req: USCCReq):
    return validate_and_parse_uscc(req.uscc)


@router.post("/api/v1/phase1/calc/court-fees", summary="《诉讼费用交纳办法》分段计费")
@router.post("/api/calc/court-fees", include_in_schema=False)
async def api_court_fees(req: CourtFeeReq):
    try:
        return calculate_court_fees(
            claim_amount=req.claim_amount,
            case_type=req.case_type,
            is_summary_procedure=req.is_summary_procedure,
            divorce_base_fee=req.divorce_base_fee if req.divorce_base_fee is not None else 100.0,
            preservation_amount=req.preservation_amount,
        )
    except ValueError as e:
        raise _http_422(str(e)) from e


@router.post("/api/v1/phase1/calc/preservation-insurance", summary="诉讼保全责任保险（诉责险）市场参考成本估算")
@router.post("/api/calc/preservation-insurance", include_in_schema=False)
async def api_preservation_insurance(req: PreservationInsuranceReq):
    try:
        kwargs: Dict[str, Any] = {"preservation_amount": req.preservation_amount}
        if req.min_rate is not None:
            kwargs["min_rate"] = req.min_rate
        if req.max_rate is not None:
            kwargs["max_rate"] = req.max_rate
        if req.min_premium is not None:
            kwargs["min_premium"] = req.min_premium
        return calculate_preservation_insurance(**kwargs)
    except ValueError as e:
        raise _http_422(str(e)) from e


@router.post("/api/v1/phase1/calc/lpr-interest", summary="央行 LPR 历期计息及诉状利息明细说明")
@router.post("/api/calc/lpr-interest", include_in_schema=False)
async def api_lpr_interest(req: LPRInterestReq):
    try:
        return calculate_lpr_interest(
            principal=req.principal,
            start_date=req.start_date,
            end_date=req.end_date,
            multiplier=req.multiplier,
            term=req.term,
            fixed_annual_rate=req.fixed_annual_rate,
            rule_mode=req.rule_mode,
            contract_date=req.contract_date,
            filing_date=req.filing_date,
            markup_ratio=req.markup_ratio,
            today_date=req.today_date,
        )
    except ValueError as e:
        raise _http_422(str(e)) from e


@router.post("/api/v1/phase1/calc/labor-arbitration", summary="劳动争议经济补偿金精算 (N / N+1 / 2N)")
@router.post("/api/calc/labor-arbitration", include_in_schema=False)
async def api_labor_calc(req: LaborCalcReq):
    try:
        return calculate_labor_compensation(
            monthly_salary=req.monthly_salary,
            hire_date=req.hire_date,
            termination_date=req.termination_date,
            local_avg_salary=req.local_avg_salary,
            termination_reason=req.termination_reason,
        )
    except ValueError as e:
        raise _http_422(str(e)) from e


class StatuteLimitationsReq(BaseModel):
    start_date: str = Field(..., examples=["2023-01-01"])
    cause_type: str = Field(default="general_3y", examples=["general_3y"])
    interruption_events: Optional[List[Dict[str, Any]]] = Field(default=None)
    suspension_events: Optional[List[Dict[str, str]]] = Field(default=None)
    today_date: Optional[str] = Field(default=None)
    act_date: Optional[str] = Field(
        default=None,
        examples=["2021-06-01"],
        description="民事法律行为发生日（撤销权校验 5 年最长期限）",
    )
    damage_date: Optional[str] = Field(default=None, description="权利受到损害之日")
    holidays: Optional[List[str]] = Field(default=None, description="额外法定节假日 YYYY-MM-DD 列表")

class PreservationRenewalReq(BaseModel):
    preservation_type: Literal["bank_deposit", "movable_property", "real_estate_equity"] = Field(
        default="bank_deposit", examples=["bank_deposit"]    )
    start_date: str = Field(..., examples=["2023-05-10"])
    custom_years: Optional[float] = Field(default=None)
    today_date: Optional[str] = Field(default=None)
    internal_reminder_days: int = Field(default=30, ge=1)

class StatuteValidityReq(BaseModel):
    citation_text: str = Field(..., examples=["《中华人民共和国合同法》第一百零七条"])
    incident_date: Optional[str] = Field(default=None, examples=["2023-05-01"])


@router.post("/api/v1/phase1/calc/statute-limitations", summary="诉讼时效与除斥期间精确核算器 (《民法典》第188、194、195、199条)")
@router.post("/api/calc/statute-limitations", include_in_schema=False)
async def api_statute_limitations(req: StatuteLimitationsReq):
    try:
        return calculate_statute_of_limitations(
            start_date=req.start_date,
            cause_type=req.cause_type,
            interruption_events=req.interruption_events,
            suspension_events=req.suspension_events,
            today_date=req.today_date,
            act_date=req.act_date,
            damage_date=req.damage_date,
            holidays=req.holidays,
        )
    except ValueError as e:
        raise _http_422(str(e)) from e


@router.post("/api/v1/phase1/calc/preservation-renewal", summary="财产保全续查封续冻结红线精算 (法释〔2016〕22号第十八条)")
@router.post("/api/calc/preservation-renewal", include_in_schema=False)
async def api_preservation_renewal(req: PreservationRenewalReq):
    try:
        return calculate_preservation_renewal(
            preservation_type=req.preservation_type,
            start_date=req.start_date,
            custom_years=req.custom_years,
            today_date=req.today_date,
            internal_reminder_days=req.internal_reminder_days,
        )
    except ValueError as e:
        raise _http_422(str(e)) from e


@router.post("/api/v1/phase1/calc/statute-validity", summary="规范性法律文件现行有效性与废止拦截核验")
@router.post("/api/calc/statute-validity", include_in_schema=False)
async def api_statute_validity(req: StatuteValidityReq):
    return verify_statutory_validity(
        citation_text=req.citation_text,
        incident_date=req.incident_date
    )


class InstallmentItem(BaseModel):
    months: int = Field(..., examples=[3], description="到期时间距今月数")
    amount: float = Field(..., examples=[500000.0], description="本期计划还款金额")
    note: Optional[str] = Field(default="", description="分期说明")

class AssetPlanInput(BaseModel):
    asset_name: str = Field(default="抵债不动产/商办房产", examples=["朝阳区商办写字楼3套"])
    nominal_asset_valuation: float = Field(..., examples=[4500000.0], description="名义评估值或抵债作价")
    prior_liens: float = Field(default=0.0, examples=[1500000.0], description="在先优先受偿权（工程款优先权、第一顺位银行抵押等）")
    auction_discount_rate: float = Field(default=0.70, examples=[0.70], description="司法拍卖/变现折价率（一拍7折/二拍56折）")
    liquidation_costs_rate: float = Field(default=0.12, examples=[0.12], description="过户税费与持有清算成本率")
    time_to_liquidate_months: int = Field(default=18, examples=[18], description="处置变现周期月数")
    encumbrance_status: Optional[str] = Field(default="", examples=["在先银行抵押150万元"])

class SettlementSimulationReq(BaseModel):
    total_claim_nominal: float = Field(..., gt=0, examples=[5000000.0], description="债权标的名义总额")
    plan_type: Literal["cash_installments", "debt_for_asset", "hybrid"] = Field(
        default="hybrid",
        examples=["hybrid"],
        description="和解方案类型: cash_installments / debt_for_asset / hybrid",
    )
    annual_discount_rate: float = Field(default=0.08, ge=0, examples=[0.08], description="年化资金折现率/资本成本")
    installments: Optional[List[InstallmentItem]] = Field(default=None, description="分期还款计划")
    asset_plan: Optional[AssetPlanInput] = Field(default=None, description="以物抵债方案")


@router.post("/api/v1/phase1/calc/settlement-simulation", summary="商业诉讼和解与以物抵债现值/回收率精算")
@router.post("/api/calc/settlement-simulation", include_in_schema=False)
async def api_settlement_simulation(req: SettlementSimulationReq):
    try:
        installments_list = [i.model_dump() if hasattr(i, "model_dump") else i.dict() for i in req.installments] if req.installments else []
        asset_dict = req.asset_plan.model_dump() if hasattr(req.asset_plan, "model_dump") else (req.asset_plan.dict() if req.asset_plan else {})
        return calculate_settlement_npv_and_recovery(
            total_claim_nominal=req.total_claim_nominal,
            plan_type=req.plan_type,
            annual_discount_rate=req.annual_discount_rate,
            installments=installments_list,
            asset_plan=asset_dict
        )
    except ValueError as e:
        raise _http_422(str(e)) from e


class AppealDeadlineReq(BaseModel):
    service_date: str = Field(..., examples=["2026-11-02"], description="一审裁判送达日期 YYYY-MM-DD")
    judgment_type: Literal["judgment", "ruling"] = Field(..., description="judgment=判决 / ruling=裁定")
    holidays: Optional[List[str]] = Field(default=None, description="额外法定节假日 YYYY-MM-DD 列表；默认读 data/legal/holidays_cn.json")
    treat_weekend_as_holiday: bool = Field(
        default=True,
        description="周六日是否顺延；口径 TODO(律师核实)，默认 True",
    )




class DelayInterestReq(BaseModel):
    principal_unpaid: float = Field(..., gt=0, examples=[100000.0], description="未清偿本金")
    start_date: str = Field(..., examples=["2026-01-01"], description="迟延履行起算日")
    end_date: str = Field(..., examples=["2026-01-31"], description="清偿日/测算截止日")
    daily_rate: Optional[float] = Field(default=None, description="日利率（可选，默认法释〔2014〕8号日万分之一点七五）")


class PenaltyAdjustmentReq(BaseModel):
    agreed_penalty: float = Field(..., ge=0, examples=[140000.0], description="约定的违约金数额")
    actual_loss: float = Field(..., gt=0, examples=[100000.0], description="造成的实际损失数额")


@router.post("/api/v1/phase1/calc/delay-interest", summary="迟延履行期间加倍部分债务利息计算器 (《民事诉讼法》第264条)")
@router.post("/api/calc/delay-interest", include_in_schema=False)
async def api_delay_interest(req: DelayInterestReq):
    try:
        return calculate_delayed_performance_interest(
            principal_unpaid=req.principal_unpaid,
            start_date=req.start_date,
            end_date=req.end_date,
            daily_rate=req.daily_rate,
        )
    except ValueError as e:
        raise _http_422(str(e)) from e


@router.post("/api/v1/phase1/calc/penalty-adjustment", summary="违约金调减测算参考 (《合同编通则解释》第65条)")
@router.post("/api/calc/penalty-adjustment", include_in_schema=False)
async def api_penalty_adjustment(req: PenaltyAdjustmentReq):
    try:
        return calculate_penalty_adjustment_reference(
            agreed_penalty=req.agreed_penalty,
            actual_loss=req.actual_loss,
        )
    except ValueError as e:
        raise _http_422(str(e)) from e


@router.post("/api/v1/phase1/calc/appeal-deadline", summary="上诉期截止日测算（《民事诉讼法》第171条）")
@router.post("/api/calc/appeal-deadline", include_in_schema=False)
async def api_appeal_deadline(req: AppealDeadlineReq):
    try:
        return calculate_appeal_deadline(
            service_date=req.service_date,
            judgment_type=req.judgment_type,
            holidays=req.holidays,
            treat_weekend_as_holiday=req.treat_weekend_as_holiday,
        )
    except ValueError as e:
        raise _http_422(str(e)) from e


# ==============================================================================
# 移动微法院在线诉讼配套支撑材料生成与公文包 ZIP 归档引擎
# ==============================================================================
_PACKAGE_ARCHIVE_REGISTRY: Dict[str, Dict[str, Tuple[str, str]]] = {}


def _create_law_firm_letter_docx(
    output_path: str,
    law_firm: str,
    lawyer_name: str,
    lawyer_license: str,
    lawyer_phone: str,
    client_name: str,
    opponent_name: str,
    cause_of_action: str,
    court_name: str,
    sign_date: str,
) -> None:
    """生成正式《律师事务所公函》（含案由、委派律师、律所、受诉法院）。"""
    import docx as python_docx
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = python_docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.1)
        s.right_margin = Inches(1.1)

    # 律所抬头红标
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_firm = p_header.add_run(law_firm or "北京市京师律师事务所")
    r_firm.font.name = "黑体"
    r_firm.font.size = Pt(20)
    r_firm.font.bold = True
    r_firm.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B)

    # 公函标题
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run("律 师 事 务 所 公 函")
    r_title.font.name = "宋体"
    r_title.font.size = Pt(16)
    r_title.font.bold = True

    # 函号预留
    current_year = datetime.now().year
    p_no = doc.add_paragraph()
    p_no.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_no = p_no.add_run(f"（{current_year}）京律函字第【待填写】号")
    r_no.font.name = "仿宋_GB2312"
    r_no.font.size = Pt(11)

    # 送达法院
    p_court = doc.add_paragraph()
    p_court.paragraph_format.space_before = Pt(8)
    p_court.paragraph_format.space_after = Pt(8)
    r_court = p_court.add_run(f"{court_name}：")
    r_court.font.name = "黑体"
    r_court.font.size = Pt(13)
    r_court.font.bold = True

    # 正文首段
    p_body1 = doc.add_paragraph()
    p_body1.paragraph_format.line_spacing = 1.35
    p_body1.paragraph_format.space_after = Pt(6)
    p_body1.paragraph_format.first_line_indent = Inches(0.35)
    r_b1 = p_body1.add_run(
        f"根据《中华人民共和国律师法》第二十八条及《中华人民共和国民事诉讼法》第六十一条之规定，"
        f"本所依法接受当事人【{client_name}】之委托，特指派本所专职律师【{lawyer_name}】"
        f"（执业证号：{lawyer_license}，联系电话：{lawyer_phone}）担任其与【{opponent_name}】"
        f"【{cause_of_action}】一案的诉讼代理人，出庭参加诉讼活动。"
    )
    r_b1.font.name = "仿宋_GB2312"
    r_b1.font.size = Pt(12)

    # 正文二段
    p_body2 = doc.add_paragraph()
    p_body2.paragraph_format.line_spacing = 1.35
    p_body2.paragraph_format.space_after = Pt(10)
    p_body2.paragraph_format.first_line_indent = Inches(0.35)
    r_b2 = p_body2.add_run(
        "代理人的具体代理权限详见委托人出具之《民事授权委托书》。"
        "特提请贵院依法准予上述指派律师出庭履行代理职责，并在立案审查、文书送达、证据质证及开庭审理各阶段予以协助为盼。"
    )
    r_b2.font.name = "仿宋_GB2312"
    r_b2.font.size = Pt(12)

    # 结束语
    p_close = doc.add_paragraph()
    p_close.paragraph_format.first_line_indent = Inches(0.35)
    r_close = p_close.add_run("特此函告。")
    r_close.font.name = "仿宋_GB2312"
    r_close.font.size = Pt(12)

    # 签章落款
    doc.add_paragraph()
    p_sign1 = doc.add_paragraph()
    p_sign1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_s1 = p_sign1.add_run(f"律师事务所（公章）：{law_firm}")
    r_s1.font.name = "仿宋_GB2312"
    r_s1.font.size = Pt(12)

    p_sign2 = doc.add_paragraph()
    p_sign2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_s2 = p_sign2.add_run(f"指派律师（签字）：{lawyer_name}")
    r_s2.font.name = "仿宋_GB2312"
    r_s2.font.size = Pt(12)

    p_sign3 = doc.add_paragraph()
    p_sign3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_s3 = p_sign3.add_run(f"出具日期：{sign_date}")
    r_s3.font.name = "仿宋_GB2312"
    r_s3.font.size = Pt(12)

    doc.save(output_path)
    from gateway.utils.disclaimer import add_disclaimer_to_docx
    add_disclaimer_to_docx(output_path)


def _create_id_credential_attachment_docx(
    output_path: str,
    plaintiff: Dict[str, Any],
    court_name: str,
    sign_date: str,
) -> None:
    """生成《原告主体资格证明材料说明页》（身份证/营业执照扫描粘贴说明页）。"""
    import docx as python_docx
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = python_docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.1)
        s.right_margin = Inches(1.1)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("原告主体资格证明材料说明页")
    r_title.font.name = "黑体"
    r_title.font.size = Pt(16)
    r_title.font.bold = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("（人民法院在线服务 · 01_起诉状及身份证明材料专页）")
    r_sub.font.name = "楷体"
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    is_legal = plaintiff.get("party_type") == "legal"
    party_type_label = "法人或其他组织" if is_legal else "自然人"
    code_label = "统一社会信用代码" if is_legal else "公民身份号码"

    # 信息表格
    table = doc.add_table(rows=4 if not is_legal else 5, cols=2)
    table.style = 'Table Grid'
    rows_data = [
        ("原告主体名称", str(plaintiff.get("name", PLACEHOLDER))),
        ("主体类型", party_type_label),
        (code_label, str(plaintiff.get("id_card_or_uscc", PLACEHOLDER))),
        ("送达住址 / 住所地", str(plaintiff.get("address", PLACEHOLDER))),
    ]
    if is_legal:
        rows_data.append((
            "法定代表人 / 负责人",
            f"{plaintiff.get('legal_rep', PLACEHOLDER)}（职务：{plaintiff.get('rep_position', '法定代表人')}）"
        ))
    for i, (k, v) in enumerate(rows_data):
        row = table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(1.8)
        c1.width = Inches(4.5)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.name = "黑体"
        r0.font.size = Pt(10.5)
        r0.font.bold = True
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.name = "仿宋_GB2312"
        r1.font.size = Pt(10.5)

    doc.add_paragraph()
    # 说明文字
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.line_spacing = 1.3
    r_desc = p_desc.add_run("【最高人民法院在线诉讼身份证明材料装订说明】\n")
    r_desc.font.name = "黑体"
    r_desc.font.size = Pt(11)
    r_desc.font.bold = True
    r_desc_b = p_desc.add_run(
        "1. 自然人原告：请将二代居民身份证人像面与国徽面（或港澳台居住证/护照）彩色复印件或清晰扫描件附于本说明页之后。\n"
        "2. 法人或其他组织原告：请在此页后附最新加盖公章的《营业执照》副本复印件，并一并提交加盖公章的《法定代表人身份证明书》及法定代表人居民身份证复印件。\n"
        "3. 本页按照人民法院在线服务电子立案规范编制，与起诉状共同作为第一类目（01_起诉状及身份证明）合卷上传。"
    )
    r_desc_b.font.name = "仿宋_GB2312"
    r_desc_b.font.size = Pt(10)

    # 粘贴/附页示意框
    t_box = doc.add_table(rows=1, cols=1)
    t_box.style = 'Table Grid'
    cell_box = t_box.rows[0].cells[0]
    cell_box.width = Inches(6.3)
    p_box = cell_box.paragraphs[0]
    p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_box.paragraph_format.space_before = Pt(36)
    p_box.paragraph_format.space_after = Pt(36)
    r_box = p_box.add_run("【此处附：身份证件或营业执照复印件 / 扫描件】")
    r_box.font.name = "楷体"
    r_box.font.size = Pt(11)
    r_box.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

    doc.add_paragraph()
    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_foot = p_foot.add_run(f"原告（签名/盖章）：____________________    日期：{sign_date}")
    r_foot.font.name = "仿宋_GB2312"
    r_foot.font.size = Pt(11)

    doc.save(output_path)
    from gateway.utils.disclaimer import add_disclaimer_to_docx
    add_disclaimer_to_docx(output_path)


def _create_lawyer_license_attachment_docx(
    output_path: str,
    lawyer_name: str,
    law_firm: str,
    lawyer_license: str,
    lawyer_phone: str,
    court_name: str,
    sign_date: str,
) -> None:
    """生成《律师执业证复印件说明页》（含年度考核说明与粘贴框）。"""
    import docx as python_docx
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = python_docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.1)
        s.right_margin = Inches(1.1)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("律师执业证复印件说明页")
    r_title.font.name = "黑体"
    r_title.font.size = Pt(16)
    r_title.font.bold = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("（人民法院在线服务 · 02_委托代理材料专页）")
    r_sub.font.name = "楷体"
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    table = doc.add_table(rows=4, cols=2)
    table.style = 'Table Grid'
    rows_data = [
        ("承办律师姓名", lawyer_name or PLACEHOLDER),
        ("执业机构", law_firm or PLACEHOLDER),
        ("律师执业证号", lawyer_license or PLACEHOLDER),
        ("联系电话", lawyer_phone or PLACEHOLDER),
    ]
    for i, (k, v) in enumerate(rows_data):
        row = table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(1.8)
        c1.width = Inches(4.5)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.name = "黑体"
        r0.font.size = Pt(10.5)
        r0.font.bold = True
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.name = "仿宋_GB2312"
        r1.font.size = Pt(10.5)

    doc.add_paragraph()
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.line_spacing = 1.3
    r_desc = p_desc.add_run("【最高人民法院在线诉讼委托代理资格说明】\n")
    r_desc.font.name = "黑体"
    r_desc.font.size = Pt(11)
    r_desc.font.bold = True
    r_desc_b = p_desc.add_run(
        "1. 按照最高人民法院《人民法院在线诉讼规则》第十一条规定，诉讼代理人应当提交律师执业证书及律师事务所公函。\n"
        "2. 请将承办律师有效执业证书基本信息页（含彩色照片）及最近年度考核称职备案盖章页复印件粘贴于下方区域。\n"
        "3. 本页与《民事授权委托书》、《律师事务所公函》共同归入第二类目（02_委托代理材料）合卷上传。"
    )
    r_desc_b.font.name = "仿宋_GB2312"
    r_desc_b.font.size = Pt(10)

    t_box = doc.add_table(rows=1, cols=1)
    t_box.style = 'Table Grid'
    cell_box = t_box.rows[0].cells[0]
    cell_box.width = Inches(6.3)
    p_box = cell_box.paragraphs[0]
    p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_box.paragraph_format.space_before = Pt(36)
    p_box.paragraph_format.space_after = Pt(36)
    r_box = p_box.add_run("【此处附：律师执业证基本信息页及年度考核备案页复印件】")
    r_box.font.name = "楷体"
    r_box.font.size = Pt(11)
    r_box.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

    doc.add_paragraph()
    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_foot = p_foot.add_run(f"承办律师（签字）：____________________    日期：{sign_date}")
    r_foot.font.name = "仿宋_GB2312"
    r_foot.font.size = Pt(11)

    doc.save(output_path)
    from gateway.utils.disclaimer import add_disclaimer_to_docx
    add_disclaimer_to_docx(output_path)


def _create_evidence_bundle_cover_docx(
    output_path: str,
    cause_of_action: str,
    plaintiff_name: str,
    opponent_name: str,
    court_name: str,
    evidences: Optional[List[Any]] = None,
    sign_date: str = "",
) -> None:
    """生成《初步证据材料合卷说明页》（附证据装订总则与目录汇总）。"""
    import docx as python_docx
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = python_docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.1)
        s.right_margin = Inches(1.1)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("初步证据材料合卷说明页")
    r_title.font.name = "黑体"
    r_title.font.size = Pt(16)
    r_title.font.bold = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("（人民法院在线服务 · 04_证据清单与证据材料专页）")
    r_sub.font.name = "楷体"
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.line_spacing = 1.3
    r_m = p_meta.add_run(
        f"案件案由：{cause_of_action}\n"
        f"原告（申请人）：{plaintiff_name}\n"
        f"被告（被申请人）：{opponent_name}\n"
        f"受诉机构：{court_name}\n"
        f"初步证据数量：共 {len(evidences or [])} 项\n"
    )
    r_m.font.name = "仿宋_GB2312"
    r_m.font.size = Pt(11)

    ev_list = evidences or []
    if ev_list:
        table = doc.add_table(rows=len(ev_list) + 1, cols=4)
        table.style = 'Table Grid'
        headers = ["序号", "证据名称", "证明目的", "形式"]
        widths = [Inches(0.8), Inches(2.2), Inches(2.5), Inches(0.8)]
        hdr_cells = table.rows[0].cells
        for idx, h_text in enumerate(headers):
            hdr_cells[idx].width = widths[idx]
            p = hdr_cells[idx].paragraphs[0]
            r = p.add_run(h_text)
            r.font.name = "黑体"
            r.font.size = Pt(10)
            r.font.bold = True
        for row_idx, ev in enumerate(ev_list, start=1):
            row_cells = table.rows[row_idx].cells
            for c_idx in range(4):
                row_cells[c_idx].width = widths[c_idx]
            ev_name = getattr(ev, "name", None) or (ev.get("name") if isinstance(ev, dict) else str(ev))
            ev_purpose = getattr(ev, "purpose", None) or (ev.get("purpose") if isinstance(ev, dict) else "")
            ev_format = getattr(ev, "original_or_copy", None) or (ev.get("original_or_copy") if isinstance(ev, dict) else "复印件")
            for c_idx, val in enumerate([str(row_idx), str(ev_name or ""), str(ev_purpose or ""), str(ev_format or "")]):
                p = row_cells[c_idx].paragraphs[0]
                r = p.add_run(val)
                r.font.name = "仿宋_GB2312"
                r.font.size = Pt(9.5)
    else:
        p_empty = doc.add_paragraph()
        r_empty = p_empty.add_run("（证据材料清单详见《证据目录》，本合卷页后附全部初步书证及电子证据打印件）")
        r_empty.font.name = "楷体"
        r_empty.font.size = Pt(10.5)

    doc.add_paragraph()
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.line_spacing = 1.3
    r_desc = p_desc.add_run("【最高人民法院在线诉讼证据装订与质证指引】\n")
    r_desc.font.name = "黑体"
    r_desc.font.size = Pt(11)
    r_desc.font.bold = True
    r_desc_b = p_desc.add_run(
        "1. 移动微法院证据上传时，证据文件序号应与本合卷说明页及《证据目录》严格对应。\n"
        "2. 多页书证应合并为单个 PDF 依次上传，单文件建议不超过 50MB。\n"
        "3. 开庭质证时原告应携带全部证据原件或原始存储介质以供核对。"
    )
    r_desc_b.font.name = "仿宋_GB2312"
    r_desc_b.font.size = Pt(10)

    doc.add_paragraph()
    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_foot = p_foot.add_run(f"提交人（签名/盖章）：____________________    日期：{sign_date}")
    r_foot.font.name = "仿宋_GB2312"
    r_foot.font.size = Pt(11)

    doc.save(output_path)
    from gateway.utils.disclaimer import add_disclaimer_to_docx
    add_disclaimer_to_docx(output_path)


def classify_litigation_file(filename: str) -> str:
    """按最高人民法院在线服务立案四大类目划分文件所属目录。"""
    if any(k in filename for k in ["起诉状", "仲裁申请书", "法定代表人", "原告身份证明", "财产保全"]):
        return "01_起诉状及身份证明"
    elif any(k in filename for k in ["委托书", "律师事务所公函", "律师执业证"]):
        return "02_委托代理材料"
    elif any(k in filename for k in ["送达地址"]):
        return "03_送达地址确认书"
    elif any(k in filename for k in ["证据", "补偿金"]):
        return "04_证据清单与证据材料"
    return "01_起诉状及身份证明"


def build_litigation_archive(
    package_id: str,
    generated_docs: List[Dict[str, Any]],
    safe_party_name: str,
    timestamp: int,
    archive_format: Literal["spc_mobile", "flat"] = "spc_mobile",
) -> Tuple[str, str]:
    """
    构建立案公文包 ZIP 归档并保存至 storage/generated/。
    - spc_mobile: 最高人民法院移动微法院四级类目归档
    - flat: 传统扁平单层 ZIP 归档
    返回: (zip_filename, zip_filepath)
    """
    generated_dir = get_generated_dir()
    if archive_format == "spc_mobile":
        zip_filename = f"立案公文包_移动微法院标准_{package_id}_{safe_party_name}_{timestamp}.zip"
    else:
        zip_filename = f"立案公文包_扁平打包_{package_id}_{safe_party_name}_{timestamp}.zip"
    zip_path = os.path.join(str(generated_dir), zip_filename)

    spc_folders = [
        "01_起诉状及身份证明",
        "02_委托代理材料",
        "03_送达地址确认书",
        "04_证据清单与证据材料",
    ]

    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        if archive_format == "spc_mobile":
            for folder in spc_folders:
                zf.writestr(f"{folder}/", "")
            for doc in generated_docs:
                cat = doc.get("category") or classify_litigation_file(doc.get("filename_docx", ""))
                # docx
                if doc.get("filename_docx"):
                    docx_path = doc.get("path") or os.path.join(str(generated_dir), doc["filename_docx"])
                    if os.path.isfile(docx_path):
                        zf.write(docx_path, arcname=f"{cat}/{doc['filename_docx']}")
                # pdf
                if doc.get("filename_pdf"):
                    pdf_path = doc.get("pdf_path") or os.path.join(str(generated_dir), doc["filename_pdf"])
                    if os.path.isfile(pdf_path):
                        zf.write(pdf_path, arcname=f"{cat}/{doc['filename_pdf']}")
        else:
            for doc in generated_docs:
                if doc.get("filename_docx"):
                    docx_path = doc.get("path") or os.path.join(str(generated_dir), doc["filename_docx"])
                    if os.path.isfile(docx_path):
                        zf.write(docx_path, arcname=doc["filename_docx"])
                if doc.get("filename_pdf"):
                    pdf_path = doc.get("pdf_path") or os.path.join(str(generated_dir), doc["filename_pdf"])
                    if os.path.isfile(pdf_path):
                        zf.write(pdf_path, arcname=doc["filename_pdf"])

    return zip_filename, zip_path


# ==============================================================================
# 核心大模型案情法言法语润色与诉讼文书包一键流水线
# ==============================================================================
@router.post("/api/v1/phase1/generate-litigation-package", summary="三大高频案由全套立案文书流水线 (.docx + .pdf)")
async def generate_litigation_package(req: GenerateLitigationPackageReq):
    """
    一键输出：
    1. 宿主机 gemma4:12b 润色的引用民法典或劳动合同法的《事实与理由》
    2. 主起诉状/仲裁申请书 (.docx + .pdf)
    3. 配套法律公文包（授权委托书、法定代表人身份证明、送达地址确认书、保全申请书或补偿金清单）
    """
    timestamp = int(time.time())
    package_id = f"pkg-{uuid.uuid4().hex[:10]}"
    filing_date_str = datetime.now().strftime("%Y年%m月%d日")
    missing_fields: List[str] = []

    # 1. 当事人解析（自然人/法人）
    plaintiff = _resolve_party(req.plaintiff, "plaintiff", missing_fields)
    defendant = _resolve_party(req.defendant, "defendant", missing_fields)

    court_full = normalize_court_name(req.court_name)
    is_labor = "劳动" in req.cause_of_action
    arbitration_full = ""
    if is_labor:
        region_hint = (req.court_name or "").replace("人民法院", "").replace("法院", "").strip()
        arbitration_full = normalize_arbitration_committee(
            req.arbitration_committee or req.court_name, region_hint=region_hint
        )

    # 退休审查（仅自然人身份证）
    p_id_info: Dict[str, Any] = {"is_retired": False}
    if plaintiff["party_type"] == "natural" and len(plaintiff["id_card_or_uscc"]) == 18:
        p_id_info = validate_and_parse_idcard(plaintiff["id_card_or_uscc"])

    # 2. 事实与理由：显式传入案由，禁止兜底猜案由
    prompt = f"""【案件案由】{req.cause_of_action}
【诉讼请求】
{chr(10).join(req.claims or [])}

【律师案情碎片笔记】
{req.raw_case_notes}

【原告】{plaintiff['name']} (住址: {plaintiff['address']})
【被告】{defendant['name']} (住址: {defendant['address']})
【受诉法院】{court_full if not is_labor else arbitration_full}

请将上述材料改写为条理清晰、因果严密、引用具体法律条文（如《民法典》或《劳动合同法》具体法条）的《事实与理由》。"""

    from gateway.utils.pii_sanitizer import build_entities_from_parties

    custom_entities = build_entities_from_parties([
        {
            "name": plaintiff.get("name"),
            "address": plaintiff.get("address"),
            "legal_rep": plaintiff.get("legal_rep"),
            "party_type": plaintiff.get("party_type"),
        },
        {
            "name": defendant.get("name"),
            "address": defendant.get("address"),
            "legal_rep": defendant.get("legal_rep"),
            "party_type": defendant.get("party_type"),
        },
    ])

    try:
        facts_polished = await ollama_client.generate(
            prompt=prompt,
            system_prompt=FACTS_POLISH_SYSTEM_PROMPT,
            temperature=0.2,
            cause_of_action=req.cause_of_action,
            custom_entities=custom_entities,
        )
    except Exception as e:
        logger.warning(f"大模型润色异常，降级使用原始笔记: {e}")
        facts_polished = req.raw_case_notes

    from gateway.utils.citation_guard import audit_citations_in_text
    citation_audit = audit_citations_in_text(facts_polished, auto_sanitize=True)
    if citation_audit.get("has_abrogated_laws") and citation_audit.get("sanitized_text"):
        facts_polished = citation_audit["sanitized_text"]

    statute_pattern = re.compile(r"《(?:中华人民共和国)?(?:民法典|劳动合同法|民事诉讼法)[^》]+》[第\d条款、及]+")
    statute_citations = statute_pattern.findall(facts_polished)

    court_fees = calculate_court_fees(req.claim_amount)

    # 利息：约定利率 + private_lending_cap 上限，不再默认 4 倍 LPR
    interest_statement = ""
    start_date_str = extract_first_date_iso(req.raw_case_notes) or req.contract_date or "2023-01-01"
    start_date_display = format_date_cn(start_date_str)
    total_interest_val = 0.0
    agreed_rate = req.agreed_annual_rate
    if agreed_rate is None:
        agreed_rate = extract_agreed_annual_rate(req.raw_case_notes)
    interest_meta: Dict[str, Any] = {"applied_rate": agreed_rate, "capped": False}

    if "借贷" in req.cause_of_action:
        contract_date = req.contract_date or start_date_str
        end_date = datetime.now().strftime("%Y-%m-%d")
        try:
            lpr_calc = calculate_lpr_interest(
                principal=req.claim_amount,
                start_date=start_date_str,
                end_date=end_date,
                fixed_annual_rate=agreed_rate if agreed_rate is not None else None,
                rule_mode="private_lending_cap",
                contract_date=contract_date,
            )
            interest_statement = lpr_calc.get("complaint_statement", "")
            total_interest_val = float(lpr_calc.get("total_interest", 0.0))
            # 从分段取实际执行年利率
            segs = lpr_calc.get("segments") or []
            applied = None
            capped = False
            for seg in segs:
                if seg.get("applied_annual_rate") is not None:
                    applied = float(seg["applied_annual_rate"])
                if seg.get("is_rate_capped"):
                    capped = True
            if applied is None and agreed_rate is not None:
                applied = float(agreed_rate)
            interest_meta = {"applied_rate": applied, "capped": capped, "agreed_rate": agreed_rate}
        except Exception as e:
            logger.warning(f"利息精算失败: {e}")
            interest_meta = {"applied_rate": agreed_rate, "capped": False, "agreed_rate": agreed_rate}

    claims_final = _build_claims_for_cause(req, agreed_rate, start_date_str, start_date_display, interest_meta)
    if not claims_final:
        claims_final = [PLACEHOLDER]
        missing_fields.append("claims")

    evidences = []
    if req.evidences:
        for ev in req.evidences:
            evidences.append(ev.model_dump() if hasattr(ev, "model_dump") else dict(ev))
    copies = req.copies if req.copies is not None else 1

    # 当事人公共模板上下文
    def _party_ctx(prefix: str, info: Dict[str, Any]) -> Dict[str, Any]:
        return {
            f"{prefix}_name": info["name"],
            f"{prefix}_gender": info.get("gender") or "",
            f"{prefix}_birth": info.get("birth") or "",
            f"{prefix}_ethnicity": info.get("ethnicity") or "",
            f"{prefix}_address": info["address"],
            f"{prefix}_id": info["id_card_or_uscc"],
            f"{prefix}_id_label": info.get("id_label") or "",
            f"{prefix}_phone": info.get("phone") or PLACEHOLDER,
            f"{prefix}_legal_rep": info.get("legal_rep") or "",
            f"{prefix}_rep_position": info.get("rep_position") or "",
            f"{prefix}_party_type": info["party_type"],
            f"{prefix}_is_natural": info["party_type"] == "natural",
            f"{prefix}_is_legal": info["party_type"] == "legal",
        }

    generated_docs = []
    common_party = {}
    common_party.update(_party_ctx("plaintiff", plaintiff))
    common_party.update(_party_ctx("defendant", defendant))
    # 劳动仲裁用语别名
    common_party.update(_party_ctx("applicant", plaintiff))
    common_party.update(_party_ctx("respondent", defendant))

    if "借贷" in req.cause_of_action:
        main_tpl = "civil_complaint_loan_template.docx"
        doc_label = "民事起诉状(民间借贷版)"
        context_main = {
            **common_party,
            "claims": claims_final,
            "loan_principal": f"{req.claim_amount:,.2f}",
            "loan_principal_capital": rmb_to_capital(req.claim_amount),
            "interest_start_date": start_date_display,
            "interest_annual_rate": (
                f"{interest_meta['applied_rate']:g}" if interest_meta.get("applied_rate") is not None else PLACEHOLDER
            ),
            "interest_base": f"{req.claim_amount:,.2f}",
            "interest_end_date": filing_date_str,
            "total_interest": f"{total_interest_val:,.2f}" if total_interest_val else PLACEHOLDER,
            "total_interest_capital": rmb_to_capital(total_interest_val) if total_interest_val else PLACEHOLDER,
            "lawyer_fee_note": "需有合同约定或法律依据",
            "facts_and_reasons": facts_polished,
            "interest_calc_statement": interest_statement,
            "court_name": court_full,
            "filing_date": filing_date_str,
            "evidences": evidences,
            "copies": copies,
        }
    elif is_labor:
        main_tpl = "labor_arbitration_template.docx"
        doc_label = "劳动人事争议仲裁申请书"
        context_main = {
            **common_party,
            "respondent_legal_rep": defendant.get("legal_rep") or PLACEHOLDER,
            "arbitration_claims": claims_final,
            "facts_and_reasons": facts_polished,
            "arbitration_committee": arbitration_full,
            "filing_date": filing_date_str,
            "evidences": evidences,
            "arbitration_limitation_note": "仲裁时效一般为1年，具体以律师核实为准",
        }
    elif "买卖" in req.cause_of_action:
        main_tpl = "civil_complaint_sales_template.docx"
        doc_label = "民事起诉状(买卖合同版)"
        context_main = {
            **common_party,
            "claims": claims_final,
            "contract_arrears": f"{req.claim_amount:,.2f}",
            "contract_arrears_capital": rmb_to_capital(req.claim_amount),
            "liquidated_damages": PLACEHOLDER,
            "overdue_rate_desc": "按照全国银行间同业拆借中心公布的一年期贷款市场报价利率（LPR）及合同约定或依法可主张的标准计算",
            "facts_and_reasons": facts_polished,
            "court_name": court_full,
            "filing_date": filing_date_str,
            "evidences": evidences,
            "copies": copies,
        }
    else:
        main_tpl = "civil_complaint_template.docx"
        doc_label = "民事起诉状"
        context_main = {
            **common_party,
            "claims": claims_final,
            "facts_and_reasons": facts_polished,
            "court_name": court_full,
            "filing_date": filing_date_str,
            "evidences": evidences,
            "copies": copies,
        }

    tpl_path = os.path.join(_templates_dir(), main_tpl)
    safe_party_name = sanitize_filename(plaintiff["name"]) or "当事人"
    main_docx_name = f"{doc_label}_{safe_party_name}_{timestamp}.docx"
    main_docx_path = os.path.join(str(get_generated_dir()), main_docx_name)
    main_pdf_name = f"{doc_label}_{safe_party_name}_{timestamp}.pdf"
    main_pdf_path = os.path.join(str(get_generated_dir()), main_pdf_name)

    doc = DocxTemplate(tpl_path)
    doc.render(context_main)
    doc.save(main_docx_path)
    _highlight_placeholders_in_docx(main_docx_path)
    from gateway.utils.disclaimer import add_disclaimer_to_docx
    add_disclaimer_to_docx(main_docx_path)

    pdf_ok = await convert_docx_to_pdf(main_docx_path, main_pdf_path, title=doc_label, text_summary=facts_polished)
    if not pdf_ok:
        logger.error(f"主文书 PDF 转换失败: {main_docx_path}")

    generated_docs.append({
        "doc_type": doc_label,
        "category": "01_起诉状及身份证明",
        "filename_docx": main_docx_name,
        "path": main_docx_path,
        "docx_url": f"/api/documents/download/{main_docx_name}",
        "filename_pdf": main_pdf_name if pdf_ok else None,
        "pdf_path": main_pdf_path if pdf_ok else None,
        "pdf_url": f"/api/documents/download/{main_pdf_name}" if pdf_ok else None,
    })

    # 附属文书列表判定（A-18：保全申请书仅在显式勾选 property_preservation 时生成）
    if req.sub_documents is not None:
        target_sub_docs = set(req.sub_documents)
    elif req.auto_generate_sub_docs:
        target_sub_docs = {"power_of_attorney", "service_address"}
        if plaintiff.get("party_type") == "legal":
            target_sub_docs.add("legal_rep_certificate")
    else:
        target_sub_docs = set()

    if "power_of_attorney" in target_sub_docs:
        poa_tpl = os.path.join(_templates_dir(), "power_of_attorney_template.docx")
        poa_docx_name = f"民事授权委托书_{safe_party_name}_{timestamp}.docx"
        poa_docx_path = os.path.join(str(get_generated_dir()), poa_docx_name)
        poa_pdf_name = f"民事授权委托书_{safe_party_name}_{timestamp}.pdf"
        poa_pdf_path = os.path.join(str(get_generated_dir()), poa_pdf_name)

        doc_poa = DocxTemplate(poa_tpl)
        doc_poa.render({
            "client_name": plaintiff["name"],
            "client_id": plaintiff["id_card_or_uscc"],
            "client_phone": plaintiff.get("phone") or PLACEHOLDER,
            "opponent_name": defendant["name"],
            "cause_of_action": req.cause_of_action,
            "lawyer_name": req.lawyer_name,
            "law_firm": req.law_firm,
            "lawyer_license": req.lawyer_license,
            "lawyer_phone": req.lawyer_phone,
            "court_name": court_full,
            "authorization_scope": "□一般代理  □特别授权（□代为承认、放弃、变更诉讼请求 □进行和解 □提起反诉或上诉）",
            "sign_date": filing_date_str,
        })
        doc_poa.save(poa_docx_path)
        poa_pdf_ok = await convert_docx_to_pdf(poa_docx_path, poa_pdf_path, title="民事授权委托书")
        generated_docs.append({
            "doc_type": "民事授权委托书",
            "category": "02_委托代理材料",
            "filename_docx": poa_docx_name,
            "path": poa_docx_path,
            "docx_url": f"/api/documents/download/{poa_docx_name}",
            "filename_pdf": poa_pdf_name if poa_pdf_ok else None,
            "pdf_path": poa_pdf_path if poa_pdf_ok else None,
            "pdf_url": f"/api/documents/download/{poa_pdf_name}" if poa_pdf_ok else None,
        })

    if "service_address" in target_sub_docs:
        addr_tpl = os.path.join(_templates_dir(), "service_address_template.docx")
        addr_docx_name = f"送达地址确认书_{safe_party_name}_{timestamp}.docx"
        addr_docx_path = os.path.join(str(get_generated_dir()), addr_docx_name)
        addr_pdf_name = f"送达地址确认书_{safe_party_name}_{timestamp}.pdf"
        addr_pdf_path = os.path.join(str(get_generated_dir()), addr_pdf_name)
        doc_addr = DocxTemplate(addr_tpl)
        doc_addr.render({
            "cause_of_action": req.cause_of_action,
            "recipient_name": plaintiff["name"],
            "recipient_id": plaintiff["id_card_or_uscc"],
            "lawyer_name": req.lawyer_name,
            "law_firm": req.law_firm,
            "lawyer_phone": req.lawyer_phone,
            "mail_address": plaintiff["address"],
            "postal_code": PLACEHOLDER,
            "mobile_phone": plaintiff.get("phone") or PLACEHOLDER,
            "email_address": PLACEHOLDER,
            "electronic_methods": "短信、电子邮件、统一微法院",
            "sign_date": filing_date_str,
        })
        doc_addr.save(addr_docx_path)
        addr_pdf_ok = await convert_docx_to_pdf(addr_docx_path, addr_pdf_path, title="送达地址确认书")
        generated_docs.append({
            "doc_type": "送达地址确认书",
            "category": "03_送达地址确认书",
            "filename_docx": addr_docx_name,
            "path": addr_docx_path,
            "docx_url": f"/api/documents/download/{addr_docx_name}",
            "filename_pdf": addr_pdf_name if addr_pdf_ok else None,
            "pdf_path": addr_pdf_path if addr_pdf_ok else None,
            "pdf_url": f"/api/documents/download/{addr_pdf_name}" if addr_pdf_ok else None,
        })

    if "legal_rep_certificate" in target_sub_docs:
        rep_tpl = os.path.join(_templates_dir(), "legal_representative_template.docx")
        if os.path.isfile(rep_tpl):
            rep_docx_name = f"法定代表人身份证明书_{safe_party_name}_{timestamp}.docx"
            rep_docx_path = os.path.join(str(get_generated_dir()), rep_docx_name)
            rep_pdf_name = f"法定代表人身份证明书_{safe_party_name}_{timestamp}.pdf"
            rep_pdf_path = os.path.join(str(get_generated_dir()), rep_pdf_name)
            doc_rep = DocxTemplate(rep_tpl)
            doc_rep.render({
                "legal_rep_name": plaintiff.get("legal_rep") or PLACEHOLDER,
                "legal_rep_gender": plaintiff.get("gender") or PLACEHOLDER,
                "legal_rep_id": PLACEHOLDER,
                "legal_rep_position": plaintiff.get("rep_position") or "法定代表人",
                "company_name": plaintiff["name"],
                "company_uscc": plaintiff["id_card_or_uscc"] or PLACEHOLDER,
                "opponent_name": defendant["name"],
                "cause_of_action": req.cause_of_action,
                "sign_date": filing_date_str,
            })
            doc_rep.save(rep_docx_path)
            rep_pdf_ok = await convert_docx_to_pdf(rep_docx_path, rep_pdf_path, title="法定代表人身份证明书")
            generated_docs.append({
                "doc_type": "法定代表人身份证明书",
                "category": "01_起诉状及身份证明",
                "filename_docx": rep_docx_name,
                "path": rep_docx_path,
                "docx_url": f"/api/documents/download/{rep_docx_name}",
                "filename_pdf": rep_pdf_name if rep_pdf_ok else None,
                "pdf_path": rep_pdf_path if rep_pdf_ok else None,
                "pdf_url": f"/api/documents/download/{rep_pdf_name}" if rep_pdf_ok else None,
            })

    if "property_preservation" in target_sub_docs:
        preserv_tpl = os.path.join(_templates_dir(), "property_preservation_template.docx")
        preserv_docx_name = f"财产保全申请书_{safe_party_name}_{timestamp}.docx"
        preserv_docx_path = os.path.join(str(get_generated_dir()), preserv_docx_name)
        preserv_pdf_name = f"财产保全申请书_{safe_party_name}_{timestamp}.pdf"
        preserv_pdf_path = os.path.join(str(get_generated_dir()), preserv_pdf_name)
        doc_preserv = DocxTemplate(preserv_tpl)
        doc_preserv.render({
            "applicant_name": plaintiff["name"],
            "applicant_id": plaintiff["id_card_or_uscc"],
            "applicant_phone": plaintiff.get("phone") or PLACEHOLDER,
            "applicant_address": plaintiff["address"],
            "applicant_legal_rep": plaintiff.get("legal_rep") or PLACEHOLDER,
            "respondent_name": defendant["name"],
            "respondent_id": defendant["id_card_or_uscc"],
            "respondent_address": defendant["address"],
            "preservation_amount": f"{req.claim_amount:,.2f}",
            "preservation_amount_capital": rmb_to_capital(req.claim_amount),
            "cause_of_action": req.cause_of_action,
            "property_clues": f"1. 被申请人名下银行基本存款账户；\n2. 被申请人名下位于 {defendant['address']} 之不动产或经营设备。",
            "guarantee_method": PLACEHOLDER,
            "guarantee_institution": PLACEHOLDER,
            "preservation_type_pre": "□",
            "preservation_type_in": "□",
            "case_number": PLACEHOLDER,
            "court_name": court_full,
            "sign_date": filing_date_str,
        })
        doc_preserv.save(preserv_docx_path)
        preserv_pdf_ok = await convert_docx_to_pdf(preserv_docx_path, preserv_pdf_path, title="财产保全申请书")
        generated_docs.append({
            "doc_type": "财产保全申请书",
            "category": "01_起诉状及身份证明",
            "filename_docx": preserv_docx_name,
            "path": preserv_docx_path,
            "docx_url": f"/api/documents/download/{preserv_docx_name}",
            "filename_pdf": preserv_pdf_name if preserv_pdf_ok else None,
            "pdf_path": preserv_pdf_path if preserv_pdf_ok else None,
            "pdf_url": f"/api/documents/download/{preserv_pdf_name}" if preserv_pdf_ok else None,
        })

    if "evidence_catalog" in target_sub_docs:
        ev_tpl = os.path.join(_templates_dir(), "evidence_catalog_template.docx")
        if os.path.isfile(ev_tpl):
            ev_docx_name = f"证据目录_{safe_party_name}_{timestamp}.docx"
            ev_docx_path = os.path.join(str(get_generated_dir()), ev_docx_name)
            ev_pdf_name = f"证据目录_{safe_party_name}_{timestamp}.pdf"
            ev_pdf_path = os.path.join(str(get_generated_dir()), ev_pdf_name)
            doc_ev = DocxTemplate(ev_tpl)
            doc_ev.render({
                "cause_of_action": req.cause_of_action,
                "plaintiff_name": plaintiff["name"],
                "filing_date": filing_date_str,
                "evidences": req.evidences or [],
            })
            doc_ev.save(ev_docx_path)
            ev_pdf_ok = await convert_docx_to_pdf(ev_docx_path, ev_pdf_path, title="证据目录")
            generated_docs.append({
                "doc_type": "证据目录",
                "category": "04_证据清单与证据材料",
                "filename_docx": ev_docx_name,
                "path": ev_docx_path,
                "docx_url": f"/api/documents/download/{ev_docx_name}",
                "filename_pdf": ev_pdf_name if ev_pdf_ok else None,
                "pdf_path": ev_pdf_path if ev_pdf_ok else None,
                "pdf_url": f"/api/documents/download/{ev_pdf_name}" if ev_pdf_ok else None,
            })

    if is_labor and (req.auto_generate_sub_docs or "labor_compensation_calc" in target_sub_docs):
        labor_tpl = os.path.join(_templates_dir(), "labor_compensation_calc_template.docx")
        labor_docx_name = f"经济补偿金计算清单_{safe_party_name}_{timestamp}.docx"
        labor_docx_path = os.path.join(str(get_generated_dir()), labor_docx_name)
        labor_calc = calculate_labor_compensation(12000, "2021-03-01", datetime.now().strftime("%Y-%m-%d"))
        doc_labor = DocxTemplate(labor_tpl)
        doc_labor.render({
            "applicant_name": plaintiff["name"],
            "respondent_name": defendant["name"],
            "hire_date": PLACEHOLDER,
            "termination_date": filing_date_str,
            "tenure_str": labor_calc["tenure_str"],
            "monthly_salary": PLACEHOLDER,
            "compensation_months": f"{labor_calc['compensation_months']}",
            "amount_n": f"{labor_calc['tiers']['N']['amount']:,.2f}",
            "amount_plus_1_notice": f"{labor_calc['tiers']['N+1']['notice_pay']:,.2f}",
            "amount_2n": f"{labor_calc['tiers']['2N']['amount']:,.2f}",
            "cap_reason": labor_calc["cap_reason"],
            "filing_date": filing_date_str,
        })
        doc_labor.save(labor_docx_path)
        generated_docs.append({
            "doc_type": "经济补偿金计算清单",
            "category": "04_证据清单与证据材料",
            "filename_docx": labor_docx_name,
            "path": labor_docx_path,
            "docx_url": f"/api/documents/download/{labor_docx_name}",
            "filename_pdf": None,
            "pdf_path": None,
            "pdf_url": None,
        })

    # 移动微法院在线诉讼配套支撑材料（当 auto_generate_sub_docs 或显式选择相应材料时生成）
    # 1. 01_起诉状及身份证明：原告主体资格证明材料说明页
    if req.auto_generate_sub_docs or "id_credential_attachment" in target_sub_docs:
        id_cred_docx_name = f"原告身份证明材料说明页_{safe_party_name}_{timestamp}.docx"
        id_cred_docx_path = os.path.join(str(get_generated_dir()), id_cred_docx_name)
        _create_id_credential_attachment_docx(
            output_path=id_cred_docx_path,
            plaintiff=plaintiff,
            court_name=court_full if not is_labor else arbitration_full,
            sign_date=filing_date_str,
        )
        generated_docs.append({
            "doc_type": "原告主体资格证明材料说明页",
            "category": "01_起诉状及身份证明",
            "filename_docx": id_cred_docx_name,
            "path": id_cred_docx_path,
            "docx_url": f"/api/documents/download/{id_cred_docx_name}",
            "filename_pdf": None,
            "pdf_path": None,
            "pdf_url": None,
        })

    # 2. 02_委托代理材料：律师事务所公函
    if req.auto_generate_sub_docs or "law_firm_letter" in target_sub_docs or ("power_of_attorney" in target_sub_docs and req.law_firm):
        firm_letter_docx_name = f"律师事务所公函_{safe_party_name}_{timestamp}.docx"
        firm_letter_docx_path = os.path.join(str(get_generated_dir()), firm_letter_docx_name)
        firm_letter_pdf_name = f"律师事务所公函_{safe_party_name}_{timestamp}.pdf"
        firm_letter_pdf_path = os.path.join(str(get_generated_dir()), firm_letter_pdf_name)
        _create_law_firm_letter_docx(
            output_path=firm_letter_docx_path,
            law_firm=req.law_firm,
            lawyer_name=req.lawyer_name,
            lawyer_license=req.lawyer_license,
            lawyer_phone=req.lawyer_phone,
            client_name=plaintiff["name"],
            opponent_name=defendant["name"],
            cause_of_action=req.cause_of_action,
            court_name=court_full if not is_labor else arbitration_full,
            sign_date=filing_date_str,
        )
        letter_pdf_ok = await convert_docx_to_pdf(firm_letter_docx_path, firm_letter_pdf_path, title="律师事务所公函")
        generated_docs.append({
            "doc_type": "律师事务所公函",
            "category": "02_委托代理材料",
            "filename_docx": firm_letter_docx_name,
            "path": firm_letter_docx_path,
            "docx_url": f"/api/documents/download/{firm_letter_docx_name}",
            "filename_pdf": firm_letter_pdf_name if letter_pdf_ok else None,
            "pdf_path": firm_letter_pdf_path if letter_pdf_ok else None,
            "pdf_url": f"/api/documents/download/{firm_letter_pdf_name}" if letter_pdf_ok else None,
        })

    # 3. 02_委托代理材料：律师执业证复印件说明页
    if req.auto_generate_sub_docs or "lawyer_license_attachment" in target_sub_docs or ("power_of_attorney" in target_sub_docs and req.lawyer_license):
        lawyer_lic_docx_name = f"律师执业证复印件说明页_{safe_party_name}_{timestamp}.docx"
        lawyer_lic_docx_path = os.path.join(str(get_generated_dir()), lawyer_lic_docx_name)
        _create_lawyer_license_attachment_docx(
            output_path=lawyer_lic_docx_path,
            lawyer_name=req.lawyer_name,
            law_firm=req.law_firm,
            lawyer_license=req.lawyer_license,
            lawyer_phone=req.lawyer_phone,
            court_name=court_full if not is_labor else arbitration_full,
            sign_date=filing_date_str,
        )
        generated_docs.append({
            "doc_type": "律师执业证复印件说明页",
            "category": "02_委托代理材料",
            "filename_docx": lawyer_lic_docx_name,
            "path": lawyer_lic_docx_path,
            "docx_url": f"/api/documents/download/{lawyer_lic_docx_name}",
            "filename_pdf": None,
            "pdf_path": None,
            "pdf_url": None,
        })

    # 4. 04_证据清单与证据材料：初步证据材料合卷说明页
    if req.auto_generate_sub_docs or "evidence_bundle_cover" in target_sub_docs:
        ev_cover_docx_name = f"初步证据材料合卷说明页_{safe_party_name}_{timestamp}.docx"
        ev_cover_docx_path = os.path.join(str(get_generated_dir()), ev_cover_docx_name)
        _create_evidence_bundle_cover_docx(
            output_path=ev_cover_docx_path,
            cause_of_action=req.cause_of_action,
            plaintiff_name=plaintiff["name"],
            opponent_name=defendant["name"],
            court_name=court_full if not is_labor else arbitration_full,
            evidences=req.evidences,
            sign_date=filing_date_str,
        )
        generated_docs.append({
            "doc_type": "初步证据材料合卷说明页",
            "category": "04_证据清单与证据材料",
            "filename_docx": ev_cover_docx_name,
            "path": ev_cover_docx_path,
            "docx_url": f"/api/documents/download/{ev_cover_docx_name}",
            "filename_pdf": None,
            "pdf_path": None,
            "pdf_url": None,
        })

    # 构建移动微法院标准 ZIP 包与传统扁平 ZIP 包
    spc_zip_name, spc_zip_path = build_litigation_archive(
        package_id=package_id,
        generated_docs=generated_docs,
        safe_party_name=safe_party_name,
        timestamp=timestamp,
        archive_format="spc_mobile",
    )
    flat_zip_name, flat_zip_path = build_litigation_archive(
        package_id=package_id,
        generated_docs=generated_docs,
        safe_party_name=safe_party_name,
        timestamp=timestamp,
        archive_format="flat",
    )

    _PACKAGE_ARCHIVE_REGISTRY[package_id] = {
        "spc_mobile": (spc_zip_name, spc_zip_path),
        "flat": (flat_zip_name, flat_zip_path),
    }

    archive_categories: Dict[str, List[str]] = {
        "01_起诉状及身份证明": [],
        "02_委托代理材料": [],
        "03_送达地址确认书": [],
        "04_证据清单与证据材料": [],
    }
    for doc in generated_docs:
        cat = doc.get("category") or classify_litigation_file(doc.get("filename_docx", ""))
        if cat not in archive_categories:
            archive_categories[cat] = []
        if doc.get("filename_docx") and doc["filename_docx"] not in archive_categories[cat]:
            archive_categories[cat].append(doc["filename_docx"])
        if doc.get("filename_pdf") and doc["filename_pdf"] not in archive_categories[cat]:
            archive_categories[cat].append(doc["filename_pdf"])

    llm_mode = getattr(ollama_client, "last_llm_mode", "rule_fallback")
    pdf_engine = get_last_pdf_engine() if pdf_ok else "local_pymupdf"
    engine_status = {
        "pdf": pdf_engine,
        "llm": llm_mode if llm_mode in ("ollama", "cloud", "rule_fallback") else "rule_fallback",
    }
    try:
        from gateway.utils.matter_vault_db import record_audit_log
        record_audit_log(
            action="document_generate",
            matter_id=getattr(req, "matter_id", None),
            target=main_docx_name,
            result="success"
        )
    except Exception as audit_err:
        logger.warning(f"记录文书生成审计日志失败: {audit_err}")

    return {
        "status": "success",
        "package_id": package_id,
        "cause_of_action": req.cause_of_action,
        "facts_and_reasons_polished": facts_polished,
        "statute_citations": list(set(statute_citations)),
        "citation_audit": citation_audit,
        "court_fees_estimate": court_fees,
        "retirement_check": p_id_info.get("retirement_note"),
        "missing_fields": missing_fields,
        "generated_documents": generated_docs,
        "main_docx_path": main_docx_path,
        "main_pdf_path": main_pdf_path if pdf_ok else None,
        "package_zip_name": spc_zip_name,
        "package_zip_url": f"/api/documents/download/{spc_zip_name}",
        "package_flat_zip_name": flat_zip_name,
        "package_flat_zip_url": f"/api/documents/download/{flat_zip_name}",
        "package_download_url": f"/api/v1/phase1/package/{package_id}/download?format={req.archive_format}",
        "archive_categories": archive_categories,
        "llm_mode": llm_mode,
        "engine_status": engine_status,
    }


# ==============================================================================
# 兼容旧版路由与下载接口
# ==============================================================================
@router.get("/api/documents/download/{filename}", summary="文书下载接口")
async def download_doc(filename: str, request: Request = None):
    safe_name = validate_download_filename(filename)
    try:
        file_path = safe_join(str(get_generated_dir()), safe_name)
    except HTTPException:
        raise
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="文件不存在")
    try:
        from gateway.utils.matter_vault_db import record_audit_log
        client_ip = request.client.host if request and request.client else None
        record_audit_log(
            action="document_download",
            target=safe_name,
            client_ip=client_ip,
            result="success"
        )
    except Exception as audit_err:
        logger.warning(f"记录文书下载审计日志失败: {audit_err}")
    if safe_name.endswith(".pdf"):
        media = "application/pdf"
    elif safe_name.endswith(".zip"):
        media = "application/zip"
    else:
        media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return FileResponse(path=file_path, filename=os.path.basename(safe_name), media_type=media)


@router.get("/api/v1/phase1/package/{package_id}/download", summary="移动微法院立案公文包 ZIP 归档直接下载")
async def download_litigation_package_zip(
    package_id: str,
    format: Literal["spc_mobile", "flat"] = Query("spc_mobile", description="归档格式：spc_mobile（四级分类）或 flat（单层扁平）"),
    request: Request = None,
):
    """
    流式下载指定立案文书包的 ZIP 归档。
    - format=spc_mobile: 最高人民法院移动微法院四级类目归档
    - format=flat: 传统扁平单层 ZIP 归档
    """
    entry = _PACKAGE_ARCHIVE_REGISTRY.get(package_id)
    zip_name, zip_path = None, None
    if entry and format in entry:
        zip_name, zip_path = entry[format]

    if not zip_path or not os.path.isfile(zip_path):
        gen_dir = str(get_generated_dir())
        prefix = "立案公文包_移动微法院标准_" if format == "spc_mobile" else "立案公文包_扁平打包_"
        if os.path.isdir(gen_dir):
            cand_files = [f for f in os.listdir(gen_dir) if f.startswith(prefix) and package_id in f and f.endswith(".zip")]
            if cand_files:
                cand_files.sort(key=lambda x: os.path.getmtime(os.path.join(gen_dir, x)), reverse=True)
                zip_name = cand_files[0]
                zip_path = os.path.join(gen_dir, zip_name)
                if package_id not in _PACKAGE_ARCHIVE_REGISTRY:
                    _PACKAGE_ARCHIVE_REGISTRY[package_id] = {}
                _PACKAGE_ARCHIVE_REGISTRY[package_id][format] = (zip_name, zip_path)

    if not zip_path or not os.path.isfile(zip_path):
        raise HTTPException(status_code=404, detail=f"未找到案卷包 {package_id} 对应的 ZIP 归档文件")

    try:
        from gateway.utils.matter_vault_db import record_audit_log
        client_ip = request.client.host if request and request.client else None
        record_audit_log(
            action="package_zip_download",
            target=zip_name,
            client_ip=client_ip,
            result="success"
        )
    except Exception as audit_err:
        logger.warning(f"记录文书包下载审计日志失败: {audit_err}")

    return FileResponse(
        path=zip_path,
        filename=os.path.basename(zip_name),
        media_type="application/zip",
    )


# ==============================================================================
# P3: 二审/再审关键阶段·一审裁决漏洞剖析与最高院类案分歧穿透接口
# ==============================================================================
class AppealAuditRequest(BaseModel):
    judgment_text: Optional[str] = Field(default="", description="一审判决书全文或裁判理由段落")
    case_title: Optional[str] = Field(default=None, description="案件名称")
    claim_amount: Optional[float] = Field(default=None, description="回购或争议标的金额")
    agreed_repurchase_rate: Optional[float] = Field(default=None, description="合同约定年化回购溢价/违约金率")
    first_instance_court: str = Field(..., description="作出一审判决之法院名称（必填，无默认）")
    first_instance_errors: Optional[List[str]] = Field(default=None, description="一审争议与错误倾向")
    use_demo: bool = Field(default=False, description="为 true 时载入 data/demo/appeal_demo.json")
    appellant: Optional[PartyInfo] = Field(default=None, description="上诉人完整身份信息")
    appellee: Optional[PartyInfo] = Field(default=None, description="被上诉人完整身份信息")
    first_instance_case_no: Optional[str] = Field(default=None, description="原审案号")
    cause_of_action: Optional[str] = Field(default=None, description="案由")
    appellate_court: Optional[str] = Field(default=None, description="二审法院；未填则文书写待填写上一级法院")
    judgment_type: Optional[Literal["judgment", "ruling"]] = Field(default=None, description="判决或裁定")
    service_date: Optional[str] = Field(default=None, description="送达日期 YYYY-MM-DD")
    include_preservation_insurance: bool = Field(
        default=False,
        description="用户明确提供时才写入保全责任险保费诉请",
    )


class AppealExportDocxRequest(BaseModel):
    title: str = Field(default="二审上诉文书", description="文档标题")
    content_text: str = Field(..., description="文书正文文本")
    doc_type: str = Field(default="appeal_pleading", description="appeal_pleading 或 appeal_memo")


@router.post("/api/v1/litigation/appeal-audit", summary="二审/再审一审判决漏洞剖析与最高院类案分歧穿透")
@router.post("/api/v1/phase1/appeal-audit", summary="二审/再审一审判决漏洞剖析与最高院类案分歧穿透（别名）")
async def api_appeal_audit(req: AppealAuditRequest):
    try:
        appeal_deadline = None
        if req.service_date and req.judgment_type:
            try:
                appeal_deadline = calculate_appeal_deadline(
                    req.service_date, req.judgment_type
                ).get("deadline")
            except ValueError:
                appeal_deadline = None
        res = analyze_first_instance_judgment_and_precedents(
            judgment_text=req.judgment_text or "",
            case_title=req.case_title,
            claim_amount=req.claim_amount,
            agreed_repurchase_rate=req.agreed_repurchase_rate,
            first_instance_court=req.first_instance_court,
            first_instance_errors=req.first_instance_errors,
            use_demo=req.use_demo,
            appellant=req.appellant.model_dump() if req.appellant else None,
            appellee=req.appellee.model_dump() if req.appellee else None,
            first_instance_case_no=req.first_instance_case_no,
            cause_of_action=req.cause_of_action,
            appellate_court=req.appellate_court,
            judgment_type=req.judgment_type,
            service_date=req.service_date,
            include_preservation_insurance=req.include_preservation_insurance,
            appeal_deadline=appeal_deadline,
        )
        if appeal_deadline:
            res["appeal_deadline"] = appeal_deadline
        return res
    except Exception as e:
        logger.error(f"二审类案审计失败: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/api/v1/litigation/appeal-export-docx", summary="二审上诉状与类案检索报告Word原生导出")
@router.post("/api/v1/phase1/appeal-export-docx", summary="二审上诉状与类案检索报告Word原生导出（别名）")
async def api_appeal_export_docx(req: AppealExportDocxRequest):
    try:
        import docx
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = docx.Document()
        
        # 页面边距设置 (司法公文标准)
        for s in doc.sections:
            s.top_margin = Inches(1.0)
            s.bottom_margin = Inches(1.0)
            s.left_margin = Inches(1.1)
            s.right_margin = Inches(1.1)

        # 标题
        h = doc.add_heading(level=1)
        h_run = h.add_run(req.title)
        h_run.font.name = "黑体"
        h_run.font.size = Pt(17)
        h_run.font.bold = True
        h_run.font.color.rgb = RGBColor(0x14, 0x23, 0x3C)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 段落渲染
        lines = req.content_text.splitlines()
        for line in lines:
            line_str = line.strip()
            if not line_str:
                doc.add_paragraph()
                continue
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.35
            p.paragraph_format.space_after = Pt(4)
            p_run = p.add_run(line)
            p_run.font.name = "仿宋_GB2312"
            p_run.font.size = Pt(12)
            if any(line.startswith(prefix) for prefix in ["一、", "二、", "三、", "四、", "【", "上诉请求", "事实与理由", "民 事 上 诉 状"]):
                p_run.font.bold = True
                p_run.font.color.rgb = RGBColor(0x14, 0x23, 0x3C)
                if line.startswith("【") or line == "民 事 上 诉 状":
                    p.paragraph_format.space_before = Pt(8)

        safe_title = sanitize_filename(req.title)
        filename = f"{safe_title}_{datetime.now().strftime('%Y%m%d%H%M%S')}_{uuid.uuid4().hex[:8]}.docx"
        filepath = safe_join(str(get_generated_dir()), filename)
        doc.save(filepath)

        return {
            "status": "success",
            "filename": filename,
            "download_url": f"/api/documents/download/{filename}",
            "file_size": os.path.getsize(filepath)
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"导出二审文书Docx失败: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))

