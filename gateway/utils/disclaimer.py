# -*- coding: utf-8 -*-
"""LawFlow 智律工作台 - 统一免责声明。"""

from __future__ import annotations

from typing import Any, Dict, MutableMapping

DISCLAIMER = "本输出由软件自动生成，仅供执业律师参考，不构成法律意见，须经律师审核后使用。"


def add_disclaimer(result: dict) -> dict:
    """在结果顶层写入 disclaimer 字段（原地修改并返回）。"""
    if not isinstance(result, MutableMapping):
        return result
    result["disclaimer"] = DISCLAIMER
    return result


def ensure_disclaimer_on_payload(payload: Any) -> Any:
    """若 payload 为 dict，则注入 disclaimer；其它类型原样返回。"""
    if isinstance(payload, dict):
        return add_disclaimer(payload)
    return payload


def add_disclaimer_to_docx(docx_path: str) -> None:
    """向 docx 各节页脚写入 DISCLAIMER（python-docx）。"""
    import docx
    from docx.shared import Pt
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    doc = docx.Document(docx_path)
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        if footer.paragraphs:
            p = footer.paragraphs[0]
            for run in list(p.runs):
                run.text = ""
        else:
            p = footer.add_paragraph()
        # 若首段已无 runs，新增
        if not p.runs:
            run = p.add_run(DISCLAIMER)
        else:
            p.runs[0].text = DISCLAIMER
            run = p.runs[0]
        run.font.size = Pt(7)
        run.font.name = "宋体"
        rPr = run._r.get_or_add_rPr()
        rFonts = parse_xml(
            f'<w:rFonts {nsdecls("w")} w:ascii="宋体" w:hAnsi="宋体" w:eastAsia="宋体"/>'
        )
        rPr.append(rFonts)
        for extra in footer.paragraphs[1:]:
            for run in list(extra.runs):
                run.text = ""
    doc.save(docx_path)


def disclaimer_pdf_footer_line(page_no: int = 0, total_pages: int = 0) -> str:
    """PDF 页脚文案（可附页码）。"""
    if total_pages > 0:
        return f"{DISCLAIMER} · 第 {page_no} 页 / 共 {total_pages} 页"
    return DISCLAIMER


def stamp_disclaimer_on_pdf(pdf_path: str) -> None:
    """在 PDF 每页底部写入 DISCLAIMER（若已含全文则跳过）。"""
    try:
        import pymupdf as fitz
    except ImportError:
        return
    if not pdf_path:
        return
    import os
    if not os.path.isfile(pdf_path):
        return
    try:
        doc = fitz.open(pdf_path)
        total = doc.page_count
        for pno in range(total):
            page = doc[pno]
            existing = page.get_text() or ""
            if DISCLAIMER in existing:
                continue
            rect = page.rect
            footer_rect = fitz.Rect(36, rect.height - 42, rect.width - 36, rect.height - 10)
            page.insert_textbox(
                footer_rect,
                disclaimer_pdf_footer_line(page_no=pno + 1, total_pages=total),
                fontsize=6.5,
                fontname="china-s",
                align=1,
            )
        doc.saveIncr()
        doc.close()
    except Exception as doc_err:
        logger.warning(f"添加 PDF 免责声明水印失败: {doc_err}")
        try:
            doc.close()
        except Exception as close_err:
            logger.debug(f"关闭 PDF 失败: {close_err}")
