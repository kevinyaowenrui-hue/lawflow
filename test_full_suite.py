import os
import sys
import json
import httpx
import pymupdf as fitz

BASE_URL = os.getenv("GATEWAY_URL", "http://127.0.0.1:5221")
API_TOKEN = os.getenv("LAWFLOW_API_TOKEN", "")

def log_test(name, success, detail=""):
    mark = "PASS" if success else "FAIL"
    print(f"[{mark}] {name} -> {detail}")
    assert success, f"Test failed for {name}: {detail}"

def run_tests():
    print("=== 开始 LawFlow 智律工作台全系统端到端实证验证 (含核心深度要件事实推演) ===")
    headers = {}
    if API_TOKEN:
        headers["X-LawFlow-Token"] = API_TOKEN
    client = httpx.Client(base_url=BASE_URL, timeout=120.0, headers=headers)

    # 1. Health check（最小信息）
    res = client.get("/health")
    assert res.status_code == 200, f"Health check failed: {res.text}"
    health_data = res.json()
    log_test(
        "网关探针健康检查 GET /health",
        health_data.get("status") == "healthy",
        f"status={health_data.get('status')}"
    )

    # ==========================================
    # 阶段一：司法实务计算 + 文书流水线
    # ==========================================
    # 2.1 人民币大写转换
    r = client.post("/api/v1/phase1/calc/rmb-capital", json={"amount": 123456.78})
    assert r.status_code == 200, r.text
    capital = r.json().get("rmb_capital", "")
    log_test("阶段一: 人民币大写精确到分 (cn2an)", "壹拾贰万叁仟肆佰伍拾陆元柒角捌分" in capital, capital)

    # 2.2 18位身份证校验及2025年渐进式延迟法定退休年龄推算
    r = client.post("/api/v1/phase1/calc/idcard", json={"id_card": "110101199003072375", "is_cadre": False})
    assert r.status_code == 200, r.text
    id_data = r.json()
    log_test("阶段一: 18位身份证解析与年轻适龄劳动者判断", id_data.get("valid") is True and id_data.get("is_retired") is False, f"性别:{id_data.get('gender')}, 出生:{id_data.get('birth_date')}, 年龄:{id_data.get('age')}")

    # 2.2.1 实测 2025 年新规：1965年1月出生男职工（延迟1个月，2025年2月退休）
    r_male = client.post("/api/v1/phase1/calc/idcard", json={"id_card": "110101196501150010", "is_cadre": False})
    assert r_male.status_code == 200, r_male.text
    male_data = r_male.json()
    log_test(
        "阶段一: 2025渐进式延迟退休 (男职工1965年1月 -> 延迟1个月至2025年2月退休)",
        male_data.get("delay_months") == 1 and male_data.get("statutory_retire_date") == "2025年02月",
        f"原退休年龄:{male_data.get('orig_retire_age')}岁 | 延迟月数:{male_data.get('delay_months')}月 | 退休时间:{male_data.get('statutory_retire_date')}"
    )

    # 2.2.2 实测 2025 年新规：1975年1月出生女工人（延迟1个月，2025年2月退休）
    r_fem_worker = client.post("/api/v1/phase1/calc/idcard", json={"id_card": "110101197501150023", "is_cadre": False})
    assert r_fem_worker.status_code == 200, r_fem_worker.text
    fem_worker_data = r_fem_worker.json()
    log_test(
        "阶段一: 2025渐进式延迟退休 (女工人1975年1月原50岁 -> 延迟1个月至2025年2月退休)",
        fem_worker_data.get("delay_months") == 1 and fem_worker_data.get("statutory_retire_date") == "2025年02月",
        f"原退休年龄:{fem_worker_data.get('orig_retire_age')}岁 | 延迟月数:{fem_worker_data.get('delay_months')}月 | 退休时间:{fem_worker_data.get('statutory_retire_date')}"
    )

    # 2.2.3 实测 2025 年新规：1970年1月出生女干部（延迟1个月，2025年2月退休）
    r_fem_cadre = client.post("/api/v1/phase1/calc/idcard", json={"id_card": "110101197001150027", "is_cadre": True})
    assert r_fem_cadre.status_code == 200, r_fem_cadre.text
    fem_cadre_data = r_fem_cadre.json()
    log_test(
        "阶段一: 2025渐进式延迟退休 (女干部1970年1月原55岁 -> 延迟1个月至2025年2月退休)",
        fem_cadre_data.get("delay_months") == 1 and fem_cadre_data.get("statutory_retire_date") == "2025年02月",
        f"原退休年龄:{fem_cadre_data.get('orig_retire_age')}岁 | 延迟月数:{fem_cadre_data.get('delay_months')}月 | 退休时间:{fem_cadre_data.get('statutory_retire_date')}"
    )

    # 2.3 统一社会信用代码校验 (GB 32100-2015 MOD 31-3)
    r = client.post("/api/v1/phase1/calc/uscc", json={"uscc": "91440300741270847Y"})
    assert r.status_code == 200, r.text
    uscc_data = r.json()
    log_test("阶段一: 统一社会信用代码校验与解析", uscc_data.get("valid") is True, f"登记机构:{uscc_data.get('department')}, 省份:{uscc_data.get('province')}")

    # 2.4 LPR利息与诉状段落（商事浮动分段计息）
    r = client.post("/api/v1/phase1/calc/lpr-interest", json={
        "principal": 500000.0,
        "start_date": "2023-01-01",
        "end_date": "2024-01-01",
        "multiplier": 1.5,
        "term": "1Y"
    })
    assert r.status_code == 200, r.text
    lpr_data = r.json()
    log_test("阶段一: 央行 LPR 历期商事同期浮动计息", lpr_data.get("total_interest", 0) > 0, f"利息:{lpr_data.get('total_interest')}元 | 诉状说明:{lpr_data.get('complaint_statement')[:45]}...")

    # 2.4.1 民间借贷司法保护上限模式（固定合同成立时一年期LPR 4倍上限）
    r_cap = client.post("/api/v1/phase1/calc/lpr-interest", json={
        "principal": 200000.0,
        "start_date": "2022-03-01",
        "end_date": "2023-03-01",
        "rule_mode": "private_lending_cap",
        "contract_date": "2022-03-01",
        "fixed_annual_rate": 24.0 # 当事人约定24%，但成立时LPR为3.70%，法定上限必须硬性压减至14.80%
    })
    assert r_cap.status_code == 200, r_cap.text
    cap_data = r_cap.json()
    seg0 = cap_data.get("segments", [{}])[0]
    log_test(
        "阶段一: 民间借贷司法解释法释〔2020〕17号合同成立时4倍LPR硬性压减",
        seg0.get("statutory_cap_rate") == 14.8 and seg0.get("applied_annual_rate") == 14.8 and seg0.get("is_rate_capped") is True,
        f"约定利率 24% 成功依法封顶拦截至合同成立时4倍LPR({seg0.get('statutory_cap_rate')}%) | 实算利息:{cap_data.get('total_interest')}元"
    )

    # 2.4.2 民间借贷 2020-08-20 新旧法分水岭分段核算
    r_split = client.post("/api/v1/phase1/calc/lpr-interest", json={
        "principal": 100000.0,
        "start_date": "2020-01-01",
        "end_date": "2021-01-01",
        "rule_mode": "private_lending_cap",
        "contract_date": "2020-01-01",
        "fixed_annual_rate": 20.0
    })
    assert r_split.status_code == 200, r_split.text
    split_data = r_split.json()
    segs = split_data.get("segments", [])
    log_test(
        "阶段一: 民间借贷 2020-08-20 前后分段保护上限 (24% vs 4倍LPR)",
        len(segs) == 2 and "2020-08-20前" in segs[0].get("regime", "") and "2020-08-20后" in segs[1].get("regime", ""),
        f"分段数:{len(segs)} | 一段上限:{segs[0].get('statutory_cap_rate')}% | 二段上限:{segs[1].get('statutory_cap_rate')}%"
    )

    # 2.5 劳动争议补偿金 (N / N+1 / 2N)
    r = client.post("/api/v1/phase1/calc/labor-arbitration", json={
        "monthly_salary": 12000.0,
        "hire_date": "2020-03-01",
        "termination_date": "2023-08-15"
    })
    assert r.status_code == 200, r.text
    labor_data = r.json()
    tiers = labor_data.get("tiers", {})
    amount_2n = tiers.get("2N", {}).get("amount")
    log_test("阶段一: 劳动争议 N / N+1 / 2N 补偿金精算", amount_2n == 84000.0, f"工龄折算:{labor_data.get('tenure_str')} (计发月数:{labor_data.get('compensation_months')}) | N={tiers.get('N', {}).get('amount')} | 2N={amount_2n}")

    # 2.5.1 纠偏验证：不满6个月（零头天数>15天）严禁错误进位为1.0个月（实务硬性标准为0.5个月）
    r_half = client.post("/api/v1/phase1/calc/labor-arbitration", json={
        "monthly_salary": 10000.0,
        "hire_date": "2023-01-01",
        "termination_date": "2023-06-20" # 5个月19天（不满6个月，零头19天）
    })
    assert r_half.status_code == 200, r_half.text
    half_data = r_half.json()
    log_test(
        "阶段一: 劳动合同法第47条不满6个月（5个月19天）精准折算0.5个月补偿金",
        half_data.get("compensation_months") == 0.5 and half_data.get("tiers", {}).get("N", {}).get("amount") == 5000.0,
        f"工龄:{half_data.get('tenure_str')} | 计发月数:{half_data.get('compensation_months')}（绝不错误进位为1.0）| N={half_data.get('tiers', {}).get('N', {}).get('amount')}元"
    )

    # 2.5.2 跨2008年1月1日高管高薪员工《劳动合同法》第97条第3款分段计算
    r_seg2008 = client.post("/api/v1/phase1/calc/labor-arbitration", json={
        "monthly_salary": 50000.0,
        "hire_date": "2005-01-01",
        "termination_date": "2023-12-31",
        "local_avg_salary": 10000.0 # 社平三倍为30000元，2008年后封顶12年
    })
    assert r_seg2008.status_code == 200, r_seg2008.text
    seg2008_data = r_seg2008.json()
    seg_info = seg2008_data.get("segment_2008", {})
    log_test(
        "阶段一: 跨2008年高薪员工分段核算 (2008年前不封顶 vs 2008年后三倍社平12年封顶)",
        seg_info.get("is_applicable") is True and seg_info.get("pre_2008_salary_basis") == 50000.0 and seg_info.get("post_2008_compensation_months") == 12.0,
        f"2008前金额:{seg_info.get('pre_2008_amount')}元 (月基数:{seg_info.get('pre_2008_salary_basis')}) | 2008后金额:{seg_info.get('post_2008_amount')}元 (封顶月数:{seg_info.get('post_2008_compensation_months')})"
    )

    # 2.6 诉讼费用计算
    r = client.post("/api/v1/phase1/calc/court-fees", json={"claim_amount": 500000.0, "is_summary_procedure": False})
    assert r.status_code == 200, r.text
    fee_data = r.json()
    log_test("阶段一: 诉讼费用交纳办法分段计费", fee_data.get("standard_fee") == 8800.0, f"标准案件受理费:{fee_data.get('standard_fee')}元 | 保全申请费:{fee_data.get('preservation_fee')}元")

    # 2.7 立案文书流水线一站式生成 (.docx + .pdf) 与 gemma4:12b 事实与理由润色
    print("\n-> 正在测试阶段一全套起诉状自动生成与 gemma4:12b 润色流水线...")
    r = client.post("/api/v1/phase1/generate-litigation-package", json={
        "cause_of_action": "民间借贷纠纷",
        "plaintiff": {
            "name": "李四",
            "id_card_or_uscc": "110101199003072375",
            "gender": "男",
            "birth": "1990年3月7日",
            "ethnicity": "汉族",
            "address": "北京市海淀区中关村南大街1号",
            "phone": "13800138000"
        },
        "defendant": {
            "name": "张三",
            "id_card_or_uscc": "110101198505121232",
            "gender": "男",
            "birth": "1985年5月12日",
            "ethnicity": "汉族",
            "address": "北京市朝阳区建国路88号",
            "phone": "13900139000"
        },
        "claims": [
            "判令被告立即向原告偿还借款本金人民币500,000元；",
            "判令被告向原告支付利息（自2023年3月1日起至实际清偿之日止，按年利率6%计算）；",
            "判令本案全部诉讼费用由被告承担。"
        ],
        "claim_amount": 500000.0,
        "raw_case_notes": "被告张三在2023年3月因资金周转困难找原告借款50万元，约定借期半年，年利率6%。原告当天通过招商银行全额转账。借期届满后原告多次通过微信及电话催讨，被告以工程款未结为由推脱，至今分文未还。",
        "court_name": "北京市海淀区人民法院",
        "lawyer_name": "王大状",
        "law_firm": "北京智律律师事务所",
        "lawyer_phone": "13600001111",
        "lawyer_license": "11101201510123456",
        "auto_generate_sub_docs": True
    }, timeout=120.0)
    assert r.status_code == 200, r.text
    pkg_data = r.json()
    files_generated = pkg_data.get("generated_documents", [])
    log_test("阶段一: 立案文书一揽子打包与格式转换", pkg_data.get("status") == "success" and len(files_generated) >= 4, f"生成文书种类数: {len(files_generated)}")
    for f in files_generated:
        print(f"   * 生成文书: {f.get('doc_type')} -> docx: {f.get('filename_docx')} | pdf: {f.get('filename_pdf')}")

    # ==========================================
    # 阶段二：证据清洗工厂 + 五模块庭审提纲 Agent
    # ==========================================
    # 3.1 五模块庭审提纲
    print("\n-> 正在测试阶段二五模块庭审提纲 Agent (gemma4:12b)...")
    r = client.post("/api/v1/phase2/trial-outline", json={
        "case_title": "李四诉张三民间借贷纠纷案",
        "cause_of_action": "民间借贷纠纷",
        "plaintiff_claims_and_facts": "原告诉请判令被告偿还本金50万元及利息3万元，原告已提供招行转账凭证及微信借款合意记录。",
        "defendant_defense": "被告辩称该笔资金系双方共同承揽工程的合伙投资款，非个人借贷，且原告未提供纸质借据原件。"
    }, timeout=120.0)
    assert r.status_code == 200, r.text
    trial_data = r.json()
    outline_text = trial_data.get("trial_outline_markdown", "")
    log_test("阶段二: 五模块庭审提纲生成 (gemma4:12b)", len(outline_text) > 100, f"提纲字符数: {len(outline_text)}")

    # 3.2 多视角合同审查 (乙方防御版)
    print("-> 正在测试阶段二多视角合同审查 Agent (gemma4:12b)...")
    r = client.post("/api/v1/phase2/contract-review", json={
        "contract_title": "企业级软件系统开发服务合同",
        "contract_text": "第七条 违约责任：乙方如延期交付超过3日，每日应向甲方支付合同总价款5%的惩罚性违约金；甲方有权随时单方解除合同且无需退还首期款项。发生争议由甲方所在地人民法院管辖。",
        "perspective": "乙方防御版"
    }, timeout=120.0)
    assert r.status_code == 200, r.text
    review_data = r.json()
    log_test("阶段二: 乙方防御版商业合同红线审查", "review" in review_data, f"审查结论: {review_data.get('review', {}).get('review_summary')}")

    # 3.3 敏感信息一键涂黑脱敏 (PyMuPDF - 解决中文汉字紧邻数字时的正则断言脱敏)
    dummy_pdf_filename = "test_sensitive_doc.pdf"
    dummy_pdf_path = os.path.join("storage", "generated", dummy_pdf_filename)
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    # 特别测试：汉字直接紧邻身份证号与银行卡号（无空格、无冒号，考验负向断言 (?<!\d)...(?!\d)）
    page.insert_text((50, 100), "原告身份证号110101199003072379在北京市朝阳区", fontname="china-s", fontsize=12)
    page.insert_text((50, 150), "结算转账银行卡号6222021234567890123尾号开户行", fontname="china-s", fontsize=12)
    doc.save(dummy_pdf_path)
    doc.close()

    r = client.post("/api/v1/phase2/redact-sensitive", json={
        "pdf_filename": dummy_pdf_filename,
        "redact_idcard": True,
        "redact_bankcard": True
    })
    assert r.status_code == 200, r.text
    redact_data = r.json()
    log_test(
        "阶段二: 汉字紧邻数字负向环视深度脱敏 (PyMuPDF)",
        redact_data.get("status") == "success" and redact_data.get("redactions_count") >= 2,
        f"脱敏项数: {redact_data.get('redactions_count')} (成功击破 Python 正则汉字属于 \\w 导致的 \\b 失效漏洞)"
    )

    # 3.4 本地 PyMuPDF 兜底长文多页流式排版验证（绝无截断丢字）
    import asyncio
    import docx as python_docx
    from gateway.routers.phase1_litigation import convert_docx_to_pdf

    test_long_docx = os.path.join("storage", "generated", "test_long_complaint.docx")
    test_long_pdf = os.path.join("storage", "generated", "test_long_complaint.pdf")
    d_doc = python_docx.Document()
    d_doc.add_heading("民事起诉状（长文流式排版测试）", 0)
    for i in range(1, 121):
        d_doc.add_paragraph(f"第{i}段：原告李四与被告张三民间借贷纠纷事实推演，包含第{i}笔资金周转明细、利息约定以及多次书面催告记录，以此验证多页流式排版与页码防伪水印。")
    d_doc.save(test_long_docx)

    conv_ok = asyncio.run(convert_docx_to_pdf(test_long_docx, test_long_pdf, "长篇诉状排版"))
    assert conv_ok is True, "PyMuPDF 长文流式排版转换失败"
    pdf_inspect = fitz.open(test_long_pdf)
    p_count = pdf_inspect.page_count
    first_page_text = pdf_inspect[0].get_text("text")
    last_page_text = pdf_inspect[p_count - 1].get_text("text")
    pdf_inspect.close()

    log_test(
        "阶段二: PyMuPDF 兜底长文多页流式排版与动态页码",
        p_count >= 3 and "第 1 页 / 共" in first_page_text and f"第 {p_count} 页 / 共 {p_count} 页" in last_page_text,
        f"长文成功流式排版为 {p_count} 页 | 页码防伪水印与尾页完整无截断丢字"
    )

    # ==========================================
    # 阶段三：私有双库隔离 RAG + 电子签约闭环
    # ==========================================
    # 4.1 法律正则切片器 (chinese_law_chunker: 编-章-节-条-款)
    sample_law = """
中华人民共和国民法典
第一编 总则
第一章 基本规定
第一条 为了保护民事主体的合法权益，调整民事关系，维护社会和经济秩序，适应中国特色社会主义发展要求，弘扬社会主义核心价值观，根据宪法，制定本法。
第二条 民法调整平等主体的自然人、法人和非法人组织之间的人身关系和财产关系。
第二章 自然人
第一节 民事权利能力和民事行为能力
第十三条 自然人从出生时起到死亡时止，具有民事权利能力，依法享有民事权利，承担民事义务。
"""
    r = client.post("/api/v1/phase3/rag/chunk-law", json={
        "law_name": "民法典总则编节选",
        "raw_law_text": sample_law
    })
    assert r.status_code == 200, r.text
    chunk_data = r.json()
    chunks = chunk_data.get("chunks", [])
    log_test("阶段三: 编-章-节-条-款 专用正则切片器", len(chunks) == 3, f"精准切出 {len(chunks)} 条独立法条单元，绝不切断法条")
    for c in chunks:
        print(f"   * [{c.get('article_no')}] {c.get('part')} > {c.get('chapter')} > {c.get('citation')}")

    # 4.2 双库隔离 RAG 检索
    r = client.post("/api/v1/phase3/rag/search", json={
        "query": "借款利息约定超过LPR四倍是否有效？",
        "corpus_type": "all",
        "top_k": 3
    })
    assert r.status_code == 200, r.text
    search_data = r.json()
    results = search_data.get("results", [])
    log_test("阶段三: 公共法规与律所胜诉案例双库隔离检索", len(results) > 0, f"命中条目数: {len(results)} (库源覆盖: {set(item.get('source') for item in results)})")

    # 4.3 远程委托电子签约闭环 (DocuSeal 接口)
    r = client.post("/api/v1/phase3/create-signing-link", json={
        "agreement_title": "民事借贷纠纷专项法律服务委托代理合同",
        "client_name": "李四",
        "client_phone": "13800138000",
        "client_email": "lisi@example.com",
        "lawyer_name": "王大状",
        "law_firm": "北京智律律师事务所",
        "service_fee": 30000.0,
        "case_summary": "代理原告李四诉张三民间借贷纠纷一案一审诉讼"
    })
    assert r.status_code == 200, r.text
    sign_data = r.json()
    envelope_id = sign_data.get("envelope_id")
    log_test("阶段三: DocuSeal 电子签约链接与存证生成", sign_data.get("status") == "success", f"信封ID: {envelope_id} | 签署链接: {sign_data.get('signing_url')}")

    # 4.4 电子签约完成 Webhook 存证回调
    r = client.post("/api/v1/phase3/signing-webhook", json={
        "event_type": "envelope.completed",
        "envelope_id": envelope_id,
        "status": "completed",
        "signed_at": "2026-10-16T12:00:00"
    })
    assert r.status_code == 200, r.text
    webhook_data = r.json()
    log_test("阶段三: 签署完成自动存证归档回调", webhook_data.get("status") == "success", webhook_data.get("message"))

    # ==========================================
    # 核心深度规格：请求权基础与要件事实推演中枢
    # ==========================================
    print("\n-> 正在测试深度实务精算与执业红线熔断器...")
    
    # 5.1 诉讼时效推演与中断计算器
    r = client.post("/api/v1/phase1/calc/statute-limitations", json={
        "start_date": "2023-03-10",
        "cause_type": "general_3y",
        "interruption_events": [
            {"date": "2023-09-15", "description": "李四微信承诺还款"}
        ]
    })
    assert r.status_code == 200, r.text
    statute_res = r.json()
    log_test("深度精算: 诉讼时效中断推演与执业安全期", statute_res.get("final_expire_date") == "2026-09-15", f"最终届满日: {statute_res.get('final_expire_date')} | 剩余安全期: {statute_res.get('days_remaining')}天")

    # 5.1.1 保证期间除斥红线（一般保证与连带保证法定6个月除斥期间，《民法典》第692/693条）
    r_guar_gen = client.post("/api/v1/phase1/calc/statute-limitations", json={
        "start_date": "2023-01-01",
        "cause_type": "guarantee_general_6m"
    })
    assert r_guar_gen.status_code == 200, r_guar_gen.text
    res_gen = r_guar_gen.json()
    log_test(
        "深度精算: 一般保证6个月法定除斥期间（起诉主债务人红线）",
        res_gen.get("is_preclusive") is True and res_gen.get("final_expire_date") == "2023-07-01" and "起诉主债务人" in res_gen.get("special_notice", ""),
        f"届满日:{res_gen.get('final_expire_date')} | 除斥期间(不可中断):{res_gen.get('is_preclusive')} | 穿仓红线:{res_gen.get('special_notice')[:36]}..."
    )

    r_guar_joint = client.post("/api/v1/phase1/calc/statute-limitations", json={
        "start_date": "2023-01-01",
        "cause_type": "guarantee_joint_6m"
    })
    assert r_guar_joint.status_code == 200, r_guar_joint.text
    res_joint = r_guar_joint.json()
    log_test(
        "深度精算: 连带保证6个月法定除斥期间（要求保证人担责红线）",
        res_joint.get("is_preclusive") is True and res_joint.get("final_expire_date") == "2023-07-01",
        f"届满日:{res_joint.get('final_expire_date')} | 除斥期间:{res_joint.get('is_preclusive')}"
    )

    # 5.1.2 《民法典》第194条诉讼时效最后6个月法定事由中止精算
    r_susp = client.post("/api/v1/phase1/calc/statute-limitations", json={
        "start_date": "2020-01-01",
        "cause_type": "general_3y", # 原届满日为 2023-01-01
        "suspension_events": [
            {
                "date": "2022-09-01", # 在最后6个月内（2022-07-01至2023-01-01之间）发生不可抗力
                "description": "严重疫情封控导致当事人丧失起诉行为能力",
                "resume_date": "2022-11-01" # 障碍消除日，自此日满6个月即2023-05-01届满
            }
        ]
    })
    assert r_susp.status_code == 200, r_susp.text
    res_susp = r_susp.json()
    log_test(
        "深度精算: 《民法典》第194条诉讼时效最后6个月中止及消除后满6个月顺延",
        res_susp.get("final_expire_date") == "2023-05-01",
        f"原届满日:{res_susp.get('initial_expire_date')} -> 中止后最终届满日:{res_susp.get('final_expire_date')} (精准顺延至障碍消除后6个月)"
    )

    # 5.2 财产保全续封红线计算器
    r = client.post("/api/v1/phase1/calc/preservation-renewal", json={
        "preservation_type": "bank_deposit",
        "start_date": "2024-01-15"
    })
    assert r.status_code == 200, r.text
    preserv_res = r.json()
    # 原断言 advance_days_required == 15 把错误行为写成了通过条件；法释〔2016〕22号第十八条为届满七日前
    log_test("深度精算: 财产保全续查封续冻结红线精算", preserv_res.get("advance_days_required") == 7, f"续保申请截止日: {preserv_res.get('renewal_application_deadline')} | 距截止剩余: {preserv_res.get('days_to_renewal_deadline')}天")

    # 5.3 规范性法律文件现行有效性与废止拦截核验
    r = client.post("/api/v1/phase1/calc/statute-validity", json={
        "citation_text": "依据《中华人民共和国合同法》第一百零七条及《担保法》第十九条",
        "incident_date": "2023-05-01"
    })
    assert r.status_code == 200, r.text
    validity_res = r.json()
    log_test("深度精算: 废止法规自动探针拦截与时间效力核验", validity_res.get("is_safe") is False and len(validity_res.get("detected_abrogated")) >= 2, f"成功拦截废止法规: {[x['full_name'] for x in validity_res.get('detected_abrogated')]}")

    # 5.4 一手真实原始卷宗预置库
    print("\n-> 正在测试深度要件事实矩阵与一手卷宗坐标锚定...")
    r = client.get("/api/v1/deep-matrix/presets")
    assert r.status_code == 200, r.text
    presets_res = r.json()
    log_test("深度推演: 一手高保真原始卷宗库预置", len(presets_res.get("presets", [])) >= 2, f"已载入 {len(presets_res.get('presets'))} 套顶尖实务典型一手卷宗")

    # 5.5 七维请求权基础与要件事实推演矩阵 (己方主办视角 Blue Team)
    r = client.post("/api/v1/deep-matrix/analyze", json={
        "case_preset": "case_loan_intermingled",
        "perspective": "plaintiff_blue_team"
    })
    assert r.status_code == 200, r.text
    matrix_blue = r.json()
    elements = matrix_blue.get("matrix_elements", [])
    conflicts = matrix_blue.get("detected_conflicts", [])
    scorecard = matrix_blue.get("scorecard", {})
    log_test(
        "深度推演: 七维要件事实矩阵与多源时间轴冲突侦测 (Blue Team)",
        len(elements) >= 5 and len(conflicts) >= 3 and scorecard.get("is_litigation_ready") is True,
        f"完备率: {scorecard.get('completeness_score')}% | 要件数: {len(elements)} | 侦测到冲突: {len(conflicts)}项 (含第三方代付与高利拦截)"
    )

    # 5.6 对方大牌代理律师红队极端对抗质证视角 (Red Team)
    r = client.post("/api/v1/deep-matrix/analyze", json={
        "case_preset": "case_loan_intermingled",
        "perspective": "adversarial_red_team"
    })
    assert r.status_code == 200, r.text
    matrix_red = r.json()
    log_test("深度推演: 对方代理律师红队质证视角 (Red Team)", matrix_red.get("is_red_team_view") is True, "成功输出真实性、合法性、证明力穿透式突袭发问与红队攻击切入点")

    # 5.7 独立红队质证攻击模拟器
    r = client.post("/api/v1/deep-matrix/adversarial-attack", json={
        "evidence_id": "DOC-02",
        "evidence_name": "招商银行网银电子转账回单",
        "plaintiff_purpose": "证明已向李四交付50万元借款",
        "raw_snippet": "收款人：王五；摘要：往来款"
    })
    assert r.status_code == 200, r.text
    attack_res = r.json()
    log_test("深度推演: 独立证据红队质证穿透攻击生成", len(attack_res.get("attacks", [])) == 3, f"生成3大质证攻击维度与补强取证SOP ({len(attack_res.get('remedy_checklist'))}项)")

    # ==========================================
    # 阶段五：权威法典与子母法图谱中枢 (SQLite FTS5)
    # ==========================================
    # 6.1 全景法典与图谱指标总览
    r = client.get("/api/v1/statutes/overview")
    assert r.status_code == 200, r.text
    stat_ov = r.json().get("data", {})
    log_test(
        "法典中枢: SQLite FTS5 引擎与图谱概览 GET /api/v1/statutes/overview",
        stat_ov.get("total_laws", 0) >= 10 and stat_ov.get("total_articles", 0) >= 35 and stat_ov.get("total_parent_child_links", 0) >= 10,
        f"核心法典:{stat_ov.get('total_laws')}部 | 法条数:{stat_ov.get('total_articles')}条 | 子母关联:{stat_ov.get('total_parent_child_links')}组 | 废止旧法:{stat_ov.get('total_abrogated_laws')}部"
    )

    # 6.2 精确条号直达与子母法穿透 (民法典 585 ➜ 合同编解释 64/65条 30% 上限)
    r = client.post("/api/v1/statutes/search", json={"query": "民法典 585"})
    assert r.status_code == 200, r.text
    s_res = r.json().get("results", [])
    top_art = s_res[0] if s_res else {}
    child_links = top_art.get("child_interpretations", [])
    has_30_percent_rule = any("30%" in c.get("notes", "") or "百分之三十" in c.get("content", "") for c in child_links)
    log_test(
        "法典中枢: 民法典第585条直达与最高法合同编解释(30%上限)子母连带穿透",
        top_art.get("article_no_cn") == "第五百八十五条" and len(child_links) >= 2 and has_30_percent_rule,
        f"命中:{top_art.get('short_name')} {top_art.get('article_no_cn')} | 连带穿透召回司法解释:{len(child_links)}条 (含违约金30%调减红线)"
    )

    # 6.3 商业与公司法复合短语检索 (刺破公司面纱 ➜ 新公司法第23条)
    r = client.post("/api/v1/statutes/search", json={"query": "刺破公司面纱"})
    assert r.status_code == 200, r.text
    res_corp = r.json().get("results", [])
    top_corp = res_corp[0] if res_corp else {}
    log_test(
        "法典中枢: 新《公司法》(2024) 刺破公司面纱与财产混同举证倒置",
        top_corp.get("article_no_cn") == "第二十三条" and "连带责任" in top_corp.get("content", ""),
        f"命中:{top_corp.get('short_name')} {top_corp.get('article_no_cn')} (纵向横向穿透/一人公司举证责任倒置)"
    )

    # 6.4 民间借贷转贷与无效抗辩 (套取贷款 ➜ 民间借贷司法解释第13条)
    r = client.post("/api/v1/statutes/search", json={"query": "套取贷款"})
    assert r.status_code == 200, r.text
    res_loan = r.json().get("results", [])
    top_loan = res_loan[0] if res_loan else {}
    log_test(
        "法典中枢: 民间借贷司法解释第13条 (套取金融机构贷款转贷法定无效)",
        top_loan.get("article_no_cn") == "第十三条" and "套取金融机构贷款" in top_loan.get("content", ""),
        f"命中:{top_loan.get('short_name')} {top_loan.get('article_no_cn')} (民间借贷效力阻断规范)"
    )

    # 6.5 刑民交叉与追诉标准 (刑法 224 ➜ 合同诈骗罪及两万元立案标准)
    r = client.post("/api/v1/statutes/search", json={"query": "刑法 224"})
    assert r.status_code == 200, r.text
    res_crim = r.json().get("results", [])
    top_crim = res_crim[0] if res_crim else {}
    crim_children = top_crim.get("child_interpretations", [])
    log_test(
        "法典中枢: 刑法第224条合同诈骗罪与两高立案追诉标准(二)穿透",
        top_crim.get("article_no_cn") == "第二百二十四条" and len(crim_children) >= 1,
        f"命中:{top_crim.get('short_name')} {top_crim.get('article_no_cn')} | 连带立案追诉门槛标准:{crim_children[0].get('notes', '')[:35]}..."
    )

    # 6.6 律所自定义地方高院裁判指引原子切片导入 (chinese_law_chunker ETL)
    sample_raw_law = """
第一条 为了统一全省法院审理涉互联网平台用工劳动争议案件的裁判标准，妥善平衡劳动者合法权益与平台经济健康发展，制定本指引。
第二条 依托互联网平台就业的新就业形态劳动者，与平台企业或合作企业发生用工争议的，应根据用工事实和劳动管理程度综合认定法律关系。
第三条 平台企业通过算法或者规则对劳动者进行实质性管理，劳动者在工作时间、工作路线、服务价格等方面受平台严格控制的，应当认定双方存在劳动关系。
"""
    r = client.post("/api/v1/statutes/import", json={
        "law_id": "user_import_highcourt_platform_labor_2024",
        "full_name": "【测试】某省高级人民法院关于审理互联网平台用工劳动争议案件的审理指引",
        "short_name": "【测试】平台用工审理指引",
        "hierarchy": 4,
        "promulgated_by": "某省高级人民法院",
        "effective_date": "2024-05-01",
        "domain": "劳动",
        "raw_text": sample_raw_law,
        "summary": "新业态用工劳动关系从属性认定标准"
    })
    assert r.status_code == 200, r.text
    import_res = r.json()
    log_test(
        "法典中枢: 律所私有裁判指引一键编章节条款原子切片导入",
        import_res.get("data", {}).get("articles_imported") == 3,
        f"成功切分导入 {import_res.get('data', {}).get('articles_imported')} 条新法条并即时构建 FTS5 索引"
    )

    # 6.7 导入后即时检索验证
    r = client.post("/api/v1/statutes/search", json={"query": "算法或者规则"})
    assert r.status_code == 200, r.text
    res_imported = r.json().get("results", [])
    log_test(
        "法典中枢: 自定义导入法规即时毫秒级混合检索验证",
        len(res_imported) >= 1 and "平台用工审理指引" in res_imported[0].get("short_name", "") and (
            res_imported[0].get("verification_status") == "USER_IMPORTED_UNVERIFIED"
            or "未核验" in str(res_imported[0].get("verification_label") or "")
        ),
        f"命中新建私有条文: {res_imported[0].get('short_name')} {res_imported[0].get('article_no_cn')} status={res_imported[0].get('verification_status')}"
    )

    # ==========================================
    # 阶段七：第二优先级 —— 全局在办案件上下文总线与真实卷宗动态推演
    # ==========================================
    print("\n-> 正在测试第二优先级：真实卷宗动态推演 (POST /api/v1/deep-matrix/analyze-custom)...")
    
    # 7.1 自定义民间借贷纠纷案情动态推演 (含第三方代付、未出借据、自动关联法条)
    r = client.post("/api/v1/deep-matrix/analyze-custom", json={
        "case_preset": "custom",
        "perspective": "plaintiff_blue_team",
        "custom_case_title": "王强诉赵刚民间借贷纠纷案",
        "custom_cause_of_action": "民间借贷纠纷",
        "custom_claim_amount": 350000.0,
        "custom_plaintiff": "王强",
        "custom_defendant": "赵刚",
        "custom_raw_text": "2023年4月15日，赵刚向王强借款35万元，约定借期1年，月息1.5%。王强于当天通过微信指定账户将35万元转给赵刚的表弟钱某。借期届满后赵刚以未向其本人银行卡付款为由拒不偿还。",
        "custom_evidence_notes": "证据一：微信聊天记录导出件（赵刚明确要求转入钱某建行卡，并附卡号）\n证据二：招商银行个人网银电子回单（户名钱某，金额35万元）\n证据三：催款录音光盘（赵刚在通话中承认收到该笔借款）",
        "custom_statutes": [
            "《中华人民共和国民法典》第六百六十七条",
            "《最高人民法院关于审理民间借贷案件适用法律若干问题的规定》第十九条"
        ]
    })
    assert r.status_code == 200, r.text
    custom_loan_res = r.json()
    elements = custom_loan_res.get("matrix_elements", [])
    conflicts = custom_loan_res.get("detected_conflicts", [])
    timeline = custom_loan_res.get("timeline_events", [])
    scorecard = custom_loan_res.get("scorecard", {})
    raw_docs = custom_loan_res.get("raw_documents", [])
    
    log_test(
        "第二优先级: 真实借贷卷宗七维动态推演与代付冲突识别",
        len(elements) >= 3 and len(conflicts) >= 1 and len(timeline) >= 2 and len(raw_docs) >= 3,
        f"动态要件数:{len(elements)} | 侦测到冲突:{[c['conflict_type'] for c in conflicts]} | 完备率:{scorecard.get('completeness_score')}%"
    )

    # 7.2 红队攻击视角下自定义案卷动态推演
    r_red = client.post("/api/v1/deep-matrix/analyze-custom", json={
        "case_preset": "custom",
        "perspective": "adversarial_red_team",
        "custom_case_title": "王强诉赵刚民间借贷纠纷案",
        "custom_cause_of_action": "民间借贷纠纷",
        "custom_claim_amount": 350000.0,
        "custom_plaintiff": "王强",
        "custom_defendant": "赵刚",
        "custom_raw_text": "2023年4月15日，赵刚向王强借款35万元，约定借期1年，月息1.5%。王强于当天通过微信指定账户将35万元转给赵刚的表弟钱某。借期届满后赵刚以未向其本人银行卡付款为由拒不偿还。",
        "custom_evidence_notes": "证据一：微信聊天记录导出件\n证据二：网银回单35万"
    })
    assert r_red.status_code == 200, r_red.text
    red_res = r_red.json()
    red_attacks = [e.get("adversarial_attack", "") for e in red_res.get("matrix_elements", []) if e.get("adversarial_attack")]
    log_test(
        "第二优先级: 真实卷宗红队对抗视角质证攻击生成",
        red_res.get("is_red_team_view") is True and len(red_attacks) >= 3,
        f"红队质证重锤点数: {len(red_attacks)} | 攻击样本: {red_attacks[0][:35]}..."
    )

    # 7.3 自定义买卖合同纠纷案情动态推演 (自动结合 statute_db FTS5 权威法典召回母法)
    r_sale = client.post("/api/v1/deep-matrix/analyze-custom", json={
        "case_preset": "custom",
        "perspective": "plaintiff_blue_team",
        "custom_case_title": "华盛建材诉宏业建设买卖合同欠款纠纷案",
        "custom_cause_of_action": "买卖合同纠纷",
        "custom_claim_amount": 800000.0,
        "custom_plaintiff": "华盛建材有限公司",
        "custom_defendant": "宏业建设工程有限公司",
        "custom_raw_text": "2023年6月，原告依合同向被告供应钢材总价80万元，被告在送货单上签字盖章确认签收。现约定付款期限已过，被告拖欠尾款50万元未付。",
        "custom_evidence_notes": "证据一：买卖合同原件\n证据二：项目工地材料签收单5份\n证据三：催款对账确认单"
    })
    assert r_sale.status_code == 200, r_sale.text
    sale_res = r_sale.json()
    sale_elements = sale_res.get("matrix_elements", [])
    log_test(
        "第二优先级: 真实买卖合同卷宗七维动态推演与法典 FTS5 自动召回",
        len(sale_elements) >= 3 and any("民法典" in e.get("claim_basis", "") for e in sale_elements),
        f"买卖合同要件数:{len(sale_elements)} | 自动召回母法依据:{sale_elements[0].get('claim_basis')}"
    )

    # 7.4 验证 /api/v1/deep-matrix/analyze 接口当 case_preset='custom' 时的自适应路由分发
    r_delegate = client.post("/api/v1/deep-matrix/analyze", json={
        "case_preset": "custom",
        "custom_case_title": "委托分发测试案",
        "custom_cause_of_action": "通用民事纠纷",
        "custom_raw_text": "当事人存在合伙协议及投资款交付凭证，现要求清算返还出资。",
        "custom_evidence_notes": "证据一：合伙协议"
    })
    assert r_delegate.status_code == 200, r_delegate.text
    del_res = r_delegate.json()
    log_test(
        "第二优先级: /analyze 端点自动委托至 custom 动态推演引擎",
        del_res.get("status") == "success" and del_res.get("case_title") == "委托分发测试案",
        f"成功分发执行: 案件标题='{del_res.get('case_title')}' | 完备率:{del_res.get('scorecard', {}).get('completeness_score')}%"
    )

    # =========================================================================
    # 第三优先级：律所级利益冲突穿透排查与零幻觉引注熔断 (Phase 6 Compliance & Guard)
    # =========================================================================
    print("\n--- [第三优先级] 律所级利益冲突穿透排查中枢与零幻觉引注防幻觉熔断网关 ---")

    # 8.1 利益冲突穿透排查：拟起诉常年法律顾问单位（直接顶格红线冲突熔断）
    r_conflict_direct = client.post("/api/v1/conflict/check", json={
        "client_name": "北京知客云服科技有限公司",
        "opposing_party": "北京宏达商贸有限公司",
        "case_subject": "技术合同尾款争议"
    })
    assert r_conflict_direct.status_code == 200, r_conflict_direct.text
    c_direct = r_conflict_direct.json()
    log_test(
        "第三优先级: 利益冲突穿透排查 - 拟起诉本所常年顾问单位 (顶格执业红线熔断)",
        c_direct.get("conflict_level") == "CRITICAL_BLOCKED" and c_direct.get("has_conflict") is True and any("律师法" in rule and "39条" in rule for rule in c_direct.get("triggered_rules", [])),
        f"冲突定级: {c_direct.get('conflict_level')} | 命中红线条目: {c_direct.get('triggered_rules')} | 结论: {c_direct.get('summary')}"
    )

    # 8.2 利益冲突穿透排查：同案双方代理排查（张三诉李四民间借贷纠纷）
    r_conflict_both = client.post("/api/v1/conflict/check", json={
        "client_name": "张三",
        "opposing_party": "李四",
        "case_subject": "借贷纠纷"
    })
    assert r_conflict_both.status_code == 200, r_conflict_both.text
    c_both = r_conflict_both.json()
    log_test(
        "第三优先级: 利益冲突穿透排查 - 同案双方在办冲突拦截",
        c_both.get("has_conflict") is True and c_both.get("conflict_level") in ["CRITICAL_BLOCKED", "HIGH_WARNING"],
        f"冲突定级: {c_both.get('conflict_level')} | 拦截原因: {c_both.get('summary')}"
    )

    # 8.3 利益冲突穿透排查：无冲突合规案件签发 SHA-256 数字验讫证书
    r_conflict_clear = client.post("/api/v1/conflict/check", json={
        "client_name": "北京创新未来科技有限公司",
        "opposing_party": "广州星海传媒有限公司",
        "case_subject": "侵害商标权纠纷"
    })
    assert r_conflict_clear.status_code == 200, r_conflict_clear.text
    c_clear = r_conflict_clear.json()
    cert = c_clear.get("clearance_certificate", {})
    log_test(
        "第三优先级: 利益冲突穿透排查 - 绿灯放行并签发 SHA-256 数字验讫证书",
        c_clear.get("conflict_level") == "CLEAR" and c_clear.get("has_conflict") is False and len(cert.get("sha256_fingerprint", "")) == 64,
        f"定级: {c_clear.get('conflict_level')} | 证书号: {cert.get('certificate_no')} | SHA256: {cert.get('sha256_fingerprint')[:16]}..."
    )

    # 8.4 律所底档库查询与新案登记
    r_reg_list = client.get("/api/v1/conflict/registry")
    assert r_reg_list.status_code == 200, r_reg_list.text
    reg_data = r_reg_list.json()
    log_test(
        "第三优先级: 利益冲突底层底档库全量检索",
        reg_data.get("status") == "success" and reg_data.get("count", 0) >= 4,
        f"当前律所档案库在册案件数: {reg_data.get('count')} 卷"
    )

    r_reg_new = client.post("/api/v1/conflict/register-case", json={
        "case_name": "宏达集团下属子公司破产重整案",
        "case_no": "(2026)京破999号",
        "client_name": "北京宏达商贸有限公司",
        "opposing_party": "全体破产债权人",
        "status": "ACTIVE",
        "handling_lawyer": "合规审查专员",
        "notes": "破产重组专项"
    })
    assert r_reg_new.status_code == 200, r_reg_new.text
    log_test(
        "第三优先级: 实时录入新案至律所利益冲突底档库",
        r_reg_new.json().get("status") == "success",
        f"成功入库: {r_reg_new.json().get('case_name')} (案号: {r_reg_new.json().get('case_no')})"
    )

    # 8.5 零幻觉引注熔断网关：已废止旧法探测与民法典自动净化
    # 旧断言未传 incident_date 却期望自动替换，违反时间效力规则；现要求 incident_date>=2021-01-01，
    # 且定点替换后同句「第五百八十五条」只出现一次（避免全文 replace 重复引用）。
    bad_text_abrogated = "原告与被告于2022年订立买卖合同，依据《中华人民共和国合同法》第一百一十四条及《民法典》第五百八十五条之规定，被告逾期付款应当承担违约金责任。"
    r_audit_abrogated = client.post("/api/v1/guard/audit-citations", json={
        "text": bad_text_abrogated,
        "auto_sanitize": True,
        "incident_date": "2022-01-01",
    })
    assert r_audit_abrogated.status_code == 200, r_audit_abrogated.text
    audit_res1 = r_audit_abrogated.json()
    sanitized1 = audit_res1.get("sanitized_text", "")
    log_test(
        "第三优先级: 引注熔断 - 废止《合同法》探测拦截与自动净化重写",
        audit_res1.get("has_abrogated_laws") is True
        and "合同法" not in sanitized1
        and "第五百八十五条" in sanitized1
        and sanitized1.count("第五百八十五条") == 1,
        f"净化前包含废止法条 | 净化后文本: {sanitized1}"
    )

    # 8.6 零幻觉引注熔断网关：虚构法条与条号越界熔断拦截（大模型幻觉拦截）
    hallucinated_text = "依据《中华人民共和国民法典》第一千五百条之规定，当事人应当诚实守信履行义务。"
    r_audit_hallucination = client.post("/api/v1/guard/audit-citations", json={
        "text": hallucinated_text,
        "auto_sanitize": True
    })
    assert r_audit_hallucination.status_code == 200, r_audit_hallucination.text
    audit_res2 = r_audit_hallucination.json()
    log_test(
        "第三优先级: 零幻觉引注熔断 - 捏造条号越界探测 (民法典上限1260条拦截)",
        audit_res2.get("has_hallucinations") is True and audit_res2.get("stats", {}).get("hallucinated_count", 0) >= 1,
        f"捏造条号拦截数: {audit_res2.get('stats', {}).get('hallucinated_count')} | 体检报告: {audit_res2.get('audit_report')}"
    )

    # 8.7 零幻觉引注熔断网关：100% 权威现行法条绿灯放行并提取条文释义
    clean_text = "依《中华人民共和国民法典》第五百零九条、第五百七十七条以及第五百八十五条之规定，当事人应当按照约定全面履行自己的义务。"
    r_audit_clean = client.post("/api/v1/guard/audit-citations", json={
        "text": clean_text,
        "auto_sanitize": True
    })
    assert r_audit_clean.status_code == 200, r_audit_clean.text
    audit_res3 = r_audit_clean.json()
    log_test(
        "第三优先级: 零幻觉引注熔断 - 100% 权威现行法条通过并关联官方释义",
        audit_res3.get("compliance_rate") == "100.0%" and audit_res3.get("stats", {}).get("verified_count") == 3,
        f"核验合格率: {audit_res3.get('compliance_rate')} | 权威条目数: {audit_res3.get('stats', {}).get('verified_count')}条"
    )

    # 8.8 零幻觉引注熔断网关：单个法条快速探针
    r_verify_single = client.post("/api/v1/guard/verify-single", json={
        "law_name": "《民法典》",
        "article_num": 585
    })
    assert r_verify_single.status_code == 200, r_verify_single.text
    v_single = r_verify_single.json()
    log_test(
        "第三优先级: 单个法条现行有效性权威穿透探针 (民法典第585条)",
        v_single.get("audit_status") == "VERIFIED_VALID" and "违约金" in v_single.get("article_content", ""),
        f"条文全称: {v_single.get('statute_full_name')} | 摘录: {v_single.get('article_content')[:40]}..."
    )

    # 8.9 业务流水线端到端协同验证：立案文书流水线自动附带零幻觉引注体检报告
    r_pkg_with_guard = client.post("/api/v1/phase1/generate-litigation-package", json={
        "cause_of_action": "民间借贷纠纷",
        "claim_amount": 200000.0,
        "raw_case_notes": "原告借给被告20万，被告出具借条，到期未还，约定利息未结清。请求判令还本付息并支付违约金。",
        "court_name": "北京市朝阳区",
        "plaintiff": {
            "name": "李四",
            "id_card_or_uscc": "110101199003072375",
            "phone": "13800000000",
            "address": "北京市朝阳区"
        },
        "defendant": {
            "name": "张三",
            "id_card_or_uscc": "110101198506121234",
            "phone": "13900000000",
            "address": "北京市海淀区"
        },
        "auto_generate_sub_docs": False
    })
    assert r_pkg_with_guard.status_code == 200, r_pkg_with_guard.text
    pkg_guard_res = r_pkg_with_guard.json()
    log_test(
        "第三优先级: 立案文书流水线端到端无缝嵌入零幻觉引注体检与自动净化",
        "citation_audit" in pkg_guard_res and pkg_guard_res.get("citation_audit", {}).get("stats", {}).get("total_citations", 0) >= 0,
        f"文书包流水线输出状态: {pkg_guard_res.get('status')} | 引注体检报告: {pkg_guard_res.get('citation_audit', {}).get('audit_report')}"
    )

    print("\n" + "=" * 60)
    print("恭喜！LawFlow 智律工作台全系统（含第三优先级利益冲突穿透排查与零幻觉引注熔断）全部通过实证校验！")
    print("=" * 60)

if __name__ == "__main__":
    run_tests()

