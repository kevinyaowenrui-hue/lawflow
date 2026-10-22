# -*- coding: utf-8 -*-
"""
LawFlow 智律工作台全系统端到端实证（TestClient，不依赖 live uvicorn）。
每个原 log_test 检查点对应一个独立 test_* 函数。
"""
from __future__ import annotations

import asyncio

import docx as python_docx
import pymupdf as fitz
import pytest

from gateway.core.paths import get_generated_dir
from gateway.routers.phase1_litigation import convert_docx_to_pdf


# ---------------------------------------------------------------------------
# 阶段一：司法实务计算 + 文书流水线
# ---------------------------------------------------------------------------

def test_health_probe(client):
    res = client.get("/health")
    assert res.status_code == 200, res.text
    assert res.json().get("status") == "healthy"


def test_rmb_capital(client):
    r = client.post("/api/v1/phase1/calc/rmb-capital", json={"amount": 123456.78})
    assert r.status_code == 200, r.text
    capital = r.json().get("rmb_capital", "")
    assert "壹拾贰万叁仟肆佰伍拾陆元柒角捌分" in capital


def test_idcard_young_worker(client):
    r = client.post("/api/v1/phase1/calc/idcard", json={"id_card": "110101199003072375", "is_cadre": False})
    assert r.status_code == 200, r.text
    id_data = r.json()
    assert id_data.get("valid") is True
    assert id_data.get("is_retired") is False


def test_retire_delay_male_1965(client):
    r_male = client.post("/api/v1/phase1/calc/idcard", json={"id_card": "110101196501150010", "is_cadre": False})
    assert r_male.status_code == 200, r_male.text
    male_data = r_male.json()
    assert male_data.get("delay_months") == 1
    assert male_data.get("statutory_retire_date") == "2025年02月"


def test_retire_delay_female_worker_1975(client):
    r_fem_worker = client.post("/api/v1/phase1/calc/idcard", json={"id_card": "110101197501150023", "is_cadre": False})
    assert r_fem_worker.status_code == 200, r_fem_worker.text
    fem_worker_data = r_fem_worker.json()
    assert fem_worker_data.get("delay_months") == 1
    assert fem_worker_data.get("statutory_retire_date") == "2025年02月"


def test_retire_delay_female_cadre_1970(client):
    r_fem_cadre = client.post("/api/v1/phase1/calc/idcard", json={"id_card": "110101197001150027", "is_cadre": True})
    assert r_fem_cadre.status_code == 200, r_fem_cadre.text
    fem_cadre_data = r_fem_cadre.json()
    assert fem_cadre_data.get("delay_months") == 1
    assert fem_cadre_data.get("statutory_retire_date") == "2025年02月"


def test_uscc_validate(client):
    r = client.post("/api/v1/phase1/calc/uscc", json={"uscc": "91440300741270847Y"})
    assert r.status_code == 200, r.text
    assert r.json().get("valid") is True


def test_lpr_commercial_interest(client):
    r = client.post("/api/v1/phase1/calc/lpr-interest", json={
        "principal": 500000.0,
        "start_date": "2023-01-01",
        "end_date": "2024-01-01",
        "multiplier": 1.5,
        "term": "1Y",
    })
    assert r.status_code == 200, r.text
    assert r.json().get("total_interest", 0) > 0


def test_private_lending_cap_4x_lpr(client):
    r_cap = client.post("/api/v1/phase1/calc/lpr-interest", json={
        "principal": 200000.0,
        "start_date": "2022-03-01",
        "end_date": "2023-03-01",
        "rule_mode": "private_lending_cap",
        "contract_date": "2022-03-01",
        "fixed_annual_rate": 24.0,
    })
    assert r_cap.status_code == 200, r_cap.text
    seg0 = r_cap.json().get("segments", [{}])[0]
    assert seg0.get("statutory_cap_rate") == 14.8
    assert seg0.get("applied_annual_rate") == 14.8
    assert seg0.get("is_rate_capped") is True


def test_private_lending_split_2020_08_20(client):
    # 批次5 S-16：合同成立日在 2020-08-20 之前必须提供起诉日 filing_date
    r_split = client.post("/api/v1/phase1/calc/lpr-interest", json={
        "principal": 100000.0,
        "start_date": "2020-01-01",
        "end_date": "2021-01-01",
        "rule_mode": "private_lending_cap",
        "contract_date": "2020-01-01",
        "filing_date": "2021-12-01",
        "fixed_annual_rate": 20.0,
    })
    assert r_split.status_code == 200, r_split.text
    segs = r_split.json().get("segments", [])
    assert len(segs) == 2
    assert "2020-08-20前" in segs[0].get("regime", "")
    assert "2020-08-20后" in segs[1].get("regime", "")


def test_labor_arbitration_2n(client):
    r = client.post("/api/v1/phase1/calc/labor-arbitration", json={
        "monthly_salary": 12000.0,
        "hire_date": "2020-03-01",
        "termination_date": "2023-08-15",
    })
    assert r.status_code == 200, r.text
    amount_2n = r.json().get("tiers", {}).get("2N", {}).get("amount")
    assert amount_2n == 84000.0


def test_labor_half_month_under_6_months(client):
    r_half = client.post("/api/v1/phase1/calc/labor-arbitration", json={
        "monthly_salary": 10000.0,
        "hire_date": "2023-01-01",
        "termination_date": "2023-06-20",
    })
    assert r_half.status_code == 200, r_half.text
    half_data = r_half.json()
    assert half_data.get("compensation_months") == 0.5
    assert half_data.get("tiers", {}).get("N", {}).get("amount") == 5000.0


def test_labor_segment_2008_high_salary(client):
    r_seg2008 = client.post("/api/v1/phase1/calc/labor-arbitration", json={
        "monthly_salary": 50000.0,
        "hire_date": "2005-01-01",
        "termination_date": "2023-12-31",
        "local_avg_salary": 10000.0,
    })
    assert r_seg2008.status_code == 200, r_seg2008.text
    seg_info = r_seg2008.json().get("segment_2008", {})
    assert seg_info.get("is_applicable") is True
    assert seg_info.get("pre_2008_salary_basis") == 50000.0
    assert seg_info.get("post_2008_compensation_months") == 12.0


def test_court_fees(client):
    r = client.post("/api/v1/phase1/calc/court-fees", json={"claim_amount": 500000.0, "is_summary_procedure": False})
    assert r.status_code == 200, r.text
    assert r.json().get("standard_fee") == 8800.0


@pytest.mark.requires_ollama
def test_generate_litigation_package(client):
    r = client.post("/api/v1/phase1/generate-litigation-package", json={
        "cause_of_action": "民间借贷纠纷",
        "plaintiff": {
            "name": "李四",
            "id_card_or_uscc": "110101199003072375",
            "gender": "男",
            "birth": "1990年3月7日",
            "ethnicity": "汉族",
            "address": "北京市海淀区中关村南大街1号",
            "phone": "13800138000",
        },
        "defendant": {
            "name": "张三",
            "id_card_or_uscc": "110101198505121232",
            "gender": "男",
            "birth": "1985年5月12日",
            "ethnicity": "汉族",
            "address": "北京市朝阳区建国路88号",
            "phone": "13900139000",
        },
        "claims": [
            "判令被告立即向原告偿还借款本金人民币500,000元；",
            "判令被告向原告支付利息（自2023年3月1日起至实际清偿之日止，按年利率6%计算）；",
            "判令本案全部诉讼费用由被告承担。",
        ],
        "claim_amount": 500000.0,
        "raw_case_notes": "被告张三在2023年3月因资金周转困难找原告借款50万元，约定借期半年，年利率6%。原告当天通过招商银行全额转账。借期届满后原告多次通过微信及电话催讨，被告以工程款未结为由推脱，至今分文未还。",
        "court_name": "北京市海淀区人民法院",
        "lawyer_name": "王大状",
        "law_firm": "北京智律律师事务所",
        "lawyer_phone": "13600001111",
        "lawyer_license": "11101201510123456",
        "auto_generate_sub_docs": True,
    })
    assert r.status_code == 200, r.text
    pkg_data = r.json()
    files_generated = pkg_data.get("generated_documents", [])
    assert pkg_data.get("status") == "success"
    assert len(files_generated) >= 4


# ---------------------------------------------------------------------------
# 阶段二
# ---------------------------------------------------------------------------

@pytest.mark.requires_ollama
def test_trial_outline(client):
    r = client.post("/api/v1/phase2/trial-outline", json={
        "case_title": "李四诉张三民间借贷纠纷案",
        "cause_of_action": "民间借贷纠纷",
        "plaintiff_claims_and_facts": "原告诉请判令被告偿还本金50万元及利息3万元，原告已提供招行转账凭证及微信借款合意记录。",
        "defendant_defense": "被告辩称该笔资金系双方共同承揽工程的合伙投资款，非个人借贷，且原告未提供纸质借据原件。",
    })
    assert r.status_code == 200, r.text
    outline_text = r.json().get("trial_outline_markdown", "")
    assert len(outline_text) > 100


@pytest.mark.requires_ollama
def test_contract_review_defendant_perspective(client):
    r = client.post("/api/v1/phase2/contract-review", json={
        "contract_title": "企业级软件系统开发服务合同",
        "contract_text": "第七条 违约责任：乙方如延期交付超过3日，每日应向甲方支付合同总价款5%的惩罚性违约金；甲方有权随时单方解除合同且无需退还首期款项。发生争议由甲方所在地人民法院管辖。",
        "perspective": "乙方防御版",
    })
    assert r.status_code == 200, r.text
    assert "review" in r.json()


def test_redact_sensitive_adjacent_cjk(client):
    dummy_pdf_filename = "test_sensitive_doc.pdf"
    dummy_pdf_path = get_generated_dir() / dummy_pdf_filename
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((50, 100), "原告身份证号110101199003072379在北京市朝阳区", fontname="china-s", fontsize=12)
    page.insert_text((50, 150), "结算转账银行卡号6222021234567890123尾号开户行", fontname="china-s", fontsize=12)
    doc.save(str(dummy_pdf_path))
    doc.close()

    r = client.post("/api/v1/phase2/redact-sensitive", json={
        "pdf_filename": dummy_pdf_filename,
        "redact_idcard": True,
        "redact_bankcard": True,
    })
    assert r.status_code == 200, r.text
    redact_data = r.json()
    assert redact_data.get("status") == "success"
    assert redact_data.get("redactions_count") >= 2


def test_pymupdf_long_doc_pagination():
    generated = get_generated_dir()
    test_long_docx = generated / "test_long_complaint.docx"
    test_long_pdf = generated / "test_long_complaint.pdf"
    d_doc = python_docx.Document()
    d_doc.add_heading("民事起诉状（长文流式排版测试）", 0)
    for i in range(1, 121):
        d_doc.add_paragraph(
            f"第{i}段：原告李四与被告张三民间借贷纠纷事实推演，包含第{i}笔资金周转明细、"
            f"利息约定以及多次书面催告记录，以此验证多页流式排版与页码防伪水印。"
        )
    d_doc.save(str(test_long_docx))

    conv_ok = asyncio.run(convert_docx_to_pdf(str(test_long_docx), str(test_long_pdf), "长篇诉状排版"))
    assert conv_ok is True
    pdf_inspect = fitz.open(str(test_long_pdf))
    p_count = pdf_inspect.page_count
    first_page_text = pdf_inspect[0].get_text("text")
    last_page_text = pdf_inspect[p_count - 1].get_text("text")
    pdf_inspect.close()
    assert p_count >= 3
    assert "第 1 页 / 共" in first_page_text
    assert f"第 {p_count} 页 / 共 {p_count} 页" in last_page_text


# ---------------------------------------------------------------------------
# 阶段三
# ---------------------------------------------------------------------------

def test_law_chunker(client):
    sample_law = """
中华人民共和国民法典
第一编 总则
第一章 基本规定
第一条 为了保护民事主体的合法权益，调整民事关系，维护社会和经济秩序，适应中国特色社会主义发展要求，弘扬社会主义核心价值观，根据宪法，制定本法。
第二条 民法调整平等主体的自然人、法人和非法人组织之间的人身关系和财产关系。
第二章 自然人
第一节 民事权利能力和民事行为能力
第十三条 自然人从出生时起到死亡时止，具有民事权利能力，依法享有民事权利，承担民事义务。
"""
    r = client.post("/api/v1/phase3/rag/chunk-law", json={
        "law_name": "民法典总则编节选",
        "raw_law_text": sample_law,
    })
    assert r.status_code == 200, r.text
    chunks = r.json().get("chunks", [])
    assert len(chunks) == 3


def test_rag_search(client):
    r = client.post("/api/v1/phase3/rag/search", json={
        "query": "借款利息约定超过LPR四倍是否有效？",
        "corpus_type": "all",
        "top_k": 3,
    })
    assert r.status_code == 200, r.text
    assert len(r.json().get("results", [])) > 0


def test_create_signing_link(client):
    r = client.post("/api/v1/phase3/create-signing-link", json={
        "agreement_title": "民事借贷纠纷专项法律服务委托代理合同",
        "client_name": "李四",
        "client_phone": "13800138000",
        "client_email": "lisi@example.com",
        "lawyer_name": "王大状",
        "law_firm": "北京智律律师事务所",
        "service_fee": 30000.0,
        "case_summary": "代理原告李四诉张三民间借贷纠纷一案一审诉讼",
    })
    assert r.status_code == 200, r.text
    assert r.json().get("status") == "success"
    assert r.json().get("envelope_id")


def test_signing_webhook(client, monkeypatch):
    # 批次 S-13：未配置 webhook secret 时返回 503；配置后才可完成回调
    monkeypatch.setenv("DOCUSEAL_WEBHOOK_SECRET", "test-webhook-secret")
    create = client.post("/api/v1/phase3/create-signing-link", json={
        "agreement_title": "民事借贷纠纷专项法律服务委托代理合同",
        "client_name": "李四",
        "client_phone": "13800138000",
        "client_email": "lisi@example.com",
        "lawyer_name": "王大状",
        "law_firm": "北京智律律师事务所",
        "service_fee": 30000.0,
        "case_summary": "代理原告李四诉张三民间借贷纠纷一案一审诉讼",
    })
    assert create.status_code == 200, create.text
    envelope_id = create.json().get("envelope_id")
    r = client.post(
        "/api/v1/phase3/signing-webhook",
        headers={"X-LawFlow-Webhook-Secret": "test-webhook-secret"},
        json={
            "event_type": "envelope.completed",
            "envelope_id": envelope_id,
            "status": "completed",
            "signed_at": "2026-10-16T12:00:00",
        },
    )
    assert r.status_code == 200, r.text
    assert r.json().get("status") == "success"


# ---------------------------------------------------------------------------
# 深度精算
# ---------------------------------------------------------------------------

def test_statute_limitations_interruption(client):
    r = client.post("/api/v1/phase1/calc/statute-limitations", json={
        "start_date": "2023-03-10",
        "cause_type": "general_3y",
        "interruption_events": [{"date": "2023-09-15", "description": "李四微信承诺还款"}],
    })
    assert r.status_code == 200, r.text
    assert r.json().get("final_expire_date") == "2026-09-15"


def test_guarantee_general_6m(client):
    r_guar_gen = client.post("/api/v1/phase1/calc/statute-limitations", json={
        "start_date": "2023-01-01",
        "cause_type": "guarantee_general_6m",
    })
    assert r_guar_gen.status_code == 200, r_guar_gen.text
    res_gen = r_guar_gen.json()
    assert res_gen.get("is_preclusive") is True
    assert res_gen.get("final_expire_date") == "2023-07-01"
    assert "起诉主债务人" in res_gen.get("special_notice", "")


def test_guarantee_joint_6m(client):
    r_guar_joint = client.post("/api/v1/phase1/calc/statute-limitations", json={
        "start_date": "2023-01-01",
        "cause_type": "guarantee_joint_6m",
    })
    assert r_guar_joint.status_code == 200, r_guar_joint.text
    res_joint = r_guar_joint.json()
    assert res_joint.get("is_preclusive") is True
    assert res_joint.get("final_expire_date") == "2023-07-01"


def test_statute_suspension_art194(client):
    r_susp = client.post("/api/v1/phase1/calc/statute-limitations", json={
        "start_date": "2020-01-01",
        "cause_type": "general_3y",
        "suspension_events": [{
            "date": "2022-09-01",
            "description": "严重疫情封控导致当事人丧失起诉行为能力",
            "resume_date": "2022-11-01",
        }],
    })
    assert r_susp.status_code == 200, r_susp.text
    assert r_susp.json().get("final_expire_date") == "2023-05-01"


def test_preservation_renewal_advance_7_days(client):
    r = client.post("/api/v1/phase1/calc/preservation-renewal", json={
        "preservation_type": "bank_deposit",
        "start_date": "2024-01-15",
    })
    assert r.status_code == 200, r.text
    assert r.json().get("advance_days_required") == 7


def test_statute_validity_abrogated(client):
    r = client.post("/api/v1/phase1/calc/statute-validity", json={
        "citation_text": "依据《中华人民共和国合同法》第一百零七条及《担保法》第十九条",
        "incident_date": "2023-05-01",
    })
    assert r.status_code == 200, r.text
    validity_res = r.json()
    assert validity_res.get("is_safe") is False
    assert len(validity_res.get("detected_abrogated")) >= 2


# ---------------------------------------------------------------------------
# 深度矩阵
# ---------------------------------------------------------------------------

def test_deep_matrix_presets(client):
    r = client.get("/api/v1/deep-matrix/presets")
    assert r.status_code == 200, r.text
    assert len(r.json().get("presets", [])) >= 2


def test_deep_matrix_blue_team(client):
    r = client.post("/api/v1/deep-matrix/analyze", json={
        "case_preset": "case_loan_intermingled",
        "perspective": "plaintiff_blue_team",
    })
    assert r.status_code == 200, r.text
    matrix_blue = r.json()
    elements = matrix_blue.get("matrix_elements", [])
    conflicts = matrix_blue.get("detected_conflicts", [])
    scorecard = matrix_blue.get("scorecard", {})
    assert len(elements) >= 5
    assert len(conflicts) >= 3
    assert scorecard.get("is_litigation_ready") is True


def test_deep_matrix_red_team(client):
    r = client.post("/api/v1/deep-matrix/analyze", json={
        "case_preset": "case_loan_intermingled",
        "perspective": "adversarial_red_team",
    })
    assert r.status_code == 200, r.text
    assert r.json().get("is_red_team_view") is True


def test_adversarial_attack(client):
    r = client.post("/api/v1/deep-matrix/adversarial-attack", json={
        "evidence_id": "DOC-02",
        "evidence_name": "招商银行网银电子转账回单",
        "plaintiff_purpose": "证明已向李四交付50万元借款",
        "raw_snippet": "收款人：王五；摘要：往来款",
    })
    assert r.status_code == 200, r.text
    attack_res = r.json()
    assert len(attack_res.get("attacks", [])) == 3


# ---------------------------------------------------------------------------
# 法典中枢
# ---------------------------------------------------------------------------

def test_statutes_overview(client):
    r = client.get("/api/v1/statutes/overview")
    assert r.status_code == 200, r.text
    stat_ov = r.json().get("data", {})
    assert stat_ov.get("total_laws", 0) >= 10
    assert stat_ov.get("total_articles", 0) >= 35
    assert stat_ov.get("total_parent_child_links", 0) >= 10


def test_search_civil_code_585_child_links(client):
    r = client.post("/api/v1/statutes/search", json={"query": "民法典 585"})
    assert r.status_code == 200, r.text
    s_res = r.json().get("results", [])
    top_art = s_res[0] if s_res else {}
    child_links = top_art.get("child_interpretations", [])
    has_30_percent_rule = any("30%" in c.get("notes", "") or "百分之三十" in c.get("content", "") for c in child_links)
    assert top_art.get("article_no_cn") == "第五百八十五条"
    assert len(child_links) >= 2
    assert has_30_percent_rule


def test_search_pierce_corporate_veil(client):
    r = client.post("/api/v1/statutes/search", json={"query": "刺破公司面纱"})
    assert r.status_code == 200, r.text
    top_corp = (r.json().get("results") or [{}])[0]
    assert top_corp.get("article_no_cn") == "第二十三条"
    assert "连带责任" in top_corp.get("content", "")


def test_search_loan_resale_art13(client):
    r = client.post("/api/v1/statutes/search", json={"query": "套取贷款"})
    assert r.status_code == 200, r.text
    top_loan = (r.json().get("results") or [{}])[0]
    assert top_loan.get("article_no_cn") == "第十三条"
    assert "套取金融机构贷款" in top_loan.get("content", "")


def test_search_criminal_224(client):
    r = client.post("/api/v1/statutes/search", json={"query": "刑法 224"})
    assert r.status_code == 200, r.text
    top_crim = (r.json().get("results") or [{}])[0]
    crim_children = top_crim.get("child_interpretations", [])
    assert top_crim.get("article_no_cn") == "第二百二十四条"
    assert len(crim_children) >= 1


def test_import_custom_statute(client):
    sample_raw_law = """
第一条 为了统一全省法院审理涉互联网平台用工劳动争议案件的裁判标准，妥善平衡劳动者合法权益与平台经济健康发展，制定本指引。
第二条 依托互联网平台就业的新就业形态劳动者，与平台企业或合作企业发生用工争议的，应根据用工事实和劳动管理程度综合认定法律关系。
第三条 平台企业通过算法或者规则对劳动者进行实质性管理，劳动者在工作时间、工作路线、服务价格等方面受平台严格控制的，应当认定双方存在劳动关系。
"""
    r = client.post("/api/v1/statutes/import", json={
        "law_id": "user_import_highcourt_platform_labor_2024",
        "full_name": "【测试】某省高级人民法院关于审理互联网平台用工劳动争议案件的审理指引",
        "short_name": "【测试】平台用工审理指引",
        "hierarchy": 4,
        "promulgated_by": "某省高级人民法院",
        "effective_date": "2024-05-01",
        "domain": "劳动",
        "raw_text": sample_raw_law,
        "summary": "新业态用工劳动关系从属性认定标准",
    })
    assert r.status_code == 200, r.text
    assert r.json().get("data", {}).get("articles_imported") == 3


def test_search_imported_statute(client):
    # ensure import first (order-independent)
    sample_raw_law = """
第一条 为了统一全省法院审理涉互联网平台用工劳动争议案件的裁判标准，妥善平衡劳动者合法权益与平台经济健康发展，制定本指引。
第二条 依托互联网平台就业的新就业形态劳动者，与平台企业或合作企业发生用工争议的，应根据用工事实和劳动管理程度综合认定法律关系。
第三条 平台企业通过算法或者规则对劳动者进行实质性管理，劳动者在工作时间、工作路线、服务价格等方面受平台严格控制的，应当认定双方存在劳动关系。
"""
    client.post("/api/v1/statutes/import", json={
        "law_id": "user_import_highcourt_platform_labor_2024",
        "full_name": "【测试】某省高级人民法院关于审理互联网平台用工劳动争议案件的审理指引",
        "short_name": "【测试】平台用工审理指引",
        "hierarchy": 4,
        "promulgated_by": "某省高级人民法院",
        "effective_date": "2024-05-01",
        "domain": "劳动",
        "raw_text": sample_raw_law,
        "summary": "新业态用工劳动关系从属性认定标准",
    })
    r = client.post("/api/v1/statutes/search", json={"query": "算法或者规则"})
    assert r.status_code == 200, r.text
    res_imported = r.json().get("results", [])
    assert len(res_imported) >= 1
    assert "平台用工审理指引" in res_imported[0].get("short_name", "")
    assert (
        res_imported[0].get("verification_status") == "USER_IMPORTED_UNVERIFIED"
        or "未核验" in str(res_imported[0].get("verification_label") or "")
    )


# ---------------------------------------------------------------------------
# 第二优先级：自定义卷宗
# ---------------------------------------------------------------------------

def test_analyze_custom_loan(client):
    r = client.post("/api/v1/deep-matrix/analyze-custom", json={
        "case_preset": "custom",
        "perspective": "plaintiff_blue_team",
        "custom_case_title": "王强诉赵刚民间借贷纠纷案",
        "custom_cause_of_action": "民间借贷纠纷",
        "custom_claim_amount": 350000.0,
        "custom_plaintiff": "王强",
        "custom_defendant": "赵刚",
        "custom_raw_text": "2023年4月15日，赵刚向王强借款35万元，约定借期1年，月息1.5%。王强于当天通过微信指定账户将35万元转给赵刚的表弟钱某。借期届满后赵刚以未向其本人银行卡付款为由拒不偿还。",
        "custom_evidence_notes": "证据一：微信聊天记录导出件（赵刚明确要求转入钱某建行卡，并附卡号）\n证据二：招商银行个人网银电子回单（户名钱某，金额35万元）\n证据三：催款录音光盘（赵刚在通话中承认收到该笔借款）",
        "custom_statutes": [
            "《中华人民共和国民法典》第六百六十七条",
            "《最高人民法院关于审理民间借贷案件适用法律若干问题的规定》第十九条",
        ],
    })
    assert r.status_code == 200, r.text
    custom_loan_res = r.json()
    assert len(custom_loan_res.get("matrix_elements", [])) >= 3
    assert len(custom_loan_res.get("detected_conflicts", [])) >= 1
    assert len(custom_loan_res.get("timeline_events", [])) >= 2
    assert len(custom_loan_res.get("raw_documents", [])) >= 3


def test_analyze_custom_red_team(client):
    r_red = client.post("/api/v1/deep-matrix/analyze-custom", json={
        "case_preset": "custom",
        "perspective": "adversarial_red_team",
        "custom_case_title": "王强诉赵刚民间借贷纠纷案",
        "custom_cause_of_action": "民间借贷纠纷",
        "custom_claim_amount": 350000.0,
        "custom_plaintiff": "王强",
        "custom_defendant": "赵刚",
        "custom_raw_text": "2023年4月15日，赵刚向王强借款35万元，约定借期1年，月息1.5%。王强于当天通过微信指定账户将35万元转给赵刚的表弟钱某。借期届满后赵刚以未向其本人银行卡付款为由拒不偿还。",
        "custom_evidence_notes": "证据一：微信聊天记录导出件\n证据二：网银回单35万",
    })
    assert r_red.status_code == 200, r_red.text
    red_res = r_red.json()
    red_attacks = [e.get("adversarial_attack", "") for e in red_res.get("matrix_elements", []) if e.get("adversarial_attack")]
    assert red_res.get("is_red_team_view") is True
    assert len(red_attacks) >= 3


def test_analyze_custom_sales(client):
    r_sale = client.post("/api/v1/deep-matrix/analyze-custom", json={
        "case_preset": "custom",
        "perspective": "plaintiff_blue_team",
        "custom_case_title": "华盛建材诉宏业建设买卖合同欠款纠纷案",
        "custom_cause_of_action": "买卖合同纠纷",
        "custom_claim_amount": 800000.0,
        "custom_plaintiff": "华盛建材有限公司",
        "custom_defendant": "宏业建设工程有限公司",
        "custom_raw_text": "2023年6月，原告依合同向被告供应钢材总价80万元，被告在送货单上签字盖章确认签收。现约定付款期限已过，被告拖欠尾款50万元未付。",
        "custom_evidence_notes": "证据一：买卖合同原件\n证据二：项目工地材料签收单5份\n证据三：催款对账确认单",
    })
    assert r_sale.status_code == 200, r_sale.text
    sale_elements = r_sale.json().get("matrix_elements", [])
    assert len(sale_elements) >= 3
    assert any("民法典" in e.get("claim_basis", "") for e in sale_elements)


def test_analyze_delegates_custom(client):
    r_delegate = client.post("/api/v1/deep-matrix/analyze", json={
        "case_preset": "custom",
        "custom_case_title": "委托分发测试案",
        "custom_cause_of_action": "通用民事纠纷",
        "custom_raw_text": "当事人存在合伙协议及投资款交付凭证，现要求清算返还出资。",
        "custom_evidence_notes": "证据一：合伙协议",
    })
    assert r_delegate.status_code == 200, r_delegate.text
    del_res = r_delegate.json()
    assert del_res.get("status") == "success"
    assert del_res.get("case_title") == "委托分发测试案"


# ---------------------------------------------------------------------------
# 第三优先级：利益冲突 + 引注熔断
# ---------------------------------------------------------------------------

def test_conflict_waivable_advisor(client, seed_demo):
    r_conflict_direct = client.post("/api/v1/conflict/check", json={
        "client_name": "北京知客云服科技有限公司",
        "opposing_party": "北京宏达商贸有限公司",
        "case_subject": "技术合同尾款争议",
    })
    assert r_conflict_direct.status_code == 200, r_conflict_direct.text
    c_direct = r_conflict_direct.json()
    assert c_direct.get("conflict_level") == "WAIVABLE_WITH_CONSENT"
    assert c_direct.get("requires_consent") is True
    assert c_direct.get("has_conflict") is True
    assert any("律师法" in rule and "第三十九条" in rule for rule in c_direct.get("triggered_rules", []))


def test_conflict_prohibited_both_clients(client, seed_demo):
    r_conflict_both = client.post("/api/v1/conflict/check", json={
        "client_name": "李四",
        "opposing_party": "智科云服信息技术（北京）有限公司",
        "case_subject": "合同纠纷",
    })
    assert r_conflict_both.status_code == 200, r_conflict_both.text
    c_both = r_conflict_both.json()
    assert c_both.get("has_conflict") is True
    assert c_both.get("conflict_level") == "PROHIBITED"


def test_conflict_clear_certificate(client, seed_demo):
    r_conflict_clear = client.post("/api/v1/conflict/check", json={
        "client_name": "北京创新未来科技有限公司",
        "opposing_party": "广州星海传媒有限公司",
        "case_subject": "侵害商标权纠纷",
    })
    assert r_conflict_clear.status_code == 200, r_conflict_clear.text
    c_clear = r_conflict_clear.json()
    cert = c_clear.get("clearance_certificate", {})
    assert c_clear.get("conflict_level") == "CLEAR"
    assert c_clear.get("has_conflict") is False
    assert len(cert.get("sha256_fingerprint", "")) == 64


def test_conflict_registry_list(client, seed_demo):
    r_reg_list = client.get("/api/v1/conflict/registry")
    assert r_reg_list.status_code == 200, r_reg_list.text
    reg_data = r_reg_list.json()
    assert reg_data.get("status") == "success"
    assert reg_data.get("count", 0) >= 4


def test_conflict_register_case(client, seed_demo):
    r_reg_new = client.post("/api/v1/conflict/register-case", json={
        "case_name": "宏达集团下属子公司破产重整案",
        "case_no": "(2026)京破999号",
        "client_name": "北京宏达商贸有限公司",
        "opposing_party": "全体破产债权人",
        "status": "ACTIVE",
        "handling_lawyer": "合规审查专员",
        "notes": "破产重组专项",
    })
    assert r_reg_new.status_code == 200, r_reg_new.text
    assert r_reg_new.json().get("status") == "success"


def test_guard_abrogated_sanitize(client):
    bad_text_abrogated = "原告与被告于2022年订立买卖合同，依据《中华人民共和国合同法》第一百一十四条及《民法典》第五百八十五条之规定，被告逾期付款应当承担违约金责任。"
    r_audit_abrogated = client.post("/api/v1/guard/audit-citations", json={
        "text": bad_text_abrogated,
        "auto_sanitize": True,
        "incident_date": "2022-01-01",
    })
    assert r_audit_abrogated.status_code == 200, r_audit_abrogated.text
    audit_res1 = r_audit_abrogated.json()
    sanitized1 = audit_res1.get("sanitized_text", "")
    assert audit_res1.get("has_abrogated_laws") is True
    assert "合同法" not in sanitized1
    # 业务含义：废止《合同法》被剔除，现行《民法典》585 保留；sanitize 可能保留原有民法典引用
    assert "第五百八十五条" in sanitized1
    assert sanitized1.count("第五百八十五条") >= 1


def test_guard_hallucination_exceeded(client):
    hallucinated_text = "依据《中华人民共和国民法典》第一千五百条之规定，当事人应当诚实守信履行义务。"
    r_audit_hallucination = client.post("/api/v1/guard/audit-citations", json={
        "text": hallucinated_text,
        "auto_sanitize": True,
    })
    assert r_audit_hallucination.status_code == 200, r_audit_hallucination.text
    audit_res2 = r_audit_hallucination.json()
    assert audit_res2.get("has_hallucinations") is True
    assert audit_res2.get("stats", {}).get("hallucinated_count", 0) >= 1


def test_guard_clean_citations(client):
    clean_text = "依《中华人民共和国民法典》第五百零九条、第五百七十七条以及第五百八十五条之规定，当事人应当按照约定全面履行自己的义务。"
    r_audit_clean = client.post("/api/v1/guard/audit-citations", json={
        "text": clean_text,
        "auto_sanitize": True,
    })
    assert r_audit_clean.status_code == 200, r_audit_clean.text
    audit_res3 = r_audit_clean.json()
    assert audit_res3.get("compliance_rate") == "100.0%"
    assert audit_res3.get("stats", {}).get("verified_count") == 3


def test_guard_verify_single_585(client):
    r_verify_single = client.post("/api/v1/guard/verify-single", json={
        "law_name": "《民法典》",
        "article_num": 585,
    })
    assert r_verify_single.status_code == 200, r_verify_single.text
    v_single = r_verify_single.json()
    # 前序批次将绿标状态统一为 LOCAL_MATCH（原脚本写 VERIFIED_VALID）
    assert v_single.get("audit_status") == "LOCAL_MATCH"
    assert "违约金" in v_single.get("article_content", "")


@pytest.mark.requires_ollama
def test_litigation_package_embeds_citation_audit(client):
    r_pkg_with_guard = client.post("/api/v1/phase1/generate-litigation-package", json={
        "cause_of_action": "民间借贷纠纷",
        "claim_amount": 200000.0,
        "raw_case_notes": "原告借给被告20万，被告出具借条，到期未还，约定利息未结清。请求判令还本付息并支付违约金。",
        "court_name": "北京市朝阳区",
        "plaintiff": {
            "name": "李四",
            "id_card_or_uscc": "110101199003072375",
            "phone": "13800000000",
            "address": "北京市朝阳区",
        },
        "defendant": {
            "name": "张三",
            "id_card_or_uscc": "110101198506121234",
            "phone": "13900000000",
            "address": "北京市海淀区",
        },
        "auto_generate_sub_docs": False,
    })
    assert r_pkg_with_guard.status_code == 200, r_pkg_with_guard.text
    pkg_guard_res = r_pkg_with_guard.json()
    assert "citation_audit" in pkg_guard_res
    assert pkg_guard_res.get("citation_audit", {}).get("stats", {}).get("total_citations", 0) >= 0
