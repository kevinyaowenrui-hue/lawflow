# -*- coding: utf-8 -*-
"""批次7：立案文书包 S-09 / S-10 / S-11 验收测试（无模型模式）。"""
from __future__ import annotations

import os
import re
import tempfile
from pathlib import Path

import docx
import pytest


def _docx_full_text(path: str) -> str:
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _pdf_full_text(path: str) -> str:
    import pymupdf as fitz
    doc = fitz.open(path)
    texts = []
    for page in doc:
        texts.append(page.get_text())
    doc.close()
    return "\n".join(texts)


def _strip_ws(s: str) -> str:
    return re.sub(r"\s+", "", s or "")


@pytest.fixture
def no_ollama(monkeypatch):
    """强制专家规则兜底（无模型）。"""
    from gateway.core.ollama_client import ollama_client

    async def _offline(*args, **kwargs):
        return False

    monkeypatch.setattr(ollama_client, "_is_ollama_online", _offline)
    monkeypatch.setattr(ollama_client, "settings", {**ollama_client.settings, "provider": "ollama", "api_key": ""})
    return ollama_client


def _base_party_plaintiff():
    return {
        "name": "测试原告甲",
        "id_card_or_uscc": "110101199003072340",
        "address": "北京市朝阳区测试路1号",
        "phone": "13800000001",
    }


def _base_party_defendant_legal():
    return {
        "name": "测试被告乙有限公司",
        "id_card_or_uscc": "91350100M000100Y43",
        "address": "福州市鼓楼区测试路2号",
        "phone": "13800000002",
        "party_type": "legal",
        "legal_rep": "王某某",
        "rep_position": "执行董事",
    }


def test_normalize_court_name():
    from gateway.utils.text_utils import normalize_court_name

    assert normalize_court_name("北京市朝阳区") == "北京市朝阳区人民法院"
    assert normalize_court_name("北京市朝阳区人民法院") == "北京市朝阳区人民法院"
    assert normalize_court_name("  上海金融法院 ") == "上海金融法院"


def test_S09_sales_no_labor_and_claims_verbatim(client, no_ollama):
    payload = {
        "cause_of_action": "买卖合同纠纷",
        "plaintiff": _base_party_plaintiff(),
        "defendant": _base_party_defendant_legal(),
        "claims": ["判令被告支付货款 123,456 元"],
        "claim_amount": 123456.0,
        "raw_case_notes": "2024年1月1日原告向被告供应货物，被告拖欠货款123456元。",
        "court_name": "北京市朝阳区人民法院",
        "auto_generate_sub_docs": False,
    }
    resp = client.post("/api/v1/phase1/generate-litigation-package", json=payload)
    assert resp.status_code == 200, resp.text
    data = resp.json()
    path = data.get("main_docx_path")
    assert path and os.path.isfile(path)
    full = _docx_full_text(path)
    assert "劳动" not in full
    assert "判令被告支付货款 123,456 元" in full
    assert "暂计" not in full
    assert "人民法院人民法院" not in full


def test_S09_loan_agreed_rate_12pct(client, no_ollama):
    payload = {
        "cause_of_action": "民间借贷纠纷",
        "plaintiff": _base_party_plaintiff(),
        "defendant": {
            "name": "测试被告丙",
            "id_card_or_uscc": "110101199003072375",
            "address": "北京市朝阳区测试路3号",
            "phone": "13800000003",
        },
        "claims": [
            "判令被告偿还借款本金人民币 100000 元；",
            "判令被告支付利息；",
        ],
        "claim_amount": 100000.0,
        "raw_case_notes": "2023年12月1日被告向原告借款10万元，约定年利率12%，合同成立于2023-12-01，至今未还。",
        "court_name": "北京市朝阳区",
        "contract_date": "2023-12-01",
        "agreed_annual_rate": 12.0,
        "auto_generate_sub_docs": False,
    }
    resp = client.post("/api/v1/phase1/generate-litigation-package", json=payload)
    assert resp.status_code == 200, resp.text
    full = _docx_full_text(resp.json()["main_docx_path"])
    assert "12%" in full
    assert "14.6" not in full
    assert "人民法院人民法院" not in full
    assert "北京市朝阳区人民法院" in full


def test_S09_labor_arbitration_wording(client, no_ollama):
    payload = {
        "cause_of_action": "劳动人事争议",
        "plaintiff": _base_party_plaintiff(),
        "defendant": _base_party_defendant_legal(),
        "claims": ["裁决被申请人支付违法解除劳动合同赔偿金"],
        "claim_amount": 50000.0,
        "raw_case_notes": "用人单位违法解除劳动合同，申请人请求赔偿。证据及来源需律师核实。",
        "court_name": "朝阳区",
        "arbitration_committee": "朝阳区",
        "auto_generate_sub_docs": False,
    }
    resp = client.post("/api/v1/phase1/generate-litigation-package", json=payload)
    assert resp.status_code == 200, resp.text
    full = _docx_full_text(resp.json()["main_docx_path"])
    assert "仲裁委员会" in full
    assert "裁决被申请人" in full
    assert "人民法院" not in full
    assert "判令" not in full


def test_S09_missing_legal_rep_placeholder(client, no_ollama):
    payload = {
        "cause_of_action": "买卖合同纠纷",
        "plaintiff": _base_party_plaintiff(),
        "defendant": {
            "name": "某公司",
            "id_card_or_uscc": "91350100M000100Y43",
            "address": "福州",
            "phone": "13800000002",
            "party_type": "legal",
            # 不传法定代表人
        },
        "claims": ["判令被告支付货款 1 元"],
        "claim_amount": 1.0,
        "raw_case_notes": "欠货款。",
        "court_name": "福州市鼓楼区人民法院",
        "auto_generate_sub_docs": False,
    }
    resp = client.post("/api/v1/phase1/generate-litigation-package", json=payload)
    assert resp.status_code == 200, resp.text
    data = resp.json()
    full = _docx_full_text(data["main_docx_path"])
    assert "【待填写】" in full
    assert "法定代表人：法定代表人" not in full
    assert "missing_fields" in data
    assert any("legal_rep" in f for f in data["missing_fields"])


def test_S10_natural_and_legal_party_blocks(client, no_ollama):
    payload = {
        "cause_of_action": "民间借贷纠纷",
        "plaintiff": {
            "name": "女原告",
            "id_card_or_uscc": "110101199003072340",
            "address": "北京",
            "phone": "13800000001",
        },
        "defendant": {
            "name": "法人被告公司",
            "id_card_or_uscc": "91350100M000100Y43",
            "address": "福州",
            "phone": "13800000002",
            "party_type": "legal",
            "legal_rep": "李总",
            "rep_position": "董事长",
        },
        "claims": ["判令被告偿还借款本金人民币 100000 元；"],
        "claim_amount": 100000.0,
        "raw_case_notes": "2023年12月1日借款，年利率12%。",
        "court_name": "北京市朝阳区人民法院",
        "auto_generate_sub_docs": False,
    }
    resp = client.post("/api/v1/phase1/generate-litigation-package", json=payload)
    assert resp.status_code == 200, resp.text
    full = _docx_full_text(resp.json()["main_docx_path"])
    assert "女" in full
    assert ("1990年3月7日" in full) or ("1990年03月07日" in full)
    assert "男，" not in full
    # 被告段
    assert "统一社会信用代码" in full
    # 被告不应被标为公民身份号码且带性别出生——检查被告段附近
    # 至少全文对法人被告不应出现「被告：…性别」类错误；用宽松断言
    assert "公民身份号码：91350100M000100Y43" not in full


def test_S10_gender_conflict_422(client, no_ollama):
    payload = {
        "cause_of_action": "民间借贷纠纷",
        "plaintiff": {
            "name": "冲突原告",
            "id_card_or_uscc": "110101199003072340",  # 女
            "gender": "男",
            "address": "北京",
            "phone": "13800000001",
        },
        "defendant": {
            "name": "被告",
            "id_card_or_uscc": "110101199003072375",
            "address": "北京",
            "phone": "13800000002",
        },
        "claims": ["判令被告还款"],
        "claim_amount": 1.0,
        "raw_case_notes": "借款。",
        "court_name": "北京市朝阳区人民法院",
        "auto_generate_sub_docs": False,
    }
    resp = client.post("/api/v1/phase1/generate-litigation-package", json=payload)
    assert resp.status_code == 422, resp.text


def test_S11_pdf_coverage_and_first_page_plaintiff(client, no_ollama, tmp_path):
    payload = {
        "cause_of_action": "买卖合同纠纷",
        "plaintiff": _base_party_plaintiff(),
        "defendant": _base_party_defendant_legal(),
        "claims": ["判令被告支付货款 123,456 元"],
        "claim_amount": 123456.0,
        "raw_case_notes": "欠货款。" + ("详细事实。" * 50),
        "court_name": "北京市朝阳区人民法院",
        "auto_generate_sub_docs": False,
    }
    resp = client.post("/api/v1/phase1/generate-litigation-package", json=payload)
    assert resp.status_code == 200, resp.text
    data = resp.json()
    docx_path = data["main_docx_path"]
    pdf_path = data.get("main_pdf_path")
    assert pdf_path and os.path.isfile(pdf_path), "PDF 未生成"
    docx_chars = len(_strip_ws(_docx_full_text(docx_path)))
    pdf_chars = len(_strip_ws(_pdf_full_text(pdf_path)))
    assert pdf_chars >= docx_chars * 0.98, f"PDF覆盖率不足: pdf={pdf_chars} docx={docx_chars}"
    import pymupdf as fitz
    doc = fitz.open(pdf_path)
    page0 = doc[0].get_text()
    doc.close()
    assert "测试原告甲" in page0


def test_S11_long_docx_pdf_multipage(tmp_path):
    """构造 5000 字长文 docx → PDF 覆盖率 ≥98% 且页数 >1。"""
    import asyncio
    from gateway.routers.phase1_litigation import convert_docx_to_pdf

    long_text = "甲" * 5000
    docx_path = tmp_path / "long.docx"
    pdf_path = tmp_path / "long.pdf"
    d = docx.Document()
    d.add_paragraph("原告：长文原告")
    # 分段写入
    for i in range(0, len(long_text), 200):
        d.add_paragraph(long_text[i : i + 200])
    d.save(str(docx_path))

    ok = asyncio.run(convert_docx_to_pdf(str(docx_path), str(pdf_path), title="长文", text_summary=long_text))
    assert ok and pdf_path.is_file()
    docx_chars = len(_strip_ws(_docx_full_text(str(docx_path))))
    pdf_chars = len(_strip_ws(_pdf_full_text(str(pdf_path))))
    assert pdf_chars >= docx_chars * 0.98, f"长文覆盖率不足: pdf={pdf_chars} docx={docx_chars}"
    import pymupdf as fitz
    doc = fitz.open(str(pdf_path))
    pages = doc.page_count
    doc.close()
    assert pages > 1


def test_S11_no_electronic_litigation_slogan():
    import subprocess
    from gateway.core.paths import get_base_dir
    root = get_base_dir() / "gateway"
    # 用 Python 扫描避免依赖 rg 在 Windows 路径
    hits = []
    for p in root.rglob("*.py"):
        try:
            text = p.read_text(encoding="utf-8")
        except Exception:
            continue
        if "电子诉讼标准" in text:
            hits.append(str(p))
    assert hits == []
