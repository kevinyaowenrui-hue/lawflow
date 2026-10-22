# -*- coding: utf-8 -*-
"""
LawFlow 智律工作台 - P0 级升级实证：模糊容错原位修订与全英文涉外采购审查。
"""
from __future__ import annotations

import os
import tempfile

import docx
import pytest

from gateway.utils.redline_builder import (
    apply_in_place_docx_redline,
    find_best_matching_span,
    normalize_text_for_matching,
)
from gateway.utils.legal_prompts import is_predominantly_english


def test_fuzzy_matching_and_non_destructive_redline():
    raw_dirty = "Section 24.3\u00a0Cross-Default:\t“Any default by Buyer or its Affiliates—”  shall apply."
    norm = normalize_text_for_matching(raw_dirty)
    assert "\u00a0" not in norm
    assert "“" not in norm
    assert '"' in norm

    full_para = (
        "Section 24.2 Remedies. The Buyer may exercise all statutory remedies under applicable law. "
        "Section 24.3 Cross-Default. In the event Buyer or any of its Affiliates defaults under any other agreement, "
        "Seller shall have the right to immediately terminate this Agreement without notice. "
        "Section 24.4 Governing Law. This Agreement shall be governed by Delaware law."
    )
    target_clause = (
        "Section 24.3 Cross-Default. In the event Buyer or any of its Affiliates defaults under any other agreement, "
        "Seller shall have the right to immediately terminate this Agreement without notice."
    )
    span = find_best_matching_span(full_para, target_clause)
    assert span is not None
    start, end = span
    matched_sub = full_para[start:end]
    assert "Cross-Default" in matched_sub

    with tempfile.TemporaryDirectory() as tmpdir:
        doc_path = os.path.join(tmpdir, "client_master_agreement.docx")
        d = docx.Document()
        d.add_paragraph(full_para)
        d.save(doc_path)

        out_path = os.path.join(tmpdir, "client_master_agreement_redlined.docx")
        clause_reviews = [
            {
                "clause_title": "Cross-Default Deletion",
                "original_clause": (
                    "Section 24.3  Cross-Default. In the event Buyer or any of its Affiliates defaults under any other agreement, "
                    "Seller shall have the right to immediately terminate this Agreement without notice."
                ),
                "suggested_clause": (
                    "Section 24.3 [Intentionally Deleted - Cross-default provisions are hereby deleted in their entirety]."
                ),
            }
        ]

        res = apply_in_place_docx_redline(doc_path, clause_reviews, out_path)
        assert os.path.exists(out_path)
        assert res.stats["injected_clauses"] == 1
        assert len(res.stats["unmatched_clauses"]) == 0

        d_out = docx.Document(out_path)
        p_text_after = d_out.paragraphs[0].text
        assert "Section 24.2 Remedies." in p_text_after
        assert "Section 24.4 Governing Law." in p_text_after
        xml = d_out._body._element.xml
        assert "w:del" in xml and "w:ins" in xml


@pytest.mark.requires_ollama
def test_english_contract_review_and_playbook(client):
    en_contract_text = """MASTER PURCHASE AGREEMENT
Section 18. Liquidated Damages. In the event Seller delays delivery for more than 3 days, Seller shall pay liquidated damages equal to 30% of the total contract value and shall compensate Buyer for all indirect, special or consequential damages.
Section 24.3 Cross-Default. In the event Buyer or any of its Affiliates defaults under any other agreement, Seller shall have the right to immediately terminate this Agreement without notice."""

    assert is_predominantly_english(en_contract_text) is True

    resp = client.post(
        "/api/v1/phase2/contract-review",
        json={
            "contract_title": "Cross-Border Master Purchase Agreement",
            "contract_text": en_contract_text,
            "perspective": "Seller-Protective 乙方防御版",
        },
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "success"
    assert data["playbook_name"] in [
        "涉外与跨国商事采购红线库 (International Master Agreement Playbook)",
        "Custom Playbook",
    ] or "涉外" in data["playbook_name"]

    review = data["review"]
    clause_reviews = review.get("clause_reviews", [])
    assert len(clause_reviews) >= 1
    suggested = clause_reviews[0].get("suggested_clause", "")
    assert len(suggested) > 10
    assert any(
        w in suggested.lower()
        for w in ["deleted", "liquidated", "liability", "consequential", "exceed", "neither party"]
    ), f"建议条款必须是规范英文: {suggested}"


def test_export_redline_with_source_docx(client):
    with tempfile.TemporaryDirectory() as tmpdir:
        client_docx = os.path.join(tmpdir, "foreign_client_draft.docx")
        d = docx.Document()
        d.add_heading("Master Purchase Agreement", level=1)
        d.add_paragraph(
            "Section 24.3 Cross-Default. In the event Buyer or any of its Affiliates defaults under any other agreement, "
            "Seller shall have the right to immediately terminate this Agreement without notice."
        )
        d.save(client_docx)

        with open(client_docx, "rb") as f_up:
            up_resp = client.post(
                "/api/v1/phase2/contract-review/upload-source",
                files={
                    "file": (
                        "foreign_client_draft.docx",
                        f_up,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
        assert up_resp.status_code == 200, up_resp.text
        source_file_id = up_resp.json()["source_file_id"]

        export_resp = client.post(
            "/api/v1/phase2/contract-review/export-redline",
            json={
                "contract_title": "Master Purchase Agreement",
                "contract_text": "Section 24.3 Cross-Default...",
                "perspective": "Seller-Protective",
                "source_file_id": source_file_id,
                "clause_reviews": [
                    {
                        "clause_title": "Cross-Default",
                        "original_clause": (
                            "Section 24.3 Cross-Default. In the event Buyer or any of its Affiliates defaults under any other agreement, "
                            "Seller shall have the right to immediately terminate this Agreement without notice."
                        ),
                        "suggested_clause": (
                            "Section 24.3 [Intentionally Deleted - Cross-default provisions are hereby deleted in their entirety]."
                        ),
                    }
                ],
            },
        )

        assert export_resp.status_code == 200, export_resp.text
        exp_data = export_resp.json()
        assert exp_data["status"] == "success"
        assert exp_data["in_place_mode"] is True
        assert exp_data["stats"]["injected_clauses"] == 1
