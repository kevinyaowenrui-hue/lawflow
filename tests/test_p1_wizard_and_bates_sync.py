"""
LawFlow P1 Improvements Verification Test Suite
Verifies:
1. Bates Stamping absolute physical page map tracking & continuity (R_BATES_SYNC)
2. Evidence Catalog Word generation with synced Bates absolute page ranges (R_CATALOG_SYNC)
3. Citation Guard substantive law time anchor: historical law validity for pre-2021 events (R_TIME_ANCHOR)
4. Pipeline Stepper UI integration & A11Y compliance (R_STEPPER_UI)
"""

import io
import os
import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient

from gateway.main import app
from gateway.core.paths import get_generated_dir
from gateway.utils.citation_guard import audit_citations_in_text

AUTH_TOKEN = os.environ.get("LAWFLOW_API_TOKEN", "test-token")
AUTH_HEADERS = {"X-LawFlow-Token": AUTH_TOKEN}


@pytest.fixture
def client(monkeypatch):
    from gateway.core.ollama_client import ollama_client
    async def _offline():
        return False
    monkeypatch.setattr(ollama_client, "_is_ollama_online", _offline)
    return TestClient(app)


# ==============================================================================
# 1. Bates Stamping file_page_map Tracking & Continuity
# ==============================================================================

def _make_dummy_pdf(page_count: int, title: str = "Test Doc") -> bytes:
    doc = fitz.open()
    for p in range(page_count):
        page = doc.new_page(width=595, height=842)
        page.insert_text(fitz.Point(50, 100), f"{title} - Page {p + 1}")
    data = doc.tobytes()
    doc.close()
    return data


def _make_dummy_png() -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    img = Image.new("RGB", (300, 300), color=(240, 240, 240))
    img.save(buf, format="PNG")
    return buf.getvalue()


def test_bates_stamping_returns_file_page_map(client):
    """验证多文件合卷时准确返回每份材料的起止物理绝对页码，且页码连续无断层"""
    pdf1_data = _make_dummy_pdf(2, "证据一微信回单")
    pdf2_data = _make_dummy_pdf(3, "证据二借据协议")
    img_data = _make_dummy_png()

    files = [
        ("files", ("01_wechat_receipt.pdf", pdf1_data, "application/pdf")),
        ("files", ("02_loan_contract.pdf", pdf2_data, "application/pdf")),
        ("files", ("03_id_card.png", img_data, "image/png")),
    ]

    resp = client.post(
        "/api/v1/phase2/process-evidence",
        data={"evidence_label": "证据"},
        files=files,
        headers=AUTH_HEADERS,
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["status"] == "success"
    assert data["total_pages"] == 6  # 2 + 3 + 1 = 6

    page_map = data.get("file_page_map")
    assert isinstance(page_map, list)
    assert len(page_map) == 3

    # 验证证据 1：第 1-2 页
    assert page_map[0]["filename"] == "01_wechat_receipt.pdf"
    assert page_map[0]["start_page"] == 1
    assert page_map[0]["end_page"] == 2
    assert page_map[0]["page_count"] == 2
    assert page_map[0]["page_range"] == "第1-2页"

    # 验证证据 2：第 3-5 页
    assert page_map[1]["filename"] == "02_loan_contract.pdf"
    assert page_map[1]["start_page"] == 3
    assert page_map[1]["end_page"] == 5
    assert page_map[1]["page_count"] == 3
    assert page_map[1]["page_range"] == "第3-5页"

    # 验证证据 3（单张图片）：第 6 页
    assert page_map[2]["filename"] == "03_id_card.png"
    assert page_map[2]["start_page"] == 6
    assert page_map[2]["end_page"] == 6
    assert page_map[2]["page_count"] == 1
    assert page_map[2]["page_range"] == "第6页"

    # 严格检验页码链式无间隙连续性
    for i in range(1, len(page_map)):
        assert page_map[i]["start_page"] == page_map[i - 1]["end_page"] + 1


# ==============================================================================
# 2. Evidence Catalog Word Generation with Synced Bates Pages
# ==============================================================================

def test_auto_catalog_with_file_page_map(client):
    """验证传入 file_page_map 时，生成的《证据目录.docx》与返回结果中准确填入绝对物理页码"""
    file_page_map = [
        {"filename": "转账电子回单.pdf", "page_range": "第1-2页"},
        {"filename": "借条借据原件.jpg", "page_range": "第3页"},
        {"filename": "催款微信记录.pdf", "page_range": "第4-8页"},
    ]

    payload = {
        "cause_of_action": "民间借贷纠纷",
        "plaintiff_name": "张三",
        "evidence_notes": "1. 微信转账回单\n2. 借据原件\n3. 催款微信聊天记录",
        "file_page_map": file_page_map,
    }

    resp = client.post("/api/v1/phase2/auto-catalog", json=payload, headers=AUTH_HEADERS)
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["status"] == "success"

    items = data["extracted_items"]
    assert len(items) == 3
    assert items[0]["pages"] == "第1-2页"
    assert items[1]["pages"] == "第3页"
    assert items[2]["pages"] == "第4-8页"

    # 检查生成的 docx 实际存在且可正常读取表格内容
    docx_name = data["catalog_filename"]
    docx_path = os.path.join(str(get_generated_dir()), docx_name)
    assert os.path.exists(docx_path)
    doc = Document(docx_path)
    assert len(doc.tables) > 0
    table_text = "".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
    assert "第1-2页" in table_text
    assert "第3页" in table_text
    assert "第4-8页" in table_text


def test_auto_catalog_with_structured_items(client):
    """验证直接传入 structured_items 时的保真度与排版生成"""
    structured_items = [
        {
            "name": "银行转账电子回单",
            "category": "电子数据",
            "source": "招商银行App导出打印件",
            "purpose": "证明原告向被告实际履行了50万元出借义务",
            "pages": "第1-3页",
        },
        {
            "name": "借条原件",
            "category": "书证",
            "source": "被告当面签署原件",
            "purpose": "证明原被告之间存在借贷合意及还款期限约定",
            "pages": "第4页",
        },
    ]

    payload = {
        "cause_of_action": "民间借贷纠纷",
        "plaintiff_name": "李四",
        "structured_items": structured_items,
    }

    resp = client.post("/api/v1/phase2/auto-catalog", json=payload, headers=AUTH_HEADERS)
    assert resp.status_code == 200, resp.text
    data = resp.json()
    items = data["extracted_items"]
    assert len(items) == 2
    assert items[0]["name"] == "银行转账电子回单"
    assert items[0]["pages"] == "第1-3页"
    assert items[1]["name"] == "借条原件"
    assert items[1]["pages"] == "第4页"


# ==============================================================================
# 3. Citation Guard: Historical Law Validity for Pre-2021 Events
# ==============================================================================

def test_citation_guard_pre_2021_historical_valid():
    """验证发生在 2021 年前的事实援引《合同法》时，判定为 HISTORICAL_VALID 而非已废止拦截"""
    text = "依据《中华人民共和国合同法》第一百一十四条之规定，被告应向原告承担逾期违约责任。"
    
    # 事实发生于 2019 年（民法典施行前）
    result_pre = audit_citations_in_text(raw_text=text, incident_date="2019-06-15")
    citations = result_pre.get("items", [])
    assert len(citations) == 1
    c = citations[0]
    assert c["status"] == "HISTORICAL_VALID"
    assert c["is_historical_valid"] is True
    assert "时间效力" in c["notes"] or "适用当时" in c["notes"]
    assert "民法典" in c["suggested_replacement"]


def test_citation_guard_post_2021_abrogated():
    """验证发生在 2021 年后的事实若错误援引旧《合同法》，正确标记为已废止拦截或自动净化"""
    text = "依据《中华人民共和国合同法》第一百一十四条之规定，被告应向原告承担违约责任。"

    # 事实发生于 2022 年（民法典已生效）
    result_post = audit_citations_in_text(raw_text=text, incident_date="2022-03-01", auto_sanitize=True)
    citations = result_post.get("items", [])
    assert len(citations) == 1
    c = citations[0]
    # 在 2021 后，合同法 114 条被自动映射替换为民法典第 585 条
    assert c.get("sanitized_to") == "《中华人民共和国民法典》第五百八十五条" or c["status"] == "ABROGATED_BLOCKED"


# ==============================================================================
# 4. Pipeline Stepper UI & A11Y Compliance
# ==============================================================================

def test_stepper_ui_dom_elements():
    """验证前端 index.html 与 main.js 具备完整的 Stepper 节点与导航绑定"""
    with open("frontend/index.html", "r", encoding="utf-8") as f:
        html = f.read()

    assert 'id="case-pipeline-stepper-wrap"' in html
    assert 'id="step-node-1"' in html
    assert 'id="step-node-2"' in html
    assert 'id="step-node-3"' in html
    assert 'id="step-node-4"' in html
    assert 'id="step-node-5"' in html
    assert 'navigateToPipelineStep(1)' in html
    assert 'navigateToPipelineStep(4)' in html

    # 验证 main.js 导出与导航逻辑
    with open("frontend/js/main.js", "r", encoding="utf-8") as f:
        main_js = f.read()

    assert "navigateToPipelineStep" in main_js
    assert "toggleStepperView" in main_js
    assert "updatePipelineStepperActive" in main_js

    # 验证 evidence.js 导出页码同步逻辑
    with open("frontend/js/modules/03_evidence.js", "r", encoding="utf-8") as f:
        ev_js = f.read()

    assert "syncBatesPagesToCatalog" in ev_js
    assert "_lastFilePageMap" in ev_js
