# -*- coding: utf-8 -*-
"""
LawFlow 智律工作台 - Top 5 核心改造项全量端到端实证测试
"""
from __future__ import annotations

import asyncio
import os
import tempfile
import uuid

import docx
import pandas as pd
import pytest

from gateway.utils.citation_guard import (
    audit_citations_in_text,
    verify_single_citation,
    get_official_max_article,
)
from gateway.utils.ocr_engine import extract_evidence_text
from gateway.utils.redline_builder import apply_in_place_docx_redline
from gateway.utils.file_converter_engine import convert_table_to_court_pdf
from gateway.utils.review_table_engine import (
    execute_cell_extraction,
    export_review_table_to_xlsx,
)
from gateway.routers.phase3_rag_sign import api_rag_search, RAGSearchReq


def test_priority1_citation_guard_and_calc():
    assert get_official_max_article("中华人民共和国民法典") == 1260
    assert get_official_max_article("民诉法") == 306
    assert get_official_max_article("劳动合同法") == 98

    text_with_high_articles = (
        "依据《中华人民共和国民法典》第一千零六十四条关于夫妻共同债务之规定，"
        "以及第一千一百六十五条侵权责任条文。"
    )
    audit_res = audit_citations_in_text(text_with_high_articles, auto_sanitize=True)
    assert audit_res["has_hallucinations"] is False
    assert "⚠️拦截到大模型捏造法条" not in audit_res["sanitized_text"]

    fake_text = "依据《中华人民共和国民法典》第一千二百六十五条及《中华人民共和国劳动合同法》第二百条。"
    fake_audit = audit_citations_in_text(fake_text, auto_sanitize=True)
    assert fake_audit["has_hallucinations"] is True

    single_valid = verify_single_citation("中华人民共和国民法典", 1064)
    assert single_valid["audit_status"] in [
        "LOCAL_MATCH",
        "UNINDEXED_UNVERIFIABLE",
        "PENDING_VERIFICATION",
    ]
    if single_valid["audit_status"] == "LOCAL_MATCH":
        assert single_valid["is_valid"] is True
    else:
        assert single_valid["is_valid"] is False

    single_fake = verify_single_citation("中华人民共和国民法典", 1300)
    assert single_fake["is_valid"] is False
    assert single_fake["audit_status"] == "HALLUCINATION_EXCEEDED"


def test_priority2_ocr_and_direct_upload(client, seed_demo):
    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = os.path.join(tmpdir, "test_evidence.docx")
        d = docx.Document()
        d.add_paragraph("这是一份原告与被告签署的居间协议书。")
        tbl = d.add_table(rows=2, cols=2)
        tbl.cell(0, 0).text = "协议编号"
        tbl.cell(0, 1).text = "XY-2026-001"
        tbl.cell(1, 0).text = "居间费用"
        tbl.cell(1, 1).text = "人民币 50,000 元"
        d.save(docx_path)

        docx_text = extract_evidence_text(docx_path)
        assert "居间协议书" in docx_text
        assert "50,000" in docx_text

        csv_path = os.path.join(tmpdir, "test_bank.csv")
        with open(csv_path, "w", encoding="utf-8-sig") as f:
            f.write("交易日期,对方账号,金额,摘要\n2023-05-01,6222021000111222333,200000.00,借款本金划转\n")

        csv_text = extract_evidence_text(csv_path)
        assert "6222021000111222333" in csv_text
        assert "借款本金划转" in csv_text

        m_resp = client.post(
            "/api/v1/vault/matters",
            json={
                "matter_id": f"MATTER-T5-{uuid.uuid4().hex[:12]}",
                "title": "测试多格式上传案件",
                "cause_of_action": "民间借贷纠纷",
                "client_role": "原告",
            },
        )
        assert m_resp.status_code == 200
        matter_id = m_resp.json()["matter"]["matter_id"]

        with open(docx_path, "rb") as f_up:
            up_resp = client.post(
                f"/api/v1/vault/matters/{matter_id}/upload-document",
                files={
                    "file": (
                        "test_evidence.docx",
                        f_up,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                data={"doc_type": "contract"},
            )
        assert up_resp.status_code == 200
        up_data = up_resp.json()
        assert up_data["status"] == "success"
        assert "居间协议书" in up_data["document"]["content_text"]


@pytest.mark.requires_ollama
def test_priority3_context_and_grounding(client, seed_demo):
    long_doc = (
        "前言说明 " * 500
        + "\n--- 第 5 页 ---\n第七条 双方约定因本协议引起的任何争议，均由原告住所地上海市浦东新区人民法院管辖。\n"
        + "后记补充 " * 500
    )
    res = asyncio.run(
        execute_cell_extraction(
            doc_title="长篇合资经营合同.pdf",
            doc_text=long_doc,
            col_name="约定管辖法院",
            instruction="提取约定的管辖法院或仲裁机构全称",
        )
    )
    assert res["status"] == "DONE"
    assert "法院" in res["value"] or "浦东" in res["value"]

    m_resp = client.post(
        "/api/v1/vault/matters",
        json={
            "matter_id": f"MATTER-T5-{uuid.uuid4().hex[:12]}",
            "title": "测试穿透定位案件",
            "cause_of_action": "买卖合同纠纷",
            "client_role": "原告",
        },
    )
    assert m_resp.status_code == 200, m_resp.text
    matter_id = m_resp.json()["matter"]["matter_id"]

    g_resp = client.get(
        f"/api/v1/vault/matters/{matter_id}/grounding/excerpt?doc_title=绝不存在的材料名&page_no=1"
    )
    assert g_resp.status_code == 200
    assert g_resp.json()["found"] is False


def test_priority4_redline_and_converter(client, seed_demo):
    with tempfile.TemporaryDirectory() as tmpdir:
        client_docx = os.path.join(tmpdir, "client_original_contract.docx")
        d = docx.Document()
        d.add_heading("第一条 违约责任", level=1)
        d.add_paragraph("乙方如逾期交付，每日应向甲方支付合同总价款5%的惩罚性违约金。")
        d.save(client_docx)

        out_redline = os.path.join(tmpdir, "in_place_redline.docx")
        clause_reviews = [
            {
                "original_clause": "乙方如逾期交付，每日应向甲方支付合同总价款5%的惩罚性违约金。",
                "suggested_clause": "乙方如逾期交付，每日应按未履约部分金额的0.05%向甲方支付违约金，总额以合同总价款的10%为限。",
            }
        ]
        apply_in_place_docx_redline(client_docx, clause_reviews, out_redline)
        assert os.path.exists(out_redline)

        d_check = docx.Document(out_redline)
        xml_content = d_check._body._element.xml
        assert "w:del" in xml_content or "w:ins" in xml_content

        csv_bank = os.path.join(tmpdir, "bank_flow.csv")
        df_bank = pd.DataFrame(
            [
                {
                    "交易时间": "2023-08-15 14:32:00",
                    "本方银行卡号": "6222021000987654321",
                    "对手方户名": "北京某科技有限公司分公司",
                    "交易流水号": "FLOW202308159988776655443322",
                    "转账金额": "1,500,000.00",
                }
            ]
        )
        df_bank.to_csv(csv_bank, index=False, encoding="utf-8-sig")
        pdf_out_name = "test_bank_court_table.pdf"
        res_pdf = convert_table_to_court_pdf(csv_bank, pdf_out_name)
        assert res_pdf["status"] == "success"

        tbl_resp = client.post(
            "/api/v1/vault/matters/MATTER-2026-001/review-tables",
            json={"title": "测试导出表"},
        )
        assert tbl_resp.status_code == 200, tbl_resp.text
        tbl_id = tbl_resp.json()["table"]["table_id"]
        xlsx_bytes = export_review_table_to_xlsx(tbl_id)
        assert xlsx_bytes is not None and len(xlsx_bytes) > 1000


def test_priority5_workflow_and_rag():
    rag_res = asyncio.run(
        api_rag_search(
            RAGSearchReq(
                query="借条丢失 微信聊天记录 借贷合意 职务行为",
                corpus_type="all",
                top_k=3,
            )
        )
    )
    assert rag_res["status"] == "success"
    assert len(rag_res["statutes"]) >= 1
    assert any("民法典" in s["citation"] or "民间借贷" in s["citation"] for s in rag_res["statutes"])
