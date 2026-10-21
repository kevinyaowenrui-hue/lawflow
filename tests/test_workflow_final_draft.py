# -*- coding: utf-8 -*-
"""批次7：S-12 工作流终稿从案件数据渲染。"""
from __future__ import annotations

import os

import docx
import pytest


def _docx_full_text(path: str) -> str:
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


@pytest.fixture
def no_ollama(monkeypatch):
    from gateway.core.ollama_client import ollama_client

    async def _offline(*args, **kwargs):
        return False

    monkeypatch.setattr(ollama_client, "_is_ollama_online", _offline)
    monkeypatch.setattr(
        ollama_client,
        "settings",
        {**ollama_client.settings, "provider": "ollama", "api_key": ""},
    )


def test_S12_workflow_final_draft_from_matter(client, no_ollama):
    # 建案
    create = client.post(
        "/api/v1/vault/matters",
        json={
            "title": "测试原告甲诉测试被告乙借款案",
            "cause_of_action": "民间借贷纠纷",
            "court_name": "北京市朝阳区人民法院",
            "plaintiff_name": "测试原告甲",
            "defendant_name": "测试被告乙",
            "principal_amount": 100000.0,
            "summary": "被告借款未还。",
        },
    )
    assert create.status_code == 200, create.text
    matter = create.json()["matter"]
    matter_id = matter["matter_id"]

    start = client.post(
        "/api/v1/vault/workflows/start",
        json={
            "matter_id": matter_id,
            "initial_facts": "原告出借10万元，被告未还。",
            "cause_of_action": "民间借贷纠纷",
        },
    )
    assert start.status_code == 200, start.text
    wf = start.json()["workflow"]
    assert wf.get("matter_id") == matter_id
    wf_id = wf["workflow_id"]

    # Step2 → Step4
    adv2 = client.post(
        f"/api/v1/vault/workflows/{wf_id}/advance",
        json={
            "human_decisions": {
                "confirmed_claims": ["判令被告返还借款 10 万元"],
                "confirmed_statutes": ["《中华人民共和国民法典》第六百七十五条"],
                "cause_of_action": "民间借贷纠纷",
            }
        },
    )
    assert adv2.status_code == 200, adv2.text
    assert adv2.json()["workflow"].get("matter_id") == matter_id

    # Step4 → COMPLETED
    adv4 = client.post(
        f"/api/v1/vault/workflows/{wf_id}/advance",
        json={"human_decisions": {"accepted_defenses": ["时效中断"], "lawyer_defense_notes": "无"}},
    )
    assert adv4.status_code == 200, adv4.text
    data = adv4.json()
    wf_done = data["workflow"]
    assert wf_done.get("matter_id") == matter_id
    assert wf_done.get("current_step") == 5 or wf_done.get("is_completed")

    step5 = (wf_done.get("step_data") or {}).get("step5") or {}
    docx_path = step5.get("docx_path")
    if not docx_path:
        from gateway.core.paths import get_generated_dir
        name = step5.get("docx_filename")
        assert name, step5
        docx_path = str(get_generated_dir() / name)
    assert os.path.isfile(docx_path), docx_path
    full = _docx_full_text(docx_path)
    assert "测试原告甲" in full
    assert "测试被告乙" in full
    assert "判令被告返还借款" in full
    assert full.count("此致") == 1
    assert "张三诉李四" not in full
    assert "海淀区" not in full
    assert "人民法院人民法院" not in full


def test_S12_missing_claims_returns_422(client, no_ollama):
    create = client.post(
        "/api/v1/vault/matters",
        json={
            "title": "缺诉请案",
            "cause_of_action": "民间借贷纠纷",
            "court_name": "北京市朝阳区人民法院",
            "plaintiff_name": "测试原告甲",
            "defendant_name": "测试被告乙",
            "summary": "无诉请",
        },
    )
    matter_id = create.json()["matter"]["matter_id"]
    start = client.post(
        "/api/v1/vault/workflows/start",
        json={"matter_id": matter_id, "initial_facts": "事实。"},
    )
    wf_id = start.json()["workflow"]["workflow_id"]

    # 推进到 step4，但诉请为空
    adv2 = client.post(
        f"/api/v1/vault/workflows/{wf_id}/advance",
        json={
            "human_decisions": {
                "confirmed_claims": [],
                "confirmed_statutes": [],
                "cause_of_action": "民间借贷纠纷",
            }
        },
    )
    assert adv2.status_code == 200, adv2.text

    adv4 = client.post(
        f"/api/v1/vault/workflows/{wf_id}/advance",
        json={"human_decisions": {"accepted_defenses": []}},
    )
    assert adv4.status_code == 422, adv4.text
    detail = adv4.json().get("detail")
    if isinstance(detail, dict):
        assert "claims" in (detail.get("missing_fields") or [])
    else:
        assert "claims" in str(detail) or "诉请" in str(detail) or "missing" in str(detail).lower()
